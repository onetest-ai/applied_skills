#!/usr/bin/env python3
"""Deterministic structure profiler for the numeric lane (relation-schema design).

Walks a folder of Excel workbooks and, using openpyxl read_only (constant memory,
never loads the whole 100MB+ book), emits a compact JSON structure map:
per workbook -> per sheet -> detected header row, column names, row-count estimate,
and a couple of sample rows. This is the "know the schema" input for designing
conformed dimensions + a governed metric semantic layer. It reads STRUCTURE, not
the full data — the numeric source of truth stays openpyxl->Parquet->DuckDB.

Usage: profile_workbooks.py --root <dir> --out <file> [--formats xlsx,xlsm]
       [--sample-rows 3] [--scan-rows 25]  (scan-rows = rows searched for a header)
"""
import argparse, json, os, sys, warnings
warnings.filterwarnings("ignore")

def looks_like_header(row):
    vals = [c for c in row if c not in (None, "")]
    if len(vals) < 2:
        return False
    strs = sum(1 for c in vals if isinstance(c, str))
    return strs >= max(2, int(0.6 * len(vals)))

def profile_sheet(ws, sample_rows, scan_rows):
    header_idx, header = None, None
    buffered = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        buffered.append(row)
        if header is None and i < scan_rows and looks_like_header(row):
            header_idx, header = i, [("" if c is None else str(c)).strip() for c in row]
        if len(buffered) >= scan_rows + sample_rows + 1:
            break
    samples = []
    if header_idx is not None:
        for r in buffered[header_idx + 1: header_idx + 1 + sample_rows]:
            samples.append([("" if c is None else str(c)) for c in r])
    return {
        "dims": ws.dimensions,
        "max_row": ws.max_row, "max_col": ws.max_column,
        "header_row_index": header_idx,
        "columns": [h for h in (header or []) if h],
        "sample_rows": samples,
    }

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--formats", default="xlsx,xlsm")
    ap.add_argument("--sample-rows", type=int, default=3)
    ap.add_argument("--scan-rows", type=int, default=25)
    a = ap.parse_args()
    allow = {"." + e.strip().lower().lstrip(".") for e in a.formats.split(",")}
    import openpyxl

    out = []
    for root, _, files in os.walk(a.root):
        for fn in sorted(files):
            if fn.startswith(("~$", ".")):
                continue
            if os.path.splitext(fn)[1].lower() not in allow:
                continue
            src = os.path.join(root, fn)
            rel = os.path.relpath(src, a.root)
            rec = {"source": rel, "size_mb": round(os.path.getsize(src)/1e6, 2), "sheets": {}}
            try:
                wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    try:
                        rec["sheets"][ws.title] = profile_sheet(ws, a.sample_rows, a.scan_rows)
                    except Exception as e:
                        rec["sheets"][ws.title] = {"error": str(e)}
                wb.close()
                print(f"[ok] {rec['size_mb']:>7} MB  {len(rec['sheets']):>2} sheets  {rel}", file=sys.stderr)
            except Exception as e:
                rec["error"] = str(e)
                print(f"[ERR] {rel}: {e}", file=sys.stderr)
            out.append(rec)
    json.dump(out, open(a.out, "w"), indent=2)
    print(f"\nprofiled {len(out)} workbooks -> {a.out}", file=sys.stderr)

if __name__ == "__main__":
    main()
