"""Split a draft .pptx into slide sections and word-bounded batches for /kb:deck-fact-check.

    python pptx_sections.py <draft.pptx> --out <dir> [--max-words 1500]
    python pptx_sections.py --check        (dependency check for step 0)

Reads the deck two ways and writes, into <dir>, the same files sections.py writes for a .docx (section_io.py,
shared byte-identical with doc-fact-check, writes them):
- python-pptx shape walk (the text and the anchors): one section per visible slide, "Slide N: <title>"; every
  non-empty paragraph of a text shape is a `p` block, p_id s<slide>-sh<shape_id>-p<k>; every table row an `r`
  block, r_id s<slide>-sh<shape_id>-r<k>. Shapes inside groups are walked, with their slide position. Each block
  carries `pos` ({x, y, w, h} in EMU, the shape box a comment is pinned to) and, when the slide supplies it,
  `ctx` (`parent` bullet, table `header` and `row_label`, the slide's `footnote` source line).
- render_pages.py (vendored byte-identical from visual-parse): soffice → PDF → one PNG and text layer per
  visible slide, and its dense-drawing flag (its thin-text flag is not used: most slides are short). A slide is
  flagged when it holds a picture, a chart or a diagram, or when its PDF text has words the shape walk cannot see (chart labels, SmartArt): any risk-tagged token, or 15 or more
  words (R8); letter-spacing fragments and pieces of words the slide holds are not such words. A flagged slide gains a figure record {figure: N, p_id: "s<N>", image, media, why}; its PNG is
  copied to media/slide_NN.png for the main session's vision step.

Only what is on the slide is analysed. Hidden slides are not in the PDF, so they are skipped and listed in
stats.json (`hidden_slides`); pages map to slides by visible index. A shape wholly outside the slide or hidden in
the Selection Pane, and any paragraph or row that is not on the rendered slide (pushed past its edge, covered), is
left out entirely: not a block, not context, not the document date. stats.json `offscreen` {blocks, slides,
by_reason} reports what was left out. Speaker notes, comments, alt text and embedded workbooks are never read. Navigation is not read either: a text shape at the same place on 3 or more visible slides (tab bars, section eyebrows); stats.json `chrome` {texts, slides} lists it. A "<source>; <time or date>" line under a statement is that statement's ctx `attribution`, not a block. A flagged slide's figure record lists its `elements` (pictures, charts, diagrams: shape id, kind, name, box), one of which each figure claim names. stats.json also lists `flagged_slides`. The batch files carry `doc_date`: a date on
the title slide, else the file's modified date (marked inferred).

Refuses (exit 2; an earlier run's files are left as they were) a .ppt or any non-.pptx file, an unreadable package,
a missing soffice, a failed render, or a PDF whose page count does not match the visible slides. Date, footer and
slide-number placeholders are furniture: never blocks, never context. A footnote is a short bottom text box set at
most 12 pt or reading as a source / as-of line.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # sibling imports work under python -I too

import render_pages  # noqa: E402
from fact_check_invariants import (MIN_PICTURE_AREA, is_footer_placeholder, is_hidden, is_picture,  # noqa: E402
                                   slide_kinds, slide_title, walk)
import chunk_claims  # noqa: E402
import section_io  # noqa: E402
from pptx_visibility import RenderedText, normalized  # noqa: E402
from section_io import _words, split_sentences, tag_risk  # noqa: E402

FOOTNOTE_ZONE = 0.85            # a text shape whose top is below 85% of the slide height is a footnote line
FOOTNOTE_MAX_WORDS = 25
FOOTNOTE_MAX_PT = 12            # a footnote is set smaller than body text …
_SOURCE_LINE = re.compile(r"^\s*(?:sources?|as of|data|note|notes)\b", re.I)   # … or reads as a source / as-of line
_LIST_MARKER = re.compile(r"^\s*(?:\d{1,2}|[a-z])[.)](?!\d)\s*", re.I | re.M)       # auto-numbering, rendered but not text
RESIDUAL_WORDS = 15             # R8: this many words of PDF text the walk cannot see flags the slide
RESIDUAL_MIN_LETTERS = 3        # shorter letter-only tokens are letter-spacing or split words, not unseen text
_LABEL = re.compile(r"^(?:\d{1,2}\s+)?[A-Z][A-Z&/()\-.,'’ ]*$")   # "2 ORDER PLACED", "COMMERCE & IPX (INTEGRATION)"
LABEL_MAX_WORDS = 6
LEDE_GAP = 1097280              # 1.2 in: a lede sits this close under its slide's title
_BADGE = re.compile(r"^\s*[#(]?\d{1,2}[.)]?\s*$")   # a bare step number or badge ("3", "10.", "(2)"): nothing to check
CHROME_SLIDES = 3               # text at the same place on this many visible slides is navigation (page furniture)
CHROME_TOL = 228600             # 0.25 in: "the same place"
_TIME = r"\d{1,2}:\d{2}(?::\d{2})?"
_MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
_SOURCE_WORDS = re.compile(r"\b(?:session|meeting|workshop|interview|call|discussion|readout|review|survey|"
                           r"transcript|report|study|deck|memo|sync|stand-?up|alignment|roadmap|kick-?off|briefing|w/)\b",
                           re.I)   # what an attribution or a dateline names
_ATTRIBUTION = re.compile(r"^[^;\d%]{3,90};\s*(?:" + _TIME + r"(?:\s*[–—-]\s*" + _TIME + r")?"
                          r"|" + _MONTH + r"\s+\d{1,2}(?:,?\s+\d{4})?|\d{1,2}\s+" + _MONTH + r"(?:\s+\d{4})?"
                          r"|\d{4}-\d{2}-\d{2})\s*$", re.I)
_FULL_DATE = (r"(?:\d{1,2}\s+" + _MONTH + r",?\s+\d{4}|" + _MONTH + r"\s+\d{1,2},?\s+\d{4}|\d{4}-\d{2}-\d{2})")
_DATELINE = re.compile(r"^(?P<name>[^,;\d%]{3,80}),\s*" + _FULL_DATE + r"\s*$", re.I)
_TOKEN = re.compile(r"\w[\w.,%]*\w|\w")
_DATE = re.compile(r"\b(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.? \d{4}"
                   r"|\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4})\b", re.I)

NEEDS = ("deck-fact-check needs Python with python-pptx and PyMuPDF, and LibreOffice (soffice), in this "
         "environment; at least one is missing here. Run it in Claude Code with them installed.")


class Refusal(Exception):
    pass


def _is_title_slide(slide) -> bool:
    from pptx.enum.shapes import PP_PLACEHOLDER
    return any(sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.SUBTITLE)
               for sh in slide.shapes)


def _small_text(shape) -> bool:
    """Every sized run of the shape is at most FOOTNOTE_MAX_PT (runs without an explicit size are not evidence)."""
    sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size is not None]
    return bool(sizes) and max(sizes) <= FOOTNOTE_MAX_PT


def _squash(text: str) -> str:
    """Letters and digits only, case-folded: a letter-spaced label ("C U R R E N T") still matches its text."""
    return normalized(str(text))


def on_page(text: str, page_sq: str | None) -> bool | None:
    """Require the complete text, including short values; None means no rendered evidence was supplied."""
    t = _squash(text)
    if page_sq is None:
        return None
    return bool(t) and t in page_sq


def row_on_page(cells: list[str], page_sq: str | None) -> bool | None:
    """A row cannot carry an invisible cell into verification."""
    judged = [on_page(c, page_sq) for c in cells if c.strip()]
    return None if not judged or any(v is None for v in judged) else all(judged)


def shape_hidden(shape) -> bool:
    """The shape, or a group around it, carries hidden="1" (Selection Pane: hidden)."""
    el = shape._element
    node = el
    while node is not None and not node.tag.endswith("}spTree"):
        if node is el or node.tag.endswith("}grpSp"):
            c = next((x for x in node.iter() if x.tag.endswith("}cNvPr")), None)
            if c is not None and c.get("hidden") in ("1", "true"):
                return True
        node = node.getparent()
    return False


def outside(pos: dict, width: int, height: int) -> bool:
    """The shape's box lies wholly outside the slide (parked beside it as a working note)."""
    return pos["x"] + pos["w"] <= 0 or pos["x"] >= width or pos["y"] + pos["h"] <= 0 or pos["y"] >= height


def _off_reason(shape, pos, width, height) -> str | None:
    if shape_hidden(shape):
        return "hidden shape"
    if outside(pos, width, height):
        return "outside the slide"
    return None


def _text_units(shape) -> int:
    if shape.has_text_frame:
        return sum(1 for p in shape.text_frame.paragraphs if p.text.strip())
    if getattr(shape, "has_table", False) and shape.has_table:
        return len(shape.table.rows)
    return 0


def inherited_text(slide, number: int, visible_index: int) -> str:
    """Text the rendered slide shows but the slide itself does not hold: its layout's and master's text (footers,
    logos with words, decoration) and the slide number. Not claims of this slide, so never residual (R8)."""
    texts = [str(number), str(visible_index)]
    for shapes in (slide.slide_layout.shapes, slide.slide_layout.slide_master.shapes):
        texts += [shape.text_frame.text for shape, _ in walk(shapes) if shape.has_text_frame]
    return " ".join(texts)


def _doc_date(prs, page_sq: str | None = None, visibility=None) -> dict | None:
    """A date written on the deck's title slide (slide 1 laid out as a title slide and shown; a quarter is a period,
    not a date), else the package's modified date (inferred). Only on-screen text counts."""
    width, height = prs.slide_width or 9144000, prs.slide_height or 6858000
    if len(prs.slides) and not is_hidden(prs.slides[0]) and _is_title_slide(prs.slides[0]):
        for shape, pos in walk(prs.slides[0].shapes):
            if not shape.has_text_frame or _off_reason(shape, pos, width, height):
                continue
            if (m := _DATE.search(shape.text_frame.text)) and (
                    visibility(1, m.group(0), pos) if visibility else on_page(m.group(0), page_sq) is not False):
                return {"value": m.group(0), "inferred": False}
    return section_io.core_date(prs.core_properties)


_CONNECTOR = re.compile(r"\b(?:to|via|into|through|from)\b|->|→", re.I)   # "ERP TO ROUTEDB VIA API" states a flow


def is_label(text: str) -> bool:
    """A step badge ("3") or a structural label: upper-case words only, at most LABEL_MAX_WORDS, naming a part of the
    slide ("2 ORDER PLACED", "COMMERCE & IPX (INTEGRATION)", "WHERE IT CAN BREAK"). It carries no risk tag (spec §11
    Y2). A leading number is a step only when two or more words follow it ("3 REGIONS" is a count); nothing else may
    be a digit, and the words may carry no risk tag but `topo` from system names listed without a connector, so "ALL
    ORDERS SHIP SAME DAY", "ERP OWNS BILLING", "TWO DATA CENTRES" and "ERP TO ROUTEDB VIA API" stay claims."""
    t = " ".join(str(text).split())
    if _BADGE.match(t):
        return True
    if not _LABEL.match(t) or len(t.split()) > LABEL_MAX_WORDS or sum(c.isalpha() for c in t) < 2:
        return False
    m = re.match(r"^\d{1,2}\s+(\S+\s+\S.*)$", t)
    core = m.group(1) if m else t
    if re.search(r"\d", core) or _CONNECTOR.search(core):
        return False
    return set(tag_risk(core)) <= {"topo"}


def _eyebrow(text: str) -> bool:
    """An eyebrow or breadcrumb over a title: mostly capitals, or parts joined by separators."""
    letters = [c for c in text if c.isalpha()]
    return bool(letters) and (sum(c.isupper() for c in letters) / len(letters) >= 0.8 or bool(re.search(r"\s[·|>—–]\s", text)))


def slide_context(shapes: list, title_shape, height: int, footnote: list[str], retained: dict[int, list[str]] | None = None) -> dict:
    """What a claim on this slide is about (spec §11 Y1): the title (the title placeholder, else the topmost text
    shape of 3+ words in the top 30% of the slide), the lede (the next text shape under it within LEDE_GAP, 6+ words)
    and the key (the slide's footnote lines)."""
    texts = [(sh, pos, " ".join((retained.get(sh.shape_id, []) if retained is not None else [sh.text_frame.text])))
             for sh, pos in shapes if sh.has_text_frame]
    texts = [x for x in texts if x[2]]
    title = next((x for x in texts if title_shape is not None and x[0].shape_id == title_shape.shape_id), None)
    if title is None:
        top = sorted((x for x in texts if x[1]["y"] < 0.3 * height and len(x[2].split()) >= 3
                      and not is_label(x[2]) and not _eyebrow(x[2])), key=lambda x: x[1]["y"])
        title = top[0] if top else None
    out = {}
    if title:
        out["title"] = title[2]
        bottom = title[1]["y"] + title[1]["h"]
        below = sorted((x for x in texts if x is not title and bottom - LOCATE_TOL <= x[1]["y"] <= bottom + LEDE_GAP
                        and len(x[2].split()) >= 6), key=lambda x: x[1]["y"])
        if below:
            out["lede"] = below[0][2]
    if footnote:
        out["key"] = "; ".join(footnote)
    return out


def is_dateline(text: str) -> bool:
    """"<session or document name>, <full date>" — where the slide's content comes from, not a statement about it.
    The name must name a source (session, workshop, review, roadmap…), so "Go-live, 15 September 2026" stays a claim."""
    parts = [x for x in re.split(r"(?<=\d)\.\s*|;\s*", " ".join(str(text).split())) if x.strip()]
    return bool(parts) and all(_one_dateline(x) for x in parts)    # a line may name several sources, each a dateline


def _one_dateline(t: str) -> bool:
    t = t.strip().rstrip(".")
    m = _DATELINE.match(t)
    return bool(m) and len(t.split()) <= 14 and bool(_SOURCE_WORDS.search(m.group("name")))


def is_attribution(text: str) -> bool:
    """"<source>; <timestamp, range or date>": where the statement above it was said or written (spec §10 X9)."""
    t = " ".join(str(text).split())
    return bool(_ATTRIBUTION.match(t)) and bool(_SOURCE_WORDS.search(t.split(";")[0]))


def chrome_shapes(slides: list[tuple[int, list]], visibility=None) -> tuple[set, dict]:
    """Navigation: a text shape whose text sits at the same place (±CHROME_TOL) on CHROME_SLIDES or more visible
    slides (tab bars, section eyebrows, scenario pickers). Returns ({(slide, shape_id)}, {text: [slides]})."""
    seen: dict[str, list[tuple[int, int, int, int]]] = {}
    for n, shapes in slides:
        for sh, pos in shapes:
            if sh.has_text_frame and (key := " ".join(p.text for p in sh.text_frame.paragraphs
                                                     if not visibility or visibility(n, p.text, pos)).strip()):
                seen.setdefault(key, []).append((n, sh.shape_id, pos["x"], pos["y"]))
    out, texts = set(), {}
    for key, occ in seen.items():
        for n, sid, x, y in occ:
            near = {m for m, _, x2, y2 in occ if abs(x2 - x) <= CHROME_TOL and abs(y2 - y) <= CHROME_TOL}
            if len(near) >= CHROME_SLIDES:
                out.add((n, sid))
                texts.setdefault(key, set()).add(n)
    return out, {k: sorted(v) for k, v in texts.items()}


def slide_elements(slide, slide_area: int) -> list[dict]:
    """The visual elements a figure claim can name: pictures of MIN_PICTURE_AREA or more, charts, SmartArt diagrams and
    embedded objects (spec §10 X2, X10)."""
    out = []
    for shape, pos in walk(slide.shapes):
        if getattr(shape, "has_chart", False) and shape.has_chart:
            kind = "chart"
        elif shape._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds") is not None:
            kind = "diagram"
        elif is_picture(shape) and pos["w"] * pos["h"] >= MIN_PICTURE_AREA * slide_area:
            kind = "picture"
        elif shape._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}oleObj") is not None:
            kind = "object"                      # an embedded workbook, drawing or document shown on the slide
        else:
            continue
        out.append({"shape_id": shape.shape_id, "kind": kind, "name": shape.name,
                    "box": {k: int(pos[k]) for k in ("x", "y", "w", "h")}})
    return out


def visible_slides(pptx_path) -> list[int]:
    """Slide numbers that are not hidden: the slides the PDF renders, in page order."""
    from pptx import Presentation
    return [n for n, slide in enumerate(Presentation(str(pptx_path)).slides, 1) if not is_hidden(slide)]


LOCATE_TOL = 91440              # 0.1 in: a statement's rendered position may sit this far outside its shape box


def _locate(locate, n: int, text: str, pos: dict, after_y: int) -> dict | None:
    """Where a kept statement renders: the first hit of its first words inside its shape's box, at or below the
    previous statement of the same shape (document order). None when it cannot be found."""
    if locate is None:
        return None
    words = text.split()
    for take in (5, 3, 1):
        if take > len(words) and take != 1:
            continue
        needle = " ".join(words[:take])
        if len(_squash(needle)) < 3 and take < len(words):    # a short needle is only safe as the whole statement
            continue
        hits = [h for h in locate(n, needle)
                if pos["x"] - LOCATE_TOL <= h[0] <= pos["x"] + pos["w"] + LOCATE_TOL
                and pos["y"] - LOCATE_TOL <= h[1] <= pos["y"] + pos["h"] + LOCATE_TOL and h[1] >= after_y - LOCATE_TOL]
        if hits:
            x, y = min(hits, key=lambda h: (h[1], h[0]))
            return {"x": int(min(max(x, pos["x"]), pos["x"] + pos["w"])), "y": int(min(max(y, pos["y"]), pos["y"] + pos["h"]))}
    return None


def read_deck(pptx_path, page_text: dict[int, str] | None = None, locate=None, visibility=None) -> dict:
    """The shape walk, on-screen text only. Returns {sections, hidden, visible, kinds, walk_text, doc_date, offscreen}.
    A shape wholly outside the slide or hidden is skipped; with page_text (slide -> its rendered text) a paragraph or
    row that is not on the rendered slide (pushed past the edge, covered) is skipped too. Skipped text is never a
    block, never context and never the document date; `offscreen` lists it per shape {slide, shape, reason, blocks}.
    With locate(slide, words) -> [(x, y) in slide EMU], every kept paragraph and row records `at`: where it renders.
    walk_text is what the slide's render may show without a vision model: its own text plus inherited_text()."""
    from pptx import Presentation
    prs = Presentation(str(pptx_path))
    height, width = prs.slide_height or 6858000, prs.slide_width or 9144000
    sections, hidden, visible, kinds, walk_text, offscreen, elements = [], [], [], {}, {}, [], {}
    on_slide = [(n, [(sh, pos) for sh, pos in walk(slide.shapes)
                     if not is_footer_placeholder(sh) and not _off_reason(sh, pos, width, height)])
                for n, slide in enumerate(prs.slides, 1) if not is_hidden(slide)]
    chrome, chrome_texts = chrome_shapes(on_slide, visibility)
    for n, slide in enumerate(prs.slides, 1):
        if is_hidden(slide):
            hidden.append(n)
            continue
        visible.append(n)
        page_sq = _squash(page_text[n]) if page_text and n in page_text else None

        def shown(text, pos, preceding=""):
            return visibility(n, text, {**pos, "preceding": preceding}) if visibility else on_page(text, page_sq) is not False
        title_shape = None
        try:
            title_shape = slide.shapes.title
        except Exception:
            pass
        blocks, texts, footnote = [], [], []
        shapes = list(walk(slide.shapes))
        furniture = [sh.text_frame.text for sh, _ in shapes if sh.has_text_frame and is_footer_placeholder(sh)]
        shapes = [(sh, pos) for sh, pos in shapes if not is_footer_placeholder(sh)]   # date, footer, slide number
        kept = []
        for sh, pos in shapes:             # not on the slide at all: never read
            why = _off_reason(sh, pos, width, height)
            if why:
                if _text_units(sh):
                    offscreen.append({"slide": n, "shape": sh.shape_id, "reason": why, "blocks": _text_units(sh)})
            else:
                kept.append((sh, pos))
        shapes = kept
        title_on = (title_shape is not None and any(sh.shape_id == title_shape.shape_id for sh, _ in shapes)
                    and shown(slide_title(slide), next(pos for sh, pos in shapes if sh.shape_id == title_shape.shape_id)))
        title = slide_title(slide) if title_on else ""
        nav = [sh.text_frame.text for sh, _ in shapes if (n, sh.shape_id) in chrome]
        shapes = [(sh, pos) for sh, pos in shapes if (n, sh.shape_id) not in chrome]   # navigation: never read
        for shape, pos in shapes:           # the slide's source / as-of line, context for every claim on it
            if (shape.has_text_frame and pos["y"] >= FOOTNOTE_ZONE * height
                    and 0 < _words(shape.text_frame.text) <= FOOTNOTE_MAX_WORDS
                    and (title_shape is None or shape.shape_id != title_shape.shape_id)
                    and (_small_text(shape) or _SOURCE_LINE.match(shape.text_frame.text))
                    and shown(shape.text_frame.text, pos)):
                footnote.append(" ".join(shape.text_frame.text.split()))
            elif (shape.has_text_frame and is_dateline(shape.text_frame.text)    # "<session>, <date>": context too
                    and shown(shape.text_frame.text, pos)):
                footnote.append(" ".join(shape.text_frame.text.split()))
        dropped: dict[int, int] = {}
        heading_p_id = None
        retained: dict[int, list[str]] = {}
        for shape, pos in shapes:
            base = f"s{n}-sh{shape.shape_id}"
            last_y = -LOCATE_TOL
            if shape.has_text_frame:
                levels: dict[int, str] = {}
                for k, para in enumerate(shape.text_frame.paragraphs, 1):
                    text = " ".join(para.text.split())
                    if not text:
                        continue
                    preceding = " ".join(p.text for p in shape.text_frame.paragraphs[:k - 1])
                    if not shown(text, pos, preceding):
                        dropped[shape.shape_id] = dropped.get(shape.shape_id, 0) + 1
                        continue
                    texts.append(text)
                    retained.setdefault(shape.shape_id, []).append(text)
                    if is_dateline(text):           # the slide's source line: in its context (footnote), not a claim
                        continue
                    if (is_attribution(text) and blocks and blocks[-1][0] == "p"
                            and blocks[-1][1]["p_id"].startswith(base + "-")
                            and "attribution" not in blocks[-1][1].get("ctx", {})):
                        blocks[-1][1].setdefault("ctx", {})["attribution"] = text   # where the statement above was said
                        continue
                    p_id = f"{base}-p{k}"
                    if title_shape is not None and shape.shape_id == title_shape.shape_id and heading_p_id is None:
                        heading_p_id = p_id
                    ctx = {}
                    level = para.level or 0
                    parent = next((levels[lv] for lv in sorted(levels, reverse=True) if lv < level), None)
                    if parent:
                        ctx["parent"] = parent
                    levels = {lv: t for lv, t in levels.items() if lv < level}
                    levels[level] = text
                    if footnote and text not in footnote:
                        ctx["footnote"] = "; ".join(footnote)
                    sents = [{"s_id": f"{p_id}s{j}", "text": s, "risk": [] if is_label(s) else tag_risk(s)}
                             for j, s in enumerate(split_sentences(text), 1)]
                    rec = {"p_id": p_id, "text": text, "sentences": sents, "pos": pos}
                    if at := _locate(locate, n, text, pos, last_y):
                        rec["at"], last_y = at, at["y"]
                    if ctx:
                        rec["ctx"] = ctx
                    blocks.append(("p", rec, _words(text)))
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = list(shape.table.rows)
                header = [" ".join(c.text.split()) for c in rows[0].cells] if rows else []
                row_y = pos["y"]
                scale_y = pos["h"] / sum(r.height for r in rows) if rows else 1
                scale_x = pos["w"] / sum(c.width for c in shape.table.columns)
                header_visible = True
                for k, row in enumerate(rows, 1):
                    cells = [" ".join(c.text.split()) for c in row.cells]
                    cell_x = pos["x"]
                    cell_checks = []
                    for c, column in zip(row.cells, shape.table.columns):
                        box = {"x": cell_x, "y": row_y, "w": column.width * scale_x, "h": row.height * scale_y}
                        cell_checks.append(not c.text.strip() or shown(c.text, box))
                        cell_x += box["w"]
                    row_y += row.height * scale_y
                    if k == 1:
                        header_visible = all(cell_checks)
                    if not header_visible or not all(cell_checks):
                        dropped[shape.shape_id] = dropped.get(shape.shape_id, 0) + 1
                        continue
                    texts.extend(cells)
                    r_id = f"{base}-r{k}"
                    tags = tag_risk(" | ".join(cells))
                    if k == 1 and len(rows) > 1:            # a header row says what the figures measure
                        tags = tags + ["label"]
                    rec = {"t_id": base, "r_id": r_id, "cells": cells, "s_id": r_id, "risk": tags, "pos": pos}
                    first = next((c for c in cells if c), "")
                    if first and (at := _locate(locate, n, first, pos, last_y)):
                        rec["at"], last_y = at, at["y"]
                    ctx = {}
                    if k > 1 and any(header):
                        ctx.update(header=header, row_label=cells[0] if cells else "")
                    if footnote:
                        ctx["footnote"] = "; ".join(footnote)
                    if ctx:
                        rec["ctx"] = ctx
                    blocks.append(("r", rec, sum(_words(c) for c in cells)))
        offscreen += [{"slide": n, "shape": sid, "reason": "not shown on the slide", "blocks": c} for sid, c in dropped.items()]
        context = slide_context(shapes, title_shape, height, footnote, retained)
        sections.append({"section_id": f"s{n:02d}", "section": f"Slide {n}: {title}" if title else f"Slide {n}",
                         "slide_context": context,
                         "heading_path": [f"Slide {n}", title] if title else [f"Slide {n}"], "level": 1,
                         "heading_p_id": heading_p_id, "blocks": blocks,
                         "words": sum(w for *_, w in blocks)})
        kinds[n] = slide_kinds(slide, width * height)
        elements[n] = slide_elements(slide, width * height)
        walk_text[n] = " ".join(texts + furniture + nav) + " " + inherited_text(slide, n, len(visible))
    first = _squash(page_text[1]) if page_text and 1 in page_text else None
    return {"sections": sections, "hidden": hidden, "visible": visible, "kinds": kinds, "walk_text": walk_text,
            "doc_date": _doc_date(prs, first, visibility), "offscreen": offscreen, "elements": elements,
            "chrome": {"texts": sorted(chrome_texts), "slides": sorted({m for v in chrome_texts.values() for m in v})}}


def offscreen_summary(offscreen: list[dict]) -> dict:
    """{blocks, slides, by_reason} for stats.json, run.json and the coverage line."""
    by_reason: dict[str, int] = {}
    for o in offscreen:
        by_reason[o["reason"]] = by_reason.get(o["reason"], 0) + o["blocks"]
    return {"blocks": sum(by_reason.values()), "slides": sorted({o["slide"] for o in offscreen}), "by_reason": by_reason}


def residual_tokens(page_text: str, walk_text: str) -> list[str]:
    """Tokens of a slide's PDF text layer that the shape walk did not return (order kept, case-folded). Letter-only
    tokens shorter than RESIDUAL_MIN_LETTERS (letter-spacing renders "C U R R E N T") and any token whose letters
    are part of the walk's text ("matte rs" of "matters") are not unseen text (spec §10 X1)."""
    seen = {t.casefold() for t in _TOKEN.findall(walk_text)}
    walk_sq = _squash(walk_text)
    page_text = _LIST_MARKER.sub("", page_text)
    out = []
    for t in (x.casefold() for x in _TOKEN.findall(page_text)):
        if t in seen or (t.isalpha() and (len(t) < RESIDUAL_MIN_LETTERS or _squash(t) in walk_sq)):
            continue
        out.append(t)
    return out


def needs_vision(residual: list[str]) -> bool:
    """R8: unseen text with a number, date, owner… (any risk tag), or a substantial amount of it."""
    return bool(tag_risk(" ".join(residual))) or len(residual) >= RESIDUAL_WORDS


def _why(page: dict, kinds: dict, residual: list[str]) -> str:
    """Why a slide goes to the vision step, or "". render_pages' thin-text flag is not used: most slides are short,
    and the shape walk already reads their text (the residual rule catches what it cannot read)."""
    drawn = page.get("flagged") and page.get("why") not in ("thin-text", "all")
    for flag, why in ((kinds.get("chart"), "chart"), (kinds.get("diagram"), "diagram"),
                      (needs_vision(residual), "text not in shapes"), (kinds.get("picture"), "picture"),
                      (drawn, page.get("why") or "visual")):
        if flag:
            return why
    return ""


def _clear_stale(out_dir: Path) -> None:
    """Remove media/ and render/ from an earlier run of this script; refuse anything else."""
    for sub in ("media", "render"):
        path = out_dir / sub
        if path.is_symlink():
            raise Refusal(f"refusing to clear {sub}/: it is a symlink")
        if path.exists() and not (out_dir / "sections.json").is_file():
            raise Refusal(f"refusing to clear {sub}/ in a folder pptx_sections.py did not create")
    for sub in ("media", "render"):
        shutil.rmtree(out_dir / sub, ignore_errors=True)


def _validate(src: Path) -> None:
    if not src.is_file():
        raise Refusal(f"no such file: {src}")
    if src.suffix.lower() != ".pptx":
        raise Refusal(f"{src.name}: deck-fact-check reads .pptx only; open a .ppt in PowerPoint and save it as .pptx")
    import zipfile
    if not zipfile.is_zipfile(src):
        raise Refusal(f"{src.name}: not a readable .pptx package")
    if render_pages.soffice_bin() is None:
        raise Refusal("LibreOffice (soffice) is needed to render the slides; install it, then run again")


def write_outputs(pptx_path, out_dir, max_words: int = 1500) -> dict:
    """Read, render, flag, then write. Returns {sections, index, hidden, flagged}. Nothing is written on a refusal."""
    src, out_dir = Path(pptx_path), Path(out_dir)
    _validate(src)
    try:
        visible = visible_slides(src)
    except Exception as e:                       # corrupt package, not a presentation
        raise Refusal(f"{src.name}: cannot read the deck ({e})") from e
    import tempfile
    out_dir.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=".render-", dir=out_dir.parent))
    pdf_tmp = None
    try:
        try:                                     # one soffice run: the PDF feeds both the page images and the locator
            pdf, pdf_tmp = render_pages.to_pdf(str(src))
            rendered = render_pages.render(pdf, staging, rel_to=str(Path(pdf).parent))
        except SystemExit as e:                  # render_pages exits with a message when soffice fails
            raise Refusal(str(e.code) if e.code not in (None, 0) else "rendering the slides failed") from e
        if len(rendered["pages"]) != len(visible):
            raise Refusal(f"{src.name}: the rendered PDF has {len(rendered['pages'])} pages for "
                          f"{len(visible)} visible slides, so pages cannot be matched to slides; "
                          "open the deck in PowerPoint, save it again and re-run")
        page_text = {n: (staging / p["image"]).with_suffix(".txt").read_text(encoding="utf-8")
                     for n, p in zip(visible, rendered["pages"])}
        import pymupdf
        from pptx import Presentation
        width = Presentation(str(src)).slide_width or 9144000
        doc = pymupdf.open(pdf)
        page_of = dict(zip(visible, range(len(visible))))
        painted = {n: RenderedText(doc[i]) for n, i in page_of.items()}
        page_text = {n: text.text for n, text in painted.items()}
        for n, page in zip(visible, rendered["pages"]):
            (staging / page["image"]).with_suffix(".txt").write_text(page_text[n], encoding="utf-8")

        def locate(n, words):
            return painted[n].locations(words, scale=width / painted[n].width)
        try:
            deck = read_deck(src, page_text, locate, lambda n, text, pos: painted[n].contains(
                text, pos, width, preceding=pos.get("preceding", "")))
        except Exception as e:
            raise Refusal(f"{src.name}: cannot read the deck ({e})") from e
        finally:
            doc.close()
        _clear_stale(out_dir)                    # only now: a failed render leaves the previous run intact
        out_dir.mkdir(parents=True, exist_ok=True)
        shutil.move(str(staging), str(out_dir / "render"))
    finally:
        shutil.rmtree(staging, ignore_errors=True)
        if pdf_tmp:
            shutil.rmtree(pdf_tmp, ignore_errors=True)
    pages = rendered["pages"]
    flagged = []
    by_slide = {sec["section_id"]: sec for sec in deck["sections"]}
    (out_dir / "media").mkdir(exist_ok=True)
    for i, n in enumerate(deck["visible"]):
        page = pages[i]
        render_dir = out_dir / "render"
        page_text = (render_dir / page["image"]).with_suffix(".txt").read_text(encoding="utf-8")
        why = _why(page, deck["kinds"][n], residual_tokens(page_text, deck["walk_text"][n]))
        if not why:
            continue
        flagged.append(n)
        rec = {"figure": n, "p_id": f"s{n}", "image": f"slide_{n:02d}.png", "why": why,
               "media": f"media/slide_{n:02d}.png", "elements": deck["elements"][n]}
        shutil.copyfile(render_dir / page["image"], out_dir / rec["media"])
        by_slide[f"s{n:02d}"]["blocks"].append(("f", rec, 0))
    off = offscreen_summary(deck["offscreen"])
    index = section_io.write_outputs(src, out_dir, deck["sections"], max_words, doc_date=deck["doc_date"], chunk=chunk_claims.auto_size,
                                     extra_stats={"hidden_slides": deck["hidden"], "flagged_slides": flagged,
                                                  "offscreen": off, "chrome": deck["chrome"]})
    return {"sections": deck["sections"], "index": index, "hidden": deck["hidden"], "flagged": flagged,
            "offscreen": off, "chrome": deck["chrome"]}


def check_dependencies() -> int:
    """Step 0: print the versions the skill needs, or the stop message (exit 2)."""
    try:
        import pptx
        import pymupdf
        versions = (getattr(pptx, "__version__", "?"), getattr(pymupdf, "VersionBind", "?"))
    except Exception:
        versions = None
    soffice = render_pages.soffice_bin()
    if versions is None or soffice is None:
        print(NEEDS, file=sys.stderr)
        return 2
    print(f"python {sys.version.split()[0]} python-pptx {versions[0]} pymupdf {versions[1]} soffice {soffice}")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("pptx", nargs="?", help="the draft .pptx")
    ap.add_argument("--out", help="directory for sections.json, batches.json, batch_<k>.json, media/ and render/")
    ap.add_argument("--check", action="store_true", help="print the versions this skill needs, then exit")
    ap.add_argument("--max-words", type=int, default=1500, help="word budget per batch (default 1500)")
    a = ap.parse_args(argv)
    if a.check:
        return check_dependencies()
    if not a.pptx or not a.out:
        ap.error("the draft .pptx and --out are required (or use --check)")
    if a.max_words < 1:
        print("deck-fact-check: --max-words must be at least 1", file=sys.stderr)
        return 2
    try:
        res = write_outputs(a.pptx, a.out, a.max_words)
    except Refusal as e:
        print(f"deck-fact-check: {e}", file=sys.stderr)
        return 2
    st = json.loads((Path(a.out) / "stats.json").read_text(encoding="utf-8"))
    print(f"{len(res['sections'])} slides, {len(res['index'])} batches -> {a.out}")
    print(f"statements {st['statements_total']}, high-risk {st['statements_risk']}, "
          f"estimate fast ~{st['estimate_minutes']['fast']} min, deep ~{st['estimate_minutes']['deep']} min")
    figs = [f for s in res["sections"] for k, f, _ in s["blocks"] if k == "f"]
    print(f"slides for the vision step: {len(figs)}" + (" (" + "; ".join(
        f"slide {f['figure']} ({f['why']}) -> {f.get('media', 'not rendered')}" for f in figs) + ")" if figs else ""))
    if res["hidden"]:
        print(f"hidden slides, not checked: {', '.join(map(str, res['hidden']))}")
    if res["offscreen"]["blocks"]:
        o = res["offscreen"]
        print(f"not on the slide, left out of the analysis: {o['blocks']} text blocks ("
              + ", ".join(f"{v} {k}" for k, v in o["by_reason"].items())
              + f"; slides {', '.join(map(str, o['slides']))})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
