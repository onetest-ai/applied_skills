"""Invariants every /kb:doc-fact-check run must satisfy, whatever the Brain or the corpus.

Library: `check_run(findings, deliverable_path, original_sha256=..., original_path=...)` -> list of
violation strings (empty = clean). CLI, for a run saved by a project:

    python fact_check_invariants.py <findings.json> <annotated.docx|.pptx> --original <draft.docx|.pptx>

The deliverable is read only through a document view (`doc_view`): headings, comments, images and the
anchor rule are the only format-specific parts, so this file ships byte-identical in doc-fact-check/ and
deck-fact-check/ (a drift test pins the copies).

A `coverage.json` beside findings.json (list of {heading, reason}) is read too, and `run.json` (written by
merge_findings.py) gives the mode and the draft's `source_sha256` recorded by sections.py before the run.
Without a recorded hash the "original document was modified" check is skipped, never faked.

Verdict agreement against a golden is deliberately not checked here: it varies between runs of
the same skill. These are the properties that must not vary.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
import zipfile
from collections import namedtuple
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # sibling imports work under python -I too

VERDICTS = {"Verified", "Incorrect", "Misleading", "Outdated", "Controversial", "No Evidence"}
DEFECTS = VERDICTS - {"Verified", "No Evidence"}
SEVERITIES = {"Blocker", "Major", "Minor"}
MAX_COMMENT_WORDS = 90                                   # round 7: the body holds a plain-language Why and a Suggested sentence
# `evidence` is what a reader sees (page, comments): a quoted source span with its document and date.
# How the claim was checked (tool calls, queries, chunk ids) belongs in `trail`.
MAX_EVIDENCE_WORDS = 60
_TRAIL_IN_EVIDENCE = re.compile(
    r"\b(search_knowledge|get_metric(_history)?|get_evidence|get_current_fact|get_question_status|get_taxonomy|"
    r"find_related_content|list_metrics|list_sources|read_document|latest_only)\b|\b\d{12,}\b")
# The step-9 findings.json keys, in SKILL.md order (a test pins this to the SKILL.md text).
# `destination` is set in step 8, so it may be absent before then (e.g. at merge_findings.py).
FINDING_KEYS = ("id", "p_id", "section", "quote", "claim", "type", "verdict", "severity", "confidence",
                "evidence", "trail", "conflicts", "fix", "why", "suggested", "source", "sources", "destination")
COMMENT_SEVERITIES = frozenset({"Blocker", "Major"})
COMMENT_POLICIES = ("major", "all")
# The destination of a finding whose comment the annotator could not place (quote or shape not found): logged.
NOT_ANCHORED = "not anchored"
LATE_KEYS = frozenset({"destination"})
# Figure findings (`I*` ids only) also carry what step 8's `annotate` needs to anchor on the drawing.
FIGURE_KEYS = ("anchor", "figure")
# A fact-check comment's first line is its category only (round 7, user 2026-10-19): no severity, no id. A comment is
# tied back to its finding by its whole text (comment_owner), which is generated from the finding.
HEADER = re.compile(r"^\[(?P<verdict>Incorrect|Misleading|Outdated|Controversial|Verified|No Evidence)\](?:\n|$)")
COMMENTED = frozenset({"Incorrect", "Misleading", "Outdated"})   # a disagreement (Controversial) stays on the findings page
NOT_A_FIX = re.compile(r"^\W*(confirm|check|verify|ask|clarify|validate|consult)\b", re.I)
MAX_WHY_WORDS = 45
_JARGON = re.compile(r"\b(Brain|chunk|chunks|verdict|finding|claim|stage V|trail)\b")
_EDIT = re.compile(r"^\W*(change|replace|remove|add|use|delete|insert|correct|update|qualify)\b", re.I)
_MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


DECK_EVIDENCE_WORDS = 90        # a deck comment carries no evidence by default (spec §10.7): the findings page only


def evidence_problem(text, deck: bool = False) -> str | None:
    """Why a finding's `evidence` is not reader-facing, or None: no tool calls, no chunk ids, at most the word limit
    (MAX_EVIDENCE_WORDS, the Brain line of a Word comment; DECK_EVIDENCE_WORDS for a deck finding)."""
    t = str(text or "")
    if _TRAIL_IN_EVIDENCE.search(t):
        return "evidence names a tool, query option or chunk id: move how it was checked to `trail`"
    limit = DECK_EVIDENCE_WORDS if deck else MAX_EVIDENCE_WORDS
    if len(t.split()) > limit:
        return f"evidence is {len(t.split())} words (limit {limit}): quote the source span, move the rest to `trail`"
    return None


_TYPO = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201a": "'", "\u2032": "'", "\u201c": '"', "\u201d": '"',
                       "\u201e": '"', "\u2013": "-", "\u2014": "-", "\u2011": "-", "\u00a0": " ", "\u2026": "..."})


def same_text(a, b) -> bool:
    """Equal once typographic variants are folded: curly and straight quotes, dashes, non-breaking spaces, ellipsis,
    runs of whitespace. A worker that re-types a quote changes these, never the words."""
    fold = lambda x: " ".join(str(x or "").translate(_TYPO).split())
    return fold(a) == fold(b)


# Claim-record fields a stage-V worker sometimes copies into its finding; the merge drops them rather than refusing.
CLAIM_ONLY_KEYS = ("s_id", "s_ids", "kind", "risk", "summary", "questions", "unresolved", "context")
# Fields a finding copies from its claim unchanged; the merge fills a missing one from the claim (a differing one is refused).
COPIED_FROM_CLAIM = ("p_id", "section", "quote", "claim", "type")
NAME_CHECK = re.compile(r"name check:", re.I)


def name_check_problem(f: dict) -> str | None:
    """A Verified finding that is not a number claim must show the name check ran (`name check:` in its trail)."""
    if f.get("verdict") == "Verified" and f.get("type") != "NUM" and not NAME_CHECK.search(str(f.get("trail") or "")):
        return "Verified without a `name check:` entry in `trail` (record the name check, or `name check: none`)"
    return None


def conflicts_problem(f: dict) -> str | None:
    """`conflicts` lists every statement found that differs from the claim for the same item; it decides
    what the verdict may be (code, not prose: spec §4.4)."""
    c = f.get("conflicts")
    if not isinstance(c, list) or any(not isinstance(x, dict) for x in c):
        return "conflicts must be a list of {span, source, date} (an empty list when no differing statement was found)"
    for n, x in enumerate(c, 1):
        empty = [k for k in ("span", "source", "date") if not str(x.get(k) or "").strip()]
        if empty:
            return (f"conflicts entry {n} has no {', '.join(empty)}: every entry is {{span, source, date}}; an undated "
                    "source takes the date in its file name or title, marked 'date inferred'")
    return _conflict_verdict_problem(f, c)


def _conflict_verdict_problem(f: dict, c: list) -> str | None:
    v = f.get("verdict")
    if v == "Verified" and c:
        return ("Verified with recorded conflicts: a differing statement for the same item blocks Verified "
                "(Controversial or Outdated, or Incorrect/Misleading when nothing supports the draft)")
    if v in ("Controversial", "Outdated") and not c:
        return f"{v} without a recorded conflict: list the differing statement in conflicts"
    if v in ("Controversial", "Outdated") and not str(f.get("evidence") or "").strip():
        return f"{v} with empty evidence: quote the supporting side in evidence, the differing one in conflicts"
    return None


def _has_comment(f: dict) -> bool:
    return "comment" in str(f.get("destination", "")).lower()


def _range_texts(docx_path: Path) -> dict[str, str]:
    xml = zipfile.ZipFile(docx_path).read("word/document.xml").decode("utf-8")
    out = {}
    # w:id may sit beside other attributes (e.g. w:displacedByCustomXml), in either order.
    start = r'<w:commentRangeStart\b[^>]*?\bw:id="(\d+)"[^>]*?/?>'
    end = r'<w:commentRangeEnd\b[^>]*?\bw:id="{}"[^>]*?/?>'
    for m in re.finditer(start, xml):
        cid = m.group(1)
        e = re.compile(end.format(cid)).search(xml, m.end())
        if not e:
            continue
        body = xml[m.end():e.start()]
        out[cid] = html.unescape("".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", body)))
    return out


COVERAGE_REASON = "no checkable statement"
# Fast mode verifies only high-risk claims; a section with none is covered in a fast run only.
RISK_COVERAGE_REASON = "no high-risk statement"
UNVERIFIED_REASON = "not verified (its chunk was refused twice)"   # a deck run's --skip-chunk: named, not hidden
SECTION_SEP = re.compile(r"\s*(?:>|›|»)\s*")


def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", str(t)).strip().casefold()


_SLIDE_SECTION = re.compile(r"^\s*(Slide \d+)\s*:\s*", re.I)     # a deck section "Slide N: title" is the path Slide N > title


def section_key(s) -> str:
    """A heading path normalised for comparison: `A > B` with each segment case- and space-folded. A deck section
    name "Slide N: title" and its heading path "Slide N > title" are the same section."""
    s = _SLIDE_SECTION.sub(r"\1 > ", str(s or ""))
    return " > ".join(_norm(seg) for seg in SECTION_SEP.split(s) if seg.strip())


def _last_heading(name: str) -> str:
    segs = [x for x in SECTION_SEP.split(name) if x.strip()]
    return segs[-1].strip() if segs else name


DECK_PID = re.compile(r"^s(\d+)(-sh\d+)?")     # a deck claim's id: s<slide>[-sh<shape>-p<k> | -r<k>]
QUOTE_WORDS = 12                                 # the statement a deck comment quotes, cut at this many words


COMMENT_FIELDS = ("quote", "why", "suggested", "brain", "source")   # the lines a comment may carry after its header
DECK_FIELDS = ("quote", "why", "suggested")   # round 7: the statement, a plain-language Why, a Suggested sentence
WORD_FIELDS = ("why", "suggested")            # a Word comment is anchored on the words, so it does not quote them


def source_label(path: str) -> tuple[str, str]:
    """A Brain source path in reader's words: (plain name, date) — the date from the name itself, else "undated".
    The model never writes these: they are read from the finding's own `source`, so a comment cannot invent them."""
    name = re.sub(r"^.*/", "", str(path or ""))
    name = re.sub(r"(\.(md|txt|pdf|pptx|docx|xlsx|xlsm|html|mp4|m4a|vtt|csv))+$", "", name, flags=re.I)
    name = name.split("__")[-1]                       # "<folder>__<sub>__<title>": the title is the last part
    date, mm = "undated", None
    m = re.search(r"(20\d\d)[-_]?([01]\d)[-_]?([0-3]\d)(?:_\d{6}(?:UTC)?)?", name)
    if m:
        mm = m
    else:
        m2 = re.search(r"^(\d\d)([01]\d)([0-3]\d)\s*-\s*", name)
        if m2:
            mm = m2
    if mm:
        y, mo, d = (mm.group(1) if len(mm.group(1)) == 4 else "20" + mm.group(1)), int(mm.group(2)), int(mm.group(3))
        if 1 <= mo <= 12 and 1 <= d <= 31:
            date = f"{d} {_MONTHS[mo - 1]} {y}"
        name = (name[:mm.start()] + " " + name[mm.end():])
    name = re.sub(r"-?\s*Meeting Recording", "", name)
    name = re.sub(r"^\s*-\s*|\s*-\s*$", "", name.strip())
    name = re.sub(r"\s+-\s+", " – ", name)
    name = re.sub(r"[-_]+", " ", name) if " – " not in name and " " not in name.strip() else re.sub(r"_+", " ", name)
    return re.sub(r"\s+", " ", name).strip(" –-") or "the source", date


def _sentence(s: str) -> str:
    s = str(s or "").strip()
    return s if not s or s[-1] in ".!?" else s + "."


def comment_text(f, fields=None):
    """The exact text of a finding's comment (Word or deck); annotate writes it, merge and annotate length-check it.
    Round 7 shape (user 2026-10-19): the category only; a deck comment then quotes the statement (it marks a point;
    a Word comment is anchored on the words); then "Why: <source> (<date>) says <why>" — the source's plain name and
    date read from the finding's `source` by source_label, the rest written by stage V — and "Suggested: "<sentence>"".
    `fields` (any of COMMENT_FIELDS) adds the Brain and Source lines back on request."""
    deck = DECK_PID.match(str(f.get("p_id") or "")) is not None
    fields = tuple(fields or (DECK_FIELDS if deck else WORD_FIELDS))
    name, date = source_label(f.get("source", ""))
    words = str(f.get("quote") or "").split()
    lines = {"why": f"Why: {name} ({date}) says {_sentence(f.get('why'))}" if str(f.get("why") or "").strip() else "",
             "suggested": f'Suggested: "{str(f.get("suggested") or "").strip()}"' if str(f.get("suggested") or "").strip() else "",
             "brain": f"Brain: {f.get('evidence', '')}", "source": f"Source: {f.get('source', '')}"}
    if deck and words and "quote" in fields:
        lines["quote"] = '"' + " ".join(words[:QUOTE_WORDS]) + ("…" if len(words) > QUOTE_WORDS else "") + '"'
    body = [lines[k] for k in COMMENT_FIELDS if k in fields and lines.get(k)]
    head = f"[{f['verdict']}]"
    if deck:
        quote = [b for b in body if b == lines.get("quote")]
        rest = [b for b in body if b != lines.get("quote")]
        return "\n".join([head, *quote]) + ("\n\n" + "\n\n".join(rest) if rest else "")
    return "\n".join([head, *body])


def comment_owner(text: str, findings) -> str | None:
    """The id of the finding whose comment this is: comments carry no id, so the text itself is matched."""
    t = str(text).strip()
    for f in findings:
        for fields in (None, COMMENT_FIELDS):
            if comment_text(f, fields).strip() == t:
                return f.get("id")
    return None


def why_problem(f) -> str | None:
    """The round-7 checks on a commented finding's plain-language Why and Suggested sentence (None when fine)."""
    why, sug = str(f.get("why") or "").strip(), str(f.get("suggested") or "").strip()
    if not why:
        return "a commented finding needs a `why`"
    if len(why.split()) > MAX_WHY_WORDS:
        return f"`why` is over {MAX_WHY_WORDS} words"
    if m := _JARGON.search(why):
        return f"`why` uses internal words ({m.group(1)}): say what the source says, in plain words"
    known = re.sub(r"[,\s]", "", f"{f.get('evidence', '')} {f.get('quote', '')}")
    for n in re.findall(r"\d[\d,.]*\d|\d", why):
        if re.sub(r",", "", n) not in known:
            return f"`why` cites {n}, which is not in the evidence or the statement"
    if not sug or _EDIT.match(sug):
        return "`suggested` must be the full sentence that replaces the statement, not an edit instruction"
    return None


def comment_words(text: str, deck: bool) -> int:
    """The words MAX_COMMENT_WORDS limits: the comment's body, not its category line or the deck's quoted statement."""
    lines = str(text).split("\n")[1:]
    return sum(len(line.split()) for line in lines if not (deck and line.startswith('"')))


def figures_accounted(findings, run) -> bool:
    """Figure claims exist (I* findings), or every one was dropped as repeating the slide's text (run.json)."""
    unverified = [c for u in (run or {}).get("unverified") or [] for c in u.get("claim_ids", [])]
    return (any(str(f.get("id", "")).startswith("I") for f in findings)
            or bool((run or {}).get("dropped_duplicate_figure")) or any(str(c).startswith("I") for c in unverified))


def is_deck_finding(f: dict) -> bool:
    return DECK_PID.match(str(f.get("p_id") or "")) is not None


def wants_comment(f: dict, policy: str = "major") -> bool:
    """Does this finding get a comment? Round 7 (user 2026-10-19): only a defect the evidence decides — Incorrect,
    Misleading or Outdated — with a plain-language `why`, and a fix that is a correction, not "confirm…/check with…".
    A disagreement (Controversial) and a finding without hard evidence stay on the findings page.
    `major` = Blocker and Major (default), `all` = every severity."""
    return (f.get("verdict") in COMMENTED and (policy == "all" or f.get("severity") in COMMENT_SEVERITIES)
            and bool(str(f.get("why") or "").strip()) and not NOT_A_FIX.match(str(f.get("fix") or "")))


def section_coverage(k, sections, claims, cov, cname) -> tuple[list[dict], list[str]]:
    """Stage-E section coverage for batch k, shared by chunk_claims.py and merge_findings.py.

    sections: the batch's `{section_id, section}` records (batches.json); claims: the batch's stage-E
    claims; cov: its `coverage_batch_<k>.json` entries (`{section_id, reason: "no checkable statement"}`).
    Returns ([{heading, section, reason}] for the sections covered by an entry, [errors]); every section
    needs a claim or a coverage entry. An entry for a section that has a claim is an error (it would hide
    that section's tagged statements); a per-statement waiver (`{s_id, reason}`, no section_id) is not a
    section entry and is validated by chunk_claims.risk_coverage_errors."""
    errors: list[str] = []
    secs = [s for s in sections if isinstance(s, dict)]
    by_id = {s.get("section_id"): s.get("section", "") for s in secs}
    claimed = {section_key(c.get("section")) for c in claims if isinstance(c, dict)}
    covered = set(claimed)
    records: list[dict] = []
    for c in cov:
        if not isinstance(c, dict):
            errors.append(f"{cname}: an entry must be an object, got {type(c).__name__}")
            continue
        if _norm(str(c.get("reason", ""))) != COVERAGE_REASON:
            errors.append(f"{cname}: reason {c.get('reason')!r} is not {COVERAGE_REASON!r}")
            continue
        if not c.get("section_id") and c.get("s_id"):
            continue
        if not c.get("section_id"):
            errors.append(f"{cname}: entry needs section_id, got {c!r}")
            continue
        name = by_id.get(c.get("section_id"))
        if name is None:
            errors.append(f"{cname}: {c.get('section_id')!r} is not a section of batch {k}")
            continue
        if section_key(name) in claimed:
            errors.append(f"{cname}: section {c.get('section_id')!r} ({name!r}) has a stage-E claim, so it cannot be "
                          f"covered as having no checkable statement (waive single sentences with "
                          f"{{s_id, reason}} instead)")
            continue
        covered.add(section_key(name))
        records.append({"heading": _last_heading(name), "section": name, "reason": COVERAGE_REASON})
    for s in secs:
        name = s.get("section", "")
        if section_key(name) not in covered:
            errors.append(f"batch {k}: section {name!r} has neither a claim nor a coverage entry "
                          f"(stage E: add a claim, or a {COVERAGE_REASON!r} entry in {cname})")
    return records, errors


def load_findings(path):
    """Load findings.json; return (list, violations). An object wrapper is unwrapped but reported."""
    data = json.loads(Path(path).read_text())
    v: list[str] = []
    if isinstance(data, dict):
        v.append("findings.json must be a JSON list")
        data = data.get("findings") or data.get("claims") or next(iter(data.values()), [])
    if not isinstance(data, list):
        v.append("findings.json must be a JSON list")
        data = []
    return data, v


def _headings(docx_path) -> list[str]:
    from docx import Document
    out = []
    for p in Document(docx_path).paragraphs:
        name = (p.style.name if p.style is not None else "") or ""
        if (name.startswith("Heading") or name == "Title") and p.text.strip():
            out.append(p.text.strip())
    return out


def _covered_headings(findings, coverage, mode=None) -> set[str]:
    reasons = {COVERAGE_REASON, UNVERIFIED_REASON} | ({RISK_COVERAGE_REASON} if mode == "fast" else set())
    cov = set()
    for f in findings:
        name = str(f.get("section") or "")
        for seg in SECTION_SEP.split(name) + SECTION_SEP.split(_SLIDE_SECTION.sub(r"\1 > ", name)):
            if seg.strip():
                cov.add(_norm(seg))
    for c in coverage or []:
        if isinstance(c, dict) and _norm(c.get("reason", "")) in reasons and c.get("heading"):
            cov.add(_norm(c["heading"]))
    return cov


def _references_images(docx_path) -> bool:
    z = zipfile.ZipFile(docx_path)
    xml = z.read("word/document.xml").decode("utf-8")
    if not re.search(r"<(?:a:blip|v:imagedata)\b[^>]*\b(?:r:embed|r:id|r:link)=", xml):
        return False
    return any(n.startswith("word/media/") for n in z.namelist())


# --- Deck (.pptx) geometry and comments: shared by pptx_sections.py (reader), pptx_annotate.py and PptxView.
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
COMMENTS_REL = R_NS + "/comments"
AUTHORS_REL = R_NS + "/commentAuthors"
EMU_PER_POS = 1587.5     # legacy comment p:pos is in 576ths of an inch; 914400 EMU per inch / 576


def slide_title(slide) -> str:
    """The text of a slide's title placeholder, or "" (the same rule the invariants' deck view uses)."""
    try:
        shape = slide.shapes.title
    except Exception:
        return ""
    return " ".join(shape.text_frame.text.split()) if shape is not None and shape.has_text_frame else ""


def is_hidden(slide) -> bool:
    return slide._element.get("show") == "0"


def _child_transform(group, outer):
    """Compose the transform that maps a group's child coordinates to slide EMU: (sx, sy, tx, ty)."""
    xfrm = group._element.grpSpPr.find(A_NS + "xfrm")
    ns = A_NS
    if xfrm is None or xfrm.find(ns + "off") is None or xfrm.find(ns + "ext") is None:
        return outer                              # the transform is optional: a group without one is the identity
    off, ext = xfrm.find(ns + "off"), xfrm.find(ns + "ext")
    ch_off, ch_ext = xfrm.find(ns + "chOff"), xfrm.find(ns + "chExt")
    ox, oy, cx, cy = (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
    hx, hy = (int(ch_off.get("x")), int(ch_off.get("y"))) if ch_off is not None else (ox, oy)
    hw, hh = (int(ch_ext.get("cx")), int(ch_ext.get("cy"))) if ch_ext is not None else (cx, cy)
    sx, sy = (cx / hw if hw else 1.0), (cy / hh if hh else 1.0)
    # child -> group space: x' = ox + (x - hx) * sx; then the outer transform
    osx, osy, otx, oty = outer
    return (osx * sx, osy * sy, osx * (ox - hx * sx) + otx, osy * (oy - hy * sy) + oty)


def _box(shape, t) -> dict:
    sx, sy, tx, ty = t
    x, y = shape.left or 0, shape.top or 0
    w, h = shape.width or 0, shape.height or 0
    return {"x": round(x * sx + tx), "y": round(y * sy + ty), "w": round(w * sx), "h": round(h * sy)}


def walk(shapes, t=(1.0, 1.0, 0.0, 0.0)):
    """Yield (shape, pos) for every non-group shape, groups recursed, positions in slide EMU."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes, _child_transform(shape, t))
        else:
            yield shape, _box(shape, t)


MIN_PICTURE_AREA = 0.01         # pictures under 1% of the slide (logos, icons) are not read by the vision step;
                                # a text callout picture is often 2-4%, so it is (spec §10)
FOOTER_PLACEHOLDERS = ("DATE", "FOOTER", "SLIDE_NUMBER")


def is_picture(shape) -> bool:
    """A raster image, wherever it sits: a p:pic (a picture, or a picture in a placeholder, whose shape_type is
    PLACEHOLDER) or a shape filled with an image (spPr/a:blipFill)."""
    el = shape._element
    if el.tag == f"{{{P_NS}}}pic":
        return True
    sppr = el.find(f"{{{P_NS}}}spPr")
    return sppr is not None and sppr.find(A_NS + "blipFill") is not None


def is_footer_placeholder(shape) -> bool:
    """The date, footer and slide-number placeholders: master furniture, never claims or context."""
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return shape.placeholder_format.type.name in FOOTER_PLACEHOLDERS
    except Exception:
        return False


def slide_kinds(slide, slide_area: int) -> dict:
    """What the shape walk cannot read on a slide: a sizeable picture, a chart, a SmartArt diagram. The reader sends
    such a slide to the vision step; the deck view expects I* findings when any visible slide has one."""
    pic = chart = diagram = False
    for shape, pos in walk(slide.shapes):
        if is_picture(shape) and pos["w"] * pos["h"] >= MIN_PICTURE_AREA * slide_area:
            pic = True
        if getattr(shape, "has_chart", False) and shape.has_chart:
            chart = True
        if shape._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds") is not None:
            diagram = True
    return {"picture": pic, "chart": chart, "diagram": diagram}


def pptx_boxes(pptx_path) -> dict[tuple[int, int], dict]:
    """(slide number, shape id) -> the shape's box in slide EMU, groups resolved, every slide (hidden included)."""
    from pptx import Presentation
    out = {}
    for n, slide in enumerate(Presentation(str(pptx_path)).slides, 1):
        for shape, pos in walk(slide.shapes):
            out[(n, shape.shape_id)] = pos
    return out


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, tuple[str, str]]:
    """Relationship id -> (type, absolute part name) for a package part."""
    d, name = part.rsplit("/", 1)
    path = f"{d}/_rels/{name}.rels"
    if path not in z.namelist():
        return {}
    import posixpath
    import xml.etree.ElementTree as ET
    out = {}
    for r in ET.fromstring(z.read(path)).iter(f"{{{PKG_REL_NS}}}Relationship"):
        out[r.get("Id")] = (r.get("Type"), posixpath.normpath(posixpath.join(d, r.get("Target"))))
    return out


def pptx_slide_parts(pptx_path) -> list[str]:
    """Slide part names in presentation order (slide n is item n-1)."""
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(pptx_path) as z:
        rels = _rels(z, "ppt/presentation.xml")
        root = ET.fromstring(z.read("ppt/presentation.xml"))
        return [rels[s.get(f"{{{R_NS}}}id")][1] for s in root.iter(f"{{{P_NS}}}sldId")]


def read_pptx_comments(pptx_path) -> list[dict]:
    """Every legacy comment: [{slide, x, y, text, author_id, idx}] in slide order (x, y in 576ths of an inch)."""
    import xml.etree.ElementTree as ET
    out = []
    parts = pptx_slide_parts(pptx_path)
    with zipfile.ZipFile(pptx_path) as z:
        for n, part in enumerate(parts, 1):
            for typ, target in _rels(z, part).values():
                if typ != COMMENTS_REL or target not in z.namelist():
                    continue
                for cm in ET.fromstring(z.read(target)).iter(f"{{{P_NS}}}cm"):
                    pos, text = cm.find(f"{{{P_NS}}}pos"), cm.find(f"{{{P_NS}}}text")
                    out.append({"slide": n, "x": int(pos.get("x")) if pos is not None else 0,
                                "y": int(pos.get("y")) if pos is not None else 0,
                                "text": (text.text or "") if text is not None else "",
                                "author_id": cm.get("authorId"), "idx": cm.get("idx")})
    return out


Comment = namedtuple("Comment", "cid text anchor")


class DocxView:
    """A .docx deliverable: Word heading styles, python-docx comments anchored on a text range."""

    def __init__(self, path):
        self.path = Path(path)

    def headings(self) -> list[str]:
        return _headings(self.path)

    def has_images(self) -> bool:
        return _references_images(self.path)

    def comments(self) -> list[Comment]:
        from docx import Document
        ranges = _range_texts(self.path)
        return [Comment(str(c.comment_id), c.text.strip(), ranges.get(str(c.comment_id), ""))
                for c in Document(self.path).comments]

    def anchor_problem(self, c: Comment, f: dict) -> str | None:
        anchored, quote = (c.anchor or "").strip(), (f.get("quote") or "").strip()
        if anchored and quote and anchored != quote:
            return f"anchor {anchored!r} is not exactly the quoted span {quote!r}"
        return None


class PptxView:
    """A .pptx deliverable: "Slide N: <title>" headings, legacy comments pinned to a point (spec §4.2)."""

    _PID = re.compile(r"^s(\d+)(?:-sh(\d+))?")

    def __init__(self, path):
        self.path = Path(path)
        self._boxes = None

    def headings(self) -> list[str]:
        """The section name of every visible slide, exactly as pptx_sections.py writes it."""
        from pptx import Presentation
        out = []
        for n, slide in enumerate(Presentation(str(self.path)).slides, 1):
            if not is_hidden(slide):
                title = slide_title(slide)
                out.append(f"Slide {n}: {title}" if title else f"Slide {n}")
        return out

    def has_images(self) -> bool:
        """True when a visible slide holds what the reader sends to the vision step (slide_kinds: a sizeable picture,
        a chart, a diagram), so I* findings must exist."""
        from pptx import Presentation
        prs = Presentation(str(self.path))
        area = (prs.slide_width or 9144000) * (prs.slide_height or 6858000)
        return any(any(slide_kinds(slide, area).values()) for slide in prs.slides if not is_hidden(slide))

    def comments(self) -> list[Comment]:
        return [Comment(f"slide {c['slide']} #{c['idx']}", c["text"].strip(), (c["slide"], c["x"], c["y"]))
                for c in read_pptx_comments(self.path)]

    def anchor_problem(self, c: Comment, f: dict) -> str | None:
        m = self._PID.match(str(f.get("p_id") or ""))
        if not m:
            return f"p_id {f.get('p_id')!r} names no slide"
        slide, x, y = c.anchor
        if slide != int(m.group(1)):
            return f"comment is on slide {slide}, its claim is on slide {m.group(1)}"
        if m.group(2) is None:                   # a whole-slide (figure) claim: any point on its slide
            return None
        if self._boxes is None:
            self._boxes = pptx_boxes(self.path)
        box = self._boxes.get((slide, int(m.group(2))))
        if box is None:
            return f"shape {m.group(2)} of slide {slide} not found for {f.get('p_id')}"
        ex, ey = x * EMU_PER_POS, y * EMU_PER_POS
        tol = EMU_PER_POS                        # one position unit of rounding
        if not (box["x"] - tol <= ex <= box["x"] + box["w"] + tol and box["y"] - tol <= ey <= box["y"] + box["h"] + tol):
            return f"comment pin is outside the box of the shape that holds {f.get('p_id')}"
        return None


VIEWS = {".docx": DocxView, ".pptx": PptxView}


def doc_view(path):
    """The document view for a deliverable, chosen by its suffix."""
    cls = VIEWS.get(Path(path).suffix.lower())
    if cls is None:
        raise ValueError(f"no document view for {Path(path).suffix or 'a file without a suffix'}: "
                         f"expected one of {', '.join(sorted(VIEWS))}")
    return cls(path)


BANNER_STALE = "findings page coverage banner is missing or stale"


def banner_violations(page_html: str, run: dict) -> list[str]:
    """The page's `<p class="coverage" data-mode=...>` must be the one findings_report renders from run.json."""
    from findings_report import coverage_line
    m = re.search(r'<p class="coverage" data-mode="([^"]*)">(.*?)</p>', page_html, re.S)
    if m and m.group(1) == str(run.get("mode", "")) and html.unescape(m.group(2)) == coverage_line(run):
        return []
    return [BANNER_STALE]


def check_run(findings, docx_path, *, original_sha256=None, original_path=None, coverage=None,
              mode=None, run=None, page=None) -> list[str]:
    docx_path = Path(docx_path)
    view = doc_view(docx_path)
    v: list[str] = []
    if run and page is not None:
        v.extend(banner_violations(page, run))
    if run and "comments" in run and run["comments"] not in COMMENT_POLICIES:
        v.append(f"run.json: comments {run['comments']!r} is not one of {', '.join(COMMENT_POLICIES)}")

    seen: set[str] = set()
    for n, f in enumerate(findings, 1):
        fid = f.get("id")
        if fid is None or fid == "":
            v.append(f"finding {n}: missing id")
            fid = f"finding {n}"
        elif fid in seen:
            v.append(f"{fid}: duplicate claim id")
        else:
            seen.add(fid)
        if not str(f.get("section") or "").strip():
            v.append(f"finding {fid}: empty section")
        if f.get("verdict") not in VERDICTS:
            v.append(f"{fid}: unknown verdict {f.get('verdict')!r}")
        if f.get("verdict") != "Verified" and f.get("severity") not in SEVERITIES:
            v.append(f"{fid}: unknown severity {f.get('severity')!r}")
        if "conflicts" in f:                          # a v2 finding (v1 runs had no `conflicts`): the merge's checks hold
            for why in (evidence_problem(f.get("evidence"), is_deck_finding(f)), conflicts_problem(f)):
                if why:
                    v.append(f"{fid}: {why}")

        verdict, commented = f.get("verdict"), _has_comment(f)
        policy = (run or {}).get("comments", "all")   # a v1 run (no key) keeps v1's rule: every defect commented
        dest = str(f.get("destination", ""))
        # a Deep reviewer's edits to the comment set (annotate.py --approved) are recorded, not violations
        if wants_comment(f, policy) and not commented and dest != "declined by reviewer" \
                and not dest.startswith(NOT_ANCHORED):
            v.append(f"{fid}: {verdict} {f.get('severity')} defect must have a comment")
        if (verdict in DEFECTS and commented and not wants_comment(f, policy)
                and "added by reviewer" not in dest):
            v.append(f"{fid}: {f.get('severity')} defect must not carry a comment (comments: {policy})")
        if verdict == "No Evidence":
            if f.get("severity") != "Minor":
                v.append(f"{fid}: No Evidence must be Minor")
            if commented:
                v.append(f"{fid}: No Evidence must not carry a comment")
        if verdict == "Verified" and commented:
            v.append(f"{fid}: Verified claim must not carry a comment")

    if original_path is not None:
        if docx_path.resolve() == Path(original_path).resolve():
            v.append("output would overwrite the original document")
        original = doc_view(original_path)
        if original.has_images() and not figures_accounted(findings, run):
            v.append("figures present but no I* findings")
        covered = _covered_headings(findings, coverage, mode)
        for h in original.headings():
            if _norm(h) not in covered:
                v.append(f"heading not covered: {h}")
        if original_sha256 and hashlib.sha256(Path(original_path).read_bytes()).hexdigest() != original_sha256:
            v.append("original document was modified")

    by_id = {}
    # comments the original already held (a reviewer's, not a fact-check header) are kept and not this run's to judge
    before = ({x.text for x in doc_view(original_path).comments() if not HEADER.match(x.text)}
              if original_path is not None and Path(original_path).resolve() != docx_path.resolve() else set())
    for c in view.comments():
        if c.text in before:
            continue
        m = HEADER.match(c.text)
        if not m:
            v.append(f"comment {c.cid}: first line is not [Category]")
            continue
        fid = comment_owner(c.text, findings) or f"comment {c.cid}"
        by_id[fid] = (c, m)
        if comment_words(c.text, isinstance(view, PptxView)) > MAX_COMMENT_WORDS:
            v.append(f"{fid}: comment is over {MAX_COMMENT_WORDS} words")

    wanted = {f["id"]: f for f in findings if _has_comment(f) and f.get("id")}
    for fid, f in wanted.items():
        if fid not in by_id:
            v.append(f"{fid}: approved comment is missing from the document")
            continue
        c, m = by_id[fid]
        if m["verdict"].strip() != f.get("verdict"):
            v.append(f"{fid}: comment category does not match the finding")
        if why := view.anchor_problem(c, f):
            v.append(f"{fid}: {why}")
    for fid in by_id:
        if fid not in wanted:
            v.append(f"{fid}: comment in the document was never approved")
    return v


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("findings")
    ap.add_argument("docx", help="the fact-checked copy (.docx or .pptx)")
    ap.add_argument("--original")
    ap.add_argument("--page", help="the findings page (default: '<name> — findings.html' beside the annotated docx)")
    a = ap.parse_args(argv)
    data, shape = load_findings(a.findings)
    beside = Path(a.findings).with_name
    coverage = json.loads(beside("coverage.json").read_text()) if beside("coverage.json").exists() else None
    run = json.loads(beside("run.json").read_text()) if beside("run.json").exists() else {}
    page = Path(a.page) if a.page else None
    if page is None and Path(a.docx).stem.endswith(" — fact-checked"):
        page = Path(a.docx).with_name(Path(a.docx).stem[:-len(" — fact-checked")] + " — findings.html")
    page_html = page.read_text(encoding="utf-8") if page is not None and page.exists() else None
    page_missing = bool(run) and page_html is None
    vs = shape + check_run(data, a.docx, original_sha256=run.get("source_sha256"), original_path=a.original,
                           coverage=coverage, mode=run.get("mode"), run=run, page=page_html)
    if run and not (beside("baseline.md").exists() and beside("baseline.sources.json").exists()):
        vs.append("baseline.md and baseline.sources.json missing beside findings.json (run baseline.py)")
    if page_missing:
        vs.append("findings page not found for banner check")
    print("\n".join(vs) or "clean")
    return 1 if vs else 0


if __name__ == "__main__":
    sys.exit(main())
