"""Invariants every /kb:doc-fact-check run must satisfy, whatever the Brain or the corpus.

Library: `check_run(findings, docx_path, original_sha256=..., original_path=...)` -> list of
violation strings (empty = clean). CLI, for a run saved by a project:

    python fact_check_invariants.py <findings.json> <annotated.docx> --original <draft.docx>

A `coverage.json` beside findings.json (list of {heading, reason}) is read too, and `run.json` (written by
merge_findings.py) gives the mode and the draft's `source_sha256` recorded by sections.py before the run.
Without a recorded hash the "original document was modified" check is skipped, never faked.

Verdict agreement against a golden is deliberately not checked here: it varies between runs of
the same skill. These are the properties that must not vary.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))  # sibling imports work under python -I too

VERDICTS = {"Verified", "Incorrect", "Misleading", "Outdated", "Controversial", "No Evidence"}
DEFECTS = VERDICTS - {"Verified", "No Evidence"}
SEVERITIES = {"Blocker", "Major", "Minor"}
MAX_COMMENT_WORDS = 60
# `evidence` is what a reader sees (page, comments): a quoted source span with its document and date.
# How the claim was checked (tool calls, queries, chunk ids) belongs in `trail`.
MAX_EVIDENCE_WORDS = 60
_TRAIL_IN_EVIDENCE = re.compile(
    r"\b(search_knowledge|get_metric(_history)?|get_evidence|get_current_fact|get_question_status|get_taxonomy|"
    r"find_related_content|list_metrics|list_sources|read_document|latest_only)\b|\b\d{12,}\b")
# The step-9 findings.json keys, in SKILL.md order (a test pins this to the SKILL.md text).
# `destination` is set in step 8, so it may be absent before then (e.g. at merge_findings.py).
FINDING_KEYS = ("id", "p_id", "section", "quote", "type", "verdict", "severity", "confidence",
                "evidence", "trail", "fix", "source", "sources", "destination")
LATE_KEYS = frozenset({"destination"})
# Figure findings (`I*` ids only) also carry what step 8's `annotate` needs to anchor on the drawing.
FIGURE_KEYS = ("anchor", "figure")
HEADER = re.compile(r"^\[(?P<verdict>[^·\]]+?) · (?P<severity>[^·\]]+?) · (?P<id>[^\]\s]+)\]")


def evidence_problem(text) -> str | None:
    """Why a finding's `evidence` is not reader-facing, or None: no tool calls, no chunk ids, at most the word limit."""
    t = str(text or "")
    if _TRAIL_IN_EVIDENCE.search(t):
        return "evidence names a tool, query option or chunk id: move how it was checked to `trail`"
    if len(t.split()) > MAX_EVIDENCE_WORDS:
        return f"evidence is {len(t.split())} words (limit {MAX_EVIDENCE_WORDS}): quote the source span, move the rest to `trail`"
    return None


# Claim-record fields a stage-V worker sometimes copies into its finding; the merge drops them rather than refusing.
CLAIM_ONLY_KEYS = ("s_id", "s_ids", "kind", "risk")
NAME_CHECK = re.compile(r"name check:", re.I)


def name_check_problem(f: dict) -> str | None:
    """A Verified finding that is not a number claim must show the name check ran (`name check:` in its trail)."""
    if f.get("verdict") == "Verified" and f.get("type") != "NUM" and not NAME_CHECK.search(str(f.get("trail") or "")):
        return "Verified without a `name check:` entry in `trail` (record the name check, or `name check: none`)"
    return None


def _has_comment(f: dict) -> bool:
    return "comment" in str(f.get("destination", "")).lower()


def _range_texts(docx_path: Path) -> dict[str, str]:
    xml = zipfile.ZipFile(docx_path).read("word/document.xml").decode("utf-8")
    out = {}
    # w:id may sit beside other attributes (e.g. w:displacedByCustomXml), in either order.
    start = r'<w:commentRangeStart\b[^>]*?\bw:id="(\d+)"[^>]*?/?>'
    end = r'<w:commentRangeEnd\b[^>]*?\bw:id="{}"[^>]*?/?>'
    for m in re.finditer(start, xml):
        cid = m.group(1)
        e = re.compile(end.format(cid)).search(xml, m.end())
        if not e:
            continue
        body = xml[m.end():e.start()]
        out[cid] = html.unescape("".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", body)))
    return out


COVERAGE_REASON = "no checkable statement"
# Fast mode verifies only high-risk claims; a section with none is covered in a fast run only.
RISK_COVERAGE_REASON = "no high-risk statement"
SECTION_SEP = re.compile(r"\s*(?:>|›|»)\s*")


def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", str(t)).strip().casefold()


def section_key(s) -> str:
    """A heading path normalised for comparison: `A > B` with each segment case- and space-folded."""
    return " > ".join(_norm(seg) for seg in SECTION_SEP.split(str(s or "")) if seg.strip())


def _last_heading(name: str) -> str:
    segs = [x for x in SECTION_SEP.split(name) if x.strip()]
    return segs[-1].strip() if segs else name


def section_coverage(k, sections, claims, cov, cname) -> tuple[list[dict], list[str]]:
    """Stage-E section coverage for batch k, shared by chunk_claims.py and merge_findings.py.

    sections: the batch's `{section_id, section}` records (batches.json); claims: the batch's stage-E
    claims; cov: its `coverage_batch_<k>.json` entries (`{section_id, reason: "no checkable statement"}`).
    Returns ([{heading, section, reason}] for the sections covered by an entry, [errors]); every section
    needs a claim or a coverage entry. An entry for a section that has a claim is an error (it would hide
    that section's tagged statements); a per-statement waiver (`{s_id, reason}`, no section_id) is not a
    section entry and is validated by chunk_claims.risk_coverage_errors."""
    errors: list[str] = []
    secs = [s for s in sections if isinstance(s, dict)]
    by_id = {s.get("section_id"): s.get("section", "") for s in secs}
    claimed = {section_key(c.get("section")) for c in claims if isinstance(c, dict)}
    covered = set(claimed)
    records: list[dict] = []
    for c in cov:
        if not isinstance(c, dict):
            errors.append(f"{cname}: an entry must be an object, got {type(c).__name__}")
            continue
        if _norm(str(c.get("reason", ""))) != COVERAGE_REASON:
            errors.append(f"{cname}: reason {c.get('reason')!r} is not {COVERAGE_REASON!r}")
            continue
        if not c.get("section_id") and c.get("s_id"):
            continue
        if not c.get("section_id"):
            errors.append(f"{cname}: entry needs section_id, got {c!r}")
            continue
        name = by_id.get(c.get("section_id"))
        if name is None:
            errors.append(f"{cname}: {c.get('section_id')!r} is not a section of batch {k}")
            continue
        if section_key(name) in claimed:
            errors.append(f"{cname}: section {c.get('section_id')!r} ({name!r}) has a stage-E claim, so it cannot be "
                          f"covered as having no checkable statement (waive single sentences with "
                          f"{{s_id, reason}} instead)")
            continue
        covered.add(section_key(name))
        records.append({"heading": _last_heading(name), "section": name, "reason": COVERAGE_REASON})
    for s in secs:
        name = s.get("section", "")
        if section_key(name) not in covered:
            errors.append(f"batch {k}: section {name!r} has neither a claim nor a coverage entry "
                          f"(stage E: add a claim, or a {COVERAGE_REASON!r} entry in {cname})")
    return records, errors


def load_findings(path):
    """Load findings.json; return (list, violations). An object wrapper is unwrapped but reported."""
    data = json.loads(Path(path).read_text())
    v: list[str] = []
    if isinstance(data, dict):
        v.append("findings.json must be a JSON list")
        data = data.get("findings") or data.get("claims") or next(iter(data.values()), [])
    if not isinstance(data, list):
        v.append("findings.json must be a JSON list")
        data = []
    return data, v


def _headings(docx_path) -> list[str]:
    from docx import Document
    out = []
    for p in Document(docx_path).paragraphs:
        name = (p.style.name if p.style is not None else "") or ""
        if (name.startswith("Heading") or name == "Title") and p.text.strip():
            out.append(p.text.strip())
    return out


def _covered_headings(findings, coverage, mode=None) -> set[str]:
    reasons = {COVERAGE_REASON} | ({RISK_COVERAGE_REASON} if mode == "fast" else set())
    cov = set()
    for f in findings:
        for seg in SECTION_SEP.split(str(f.get("section") or "")):
            if seg.strip():
                cov.add(_norm(seg))
    for c in coverage or []:
        if isinstance(c, dict) and _norm(c.get("reason", "")) in reasons and c.get("heading"):
            cov.add(_norm(c["heading"]))
    return cov


def _references_images(docx_path) -> bool:
    z = zipfile.ZipFile(docx_path)
    xml = z.read("word/document.xml").decode("utf-8")
    if not re.search(r"<(?:a:blip|v:imagedata)\b[^>]*\b(?:r:embed|r:id|r:link)=", xml):
        return False
    return any(n.startswith("word/media/") for n in z.namelist())


BANNER_STALE = "findings page coverage banner is missing or stale"


def banner_violations(page_html: str, run: dict) -> list[str]:
    """The page's `<p class="coverage" data-mode=...>` must be the one findings_report renders from run.json."""
    from findings_report import coverage_line
    m = re.search(r'<p class="coverage" data-mode="([^"]*)">(.*?)</p>', page_html, re.S)
    if m and m.group(1) == str(run.get("mode", "")) and html.unescape(m.group(2)) == coverage_line(run):
        return []
    return [BANNER_STALE]


def check_run(findings, docx_path, *, original_sha256=None, original_path=None, coverage=None,
              mode=None, run=None, page=None) -> list[str]:
    from docx import Document

    docx_path = Path(docx_path)
    v: list[str] = []
    if run and page is not None:
        v.extend(banner_violations(page, run))

    seen: set[str] = set()
    for n, f in enumerate(findings, 1):
        fid = f.get("id")
        if fid is None or fid == "":
            v.append(f"finding {n}: missing id")
            fid = f"finding {n}"
        elif fid in seen:
            v.append(f"{fid}: duplicate claim id")
        else:
            seen.add(fid)
        if not str(f.get("section") or "").strip():
            v.append(f"finding {fid}: empty section")
        if f.get("verdict") not in VERDICTS:
            v.append(f"{fid}: unknown verdict {f.get('verdict')!r}")
        if f.get("verdict") != "Verified" and f.get("severity") not in SEVERITIES:
            v.append(f"{fid}: unknown severity {f.get('severity')!r}")

        verdict, commented = f.get("verdict"), _has_comment(f)
        if verdict in DEFECTS and not commented:
            v.append(f"{fid}: {verdict} defect must have a Word comment")
        if verdict == "No Evidence":
            if f.get("severity") != "Minor":
                v.append(f"{fid}: No Evidence must be Minor")
            if commented:
                v.append(f"{fid}: No Evidence must not carry a comment")
        if verdict == "Verified" and commented:
            v.append(f"{fid}: Verified claim must not carry a comment")

    if original_path is not None:
        if docx_path.resolve() == Path(original_path).resolve():
            v.append("output would overwrite the original document")
        if _references_images(original_path) and not any(
                str(f.get("id", "")).startswith("I") for f in findings):
            v.append("figures present but no I* findings")
        covered = _covered_headings(findings, coverage, mode)
        for h in _headings(original_path):
            if _norm(h) not in covered:
                v.append(f"heading not covered: {h}")
        if original_sha256 and hashlib.sha256(Path(original_path).read_bytes()).hexdigest() != original_sha256:
            v.append("original document was modified")

    doc = Document(docx_path)
    by_id = {}
    ranges = _range_texts(docx_path)
    for c in doc.comments:
        m = HEADER.match(c.text.strip())
        if not m:
            v.append(f"comment {c.comment_id}: header is not [Verdict · Severity · id]")
            continue
        by_id[m["id"]] = (c, m)
        if len(c.text.split()) > MAX_COMMENT_WORDS:
            v.append(f"{m['id']}: comment is over {MAX_COMMENT_WORDS} words")

    wanted = {f["id"]: f for f in findings if _has_comment(f) and f.get("id")}
    for fid, f in wanted.items():
        if fid not in by_id:
            v.append(f"{fid}: approved comment is missing from the document")
            continue
        c, m = by_id[fid]
        if m["verdict"].strip() != f.get("verdict") or m["severity"].strip() != f.get("severity"):
            v.append(f"{fid}: comment header does not match the finding")
        anchored = ranges.get(str(c.comment_id), "")
        quote = (f.get("quote") or "").strip()
        if anchored and quote and anchored.strip() != quote:
            v.append(f"{fid}: anchor {anchored.strip()!r} is not exactly the quoted span {quote!r}")
    for fid in by_id:
        if fid not in wanted:
            v.append(f"{fid}: comment in the document was never approved")
    return v


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("findings")
    ap.add_argument("docx")
    ap.add_argument("--original")
    ap.add_argument("--page", help="the findings page (default: '<name> — findings.html' beside the annotated docx)")
    a = ap.parse_args(argv)
    data, shape = load_findings(a.findings)
    beside = Path(a.findings).with_name
    coverage = json.loads(beside("coverage.json").read_text()) if beside("coverage.json").exists() else None
    run = json.loads(beside("run.json").read_text()) if beside("run.json").exists() else {}
    page = Path(a.page) if a.page else None
    if page is None and Path(a.docx).stem.endswith(" — fact-checked"):
        page = Path(a.docx).with_name(Path(a.docx).stem[:-len(" — fact-checked")] + " — findings.html")
    page_html = page.read_text(encoding="utf-8") if page is not None and page.exists() else None
    page_missing = bool(run) and page_html is None
    vs = shape + check_run(data, a.docx, original_sha256=run.get("source_sha256"), original_path=a.original,
                           coverage=coverage, mode=run.get("mode"), run=run, page=page_html)
    if page_missing:
        vs.append("findings page not found for banner check")
    print("\n".join(vs) or "clean")
    return 1 if vs else 0


if __name__ == "__main__":
    sys.exit(main())
