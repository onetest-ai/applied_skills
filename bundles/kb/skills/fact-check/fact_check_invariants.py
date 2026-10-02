"""Invariants every /kb:fact-check run must satisfy, whatever the Brain or the corpus.

Library: `check_run(findings, docx_path, original_sha256=..., original_path=...)` -> list of
violation strings (empty = clean). CLI, for a run saved by a project:

    python fact_check_invariants.py <findings.json> <annotated.docx> --original <draft.docx>

A `coverage.json` beside findings.json (list of {heading, reason: "no checkable statement"}) is read too.

Verdict agreement against a golden is deliberately not checked here: it varies between runs of
the same skill. These are the properties that must not vary.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
import zipfile
from pathlib import Path

VERDICTS = {"Verified", "Incorrect", "Misleading", "Outdated", "Controversial", "No Evidence"}
DEFECTS = VERDICTS - {"Verified", "No Evidence"}
SEVERITIES = {"Blocker", "Major", "Minor"}
MAX_COMMENT_WORDS = 60
# The step-9 findings.json keys, in SKILL.md order (a test pins this to the SKILL.md text).
# `destination` is set in step 8, so it may be absent before then (e.g. at merge_findings.py).
FINDING_KEYS = ("id", "p_id", "section", "quote", "type", "verdict", "severity", "confidence",
                "evidence", "fix", "source", "sources", "destination")
LATE_KEYS = frozenset({"destination"})
HEADER = re.compile(r"^\[(?P<verdict>[^·\]]+?) · (?P<severity>[^·\]]+?) · (?P<id>[^\]\s]+)\]")


def _has_comment(f: dict) -> bool:
    return "comment" in str(f.get("destination", "")).lower()


def _range_texts(docx_path: Path) -> dict[str, str]:
    xml = zipfile.ZipFile(docx_path).read("word/document.xml").decode("utf-8")
    out = {}
    # w:id may sit beside other attributes (e.g. w:displacedByCustomXml), in either order.
    start = r'<w:commentRangeStart\b[^>]*?\bw:id="(\d+)"[^>]*?/?>'
    end = r'<w:commentRangeEnd\b[^>]*?\bw:id="{}"[^>]*?/?>'
    for m in re.finditer(start, xml):
        cid = m.group(1)
        e = re.compile(end.format(cid)).search(xml, m.end())
        if not e:
            continue
        body = xml[m.end():e.start()]
        out[cid] = html.unescape("".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", body)))
    return out


COVERAGE_REASON = "no checkable statement"
SECTION_SEP = re.compile(r"\s*(?:>|›|»)\s*")


def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", str(t)).strip().casefold()


def load_findings(path):
    """Load findings.json; return (list, violations). An object wrapper is unwrapped but reported."""
    data = json.load(open(path))
    v: list[str] = []
    if isinstance(data, dict):
        v.append("findings.json must be a JSON list")
        data = data.get("findings") or data.get("claims") or next(iter(data.values()), [])
    if not isinstance(data, list):
        v.append("findings.json must be a JSON list")
        data = []
    return data, v


def _headings(docx_path) -> list[str]:
    from docx import Document
    out = []
    for p in Document(docx_path).paragraphs:
        name = (p.style.name if p.style is not None else "") or ""
        if (name.startswith("Heading") or name == "Title") and p.text.strip():
            out.append(p.text.strip())
    return out


def _covered_headings(findings, coverage) -> set[str]:
    cov = set()
    for f in findings:
        for seg in SECTION_SEP.split(str(f.get("section") or "")):
            if seg.strip():
                cov.add(_norm(seg))
    for c in coverage or []:
        if isinstance(c, dict) and _norm(c.get("reason", "")) == COVERAGE_REASON and c.get("heading"):
            cov.add(_norm(c["heading"]))
    return cov


def _references_images(docx_path) -> bool:
    z = zipfile.ZipFile(docx_path)
    xml = z.read("word/document.xml").decode("utf-8")
    if not re.search(r"<(?:a:blip|v:imagedata)\b[^>]*\b(?:r:embed|r:id|r:link)=", xml):
        return False
    return any(n.startswith("word/media/") for n in z.namelist())


def check_run(findings, docx_path, *, original_sha256=None, original_path=None, coverage=None) -> list[str]:
    from docx import Document

    docx_path = Path(docx_path)
    v: list[str] = []

    seen: set[str] = set()
    for n, f in enumerate(findings, 1):
        fid = f.get("id")
        if fid is None or fid == "":
            v.append(f"finding {n}: missing id")
            fid = f"finding {n}"
        elif fid in seen:
            v.append(f"{fid}: duplicate claim id")
        else:
            seen.add(fid)
        if not str(f.get("section") or "").strip():
            v.append(f"finding {fid}: empty section")
        if f.get("verdict") not in VERDICTS:
            v.append(f"{fid}: unknown verdict {f.get('verdict')!r}")
        if f.get("verdict") != "Verified" and f.get("severity") not in SEVERITIES:
            v.append(f"{fid}: unknown severity {f.get('severity')!r}")

        verdict, commented = f.get("verdict"), _has_comment(f)
        if verdict in DEFECTS and not commented:
            v.append(f"{fid}: {verdict} defect must have a Word comment")
        if verdict == "No Evidence":
            if f.get("severity") != "Minor":
                v.append(f"{fid}: No Evidence must be Minor")
            if commented:
                v.append(f"{fid}: No Evidence must not carry a comment")
        if verdict == "Verified" and commented:
            v.append(f"{fid}: Verified claim must not carry a comment")

    if original_path is not None:
        if docx_path.resolve() == Path(original_path).resolve():
            v.append("output would overwrite the original document")
        if _references_images(original_path) and not any(
                str(f.get("id", "")).startswith("I") for f in findings):
            v.append("figures present but no I* findings")
        covered = _covered_headings(findings, coverage)
        for h in _headings(original_path):
            if _norm(h) not in covered:
                v.append(f"heading not covered: {h}")
        if original_sha256 and hashlib.sha256(Path(original_path).read_bytes()).hexdigest() != original_sha256:
            v.append("original document was modified")

    doc = Document(docx_path)
    by_id = {}
    ranges = _range_texts(docx_path)
    for c in doc.comments:
        m = HEADER.match(c.text.strip())
        if not m:
            v.append(f"comment {c.comment_id}: header is not [Verdict · Severity · id]")
            continue
        by_id[m["id"]] = (c, m)
        if len(c.text.split()) > MAX_COMMENT_WORDS:
            v.append(f"{m['id']}: comment is over {MAX_COMMENT_WORDS} words")

    wanted = {f["id"]: f for f in findings if _has_comment(f) and f.get("id")}
    for fid, f in wanted.items():
        if fid not in by_id:
            v.append(f"{fid}: approved comment is missing from the document")
            continue
        c, m = by_id[fid]
        if m["verdict"].strip() != f.get("verdict") or m["severity"].strip() != f.get("severity"):
            v.append(f"{fid}: comment header does not match the finding")
        anchored = ranges.get(str(c.comment_id), "")
        quote = (f.get("quote") or "").strip()
        if anchored and quote and anchored.strip() != quote:
            v.append(f"{fid}: anchor {anchored.strip()!r} is not exactly the quoted span {quote!r}")
    for fid in by_id:
        if fid not in wanted:
            v.append(f"{fid}: comment in the document was never approved")
    return v


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("findings")
    ap.add_argument("docx")
    ap.add_argument("--original")
    a = ap.parse_args(argv)
    data, shape = load_findings(a.findings)
    cov_path = Path(a.findings).with_name("coverage.json")
    coverage = json.load(open(cov_path)) if cov_path.exists() else None
    sha = hashlib.sha256(Path(a.original).read_bytes()).hexdigest() if a.original else None
    vs = shape + check_run(data, a.docx, original_sha256=sha, original_path=a.original, coverage=coverage)
    print("\n".join(vs) or "clean")
    return 1 if vs else 0


if __name__ == "__main__":
    sys.exit(main())
