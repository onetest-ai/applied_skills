"""Word comments need a quote that is verbatim in ONE paragraph and a p_id equal to that paragraph's location id.
chunk_claims.py --repair-anchors repairs both for the claims it can (a quote cut to its verbatim run; a table row's p_id set to its row id)."""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from test_plugin_structure import KB_ROOT

DOC = KB_ROOT / "skills" / "doc-fact-check"
sys.path.insert(0, str(DOC))

try:
    import docx
    import annotate as A
    import chunk_claims as C
    HAVE = hasattr(docx.Document(), "comments")
except Exception:  # python-docx missing or older than 1.2
    HAVE = False

SENTENCE = "The nightly import runs at 02:00 and retries once; after two failures a priority ticket is raised."
CELL = "Backup jobs run at 04:30 and retry twice"        # distinct from the body sentence, so a comment cannot land there by accident


def finding(fid, quote, p_id):
    return {"id": fid, "p_id": p_id, "section": "Scope", "quote": quote, "type": "NUM", "verdict": "Incorrect",
            "severity": "Major", "confidence": "High", "evidence": "e", "trail": "t", "conflicts": [], "fix": "f",
            "why": "the source states another value", "suggested": "The statement with the source's value.", "source": "s",
            "sources": [{"name": "s", "link": "", "folder": ""}], "destination": ""}


@unittest.skipUnless(HAVE, "python-docx >= 1.2 is needed to write Word comments")
class RepairedAnchorsGetTheirComment(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())
        doc = docx.Document()
        doc.add_paragraph(SENTENCE)
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Backup schedule", CELL
        self.src = self.d / "in.docx"
        doc.save(self.src)
        self.ids = {p.text: i for p, i in A.word_paragraph_ids(docx.Document(self.src)).items()}
        self.row_id, self.body_id = self.ids[CELL], self.ids[SENTENCE]

    def run_annotate(self, quote, p_id):
        out = self.d / "out.docx"
        written, skipped = A.annotate(self.src, out, [finding("C01", quote, p_id)])
        return written, skipped, list(A._range_texts(out).values())

    def test_a_composite_quote_is_skipped_today(self):
        written, skipped, _ = self.run_annotate('He says: "the nightly import runs at 02:00 and retries once"', self.body_id)
        self.assertEqual((written, skipped), ([], ["C01"]))

    def test_the_snapped_paragraph_quote_gets_its_comment_on_exactly_that_span(self):
        snapped = C.snap_quote('He says: "the nightly import runs at 02:00 and retries once"', SENTENCE)
        written, skipped, ranges = self.run_annotate(snapped, self.body_id)
        self.assertEqual((written, skipped, ranges), (["C01"], [], [snapped]))

    def test_the_snapped_table_cell_quote_gets_its_comment_inside_that_cell(self):
        snapped = C.snap_quote("Backup schedule: backup jobs run at 04:30 and retry twice", CELL)
        written, skipped, ranges = self.run_annotate(snapped, self.row_id)
        self.assertEqual((written, skipped, ranges), (["C01"], [], [CELL]))

    def test_a_row_claim_with_a_body_paragraph_id_is_skipped_today(self):
        written, skipped, _ = self.run_annotate(CELL, self.body_id)
        self.assertEqual((written, skipped), ([], ["C01"]))

    def test_the_same_claim_with_its_row_id_gets_its_comment(self):
        written, skipped, ranges = self.run_annotate(CELL, self.row_id)
        self.assertEqual((written, skipped, ranges), (["C01"], [], [CELL]))
