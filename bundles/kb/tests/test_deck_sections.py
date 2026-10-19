"""deck-fact-check reader (spec §4.1, §2.3–2.4, R8, R9): shape walk, context, render integration, refusals."""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from test_plugin_structure import KB_ROOT

DECK = KB_ROOT / "skills" / "deck-fact-check"
sys.path.insert(0, str(DECK))

try:
    import pptx  # noqa: F401
    import pymupdf  # noqa: F401
    HAVE_LIBS = True
except ImportError:
    HAVE_LIBS = False

if HAVE_LIBS:
    import deck_fixtures as F
    import pptx_sections as P
    import render_pages
    HAVE_SOFFICE = render_pages.soffice_bin() is not None
else:
    HAVE_SOFFICE = False


def blocks(section, kind=None):
    return [(k, r) for k, r, _ in section["blocks"] if kind is None or k == kind]


@unittest.skipUnless(HAVE_LIBS, "needs python-pptx and PyMuPDF")
class ShapeWalk(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())

    def test_u1_title_and_bullets(self):
        sec = P.read_deck(F.basic_deck(self.d))["sections"][0]
        self.assertEqual(sec["section"], "Slide 1: Delivery plan")
        self.assertEqual(sec["heading_path"], ["Slide 1", "Delivery plan"])
        body = [r for k, r in blocks(sec, "p") if r["p_id"].startswith("s1-sh3-")]
        self.assertEqual([r["p_id"] for r in body], ["s1-sh3-p1", "s1-sh3-p2", "s1-sh3-p3"])
        self.assertEqual(body[2]["sentences"][0]["s_id"], "s1-sh3-p3s1")
        self.assertIn("num", body[2]["sentences"][0]["risk"])
        self.assertEqual(sec["heading_p_id"], "s1-sh2-p1")

    def test_u2_ids_are_stable_and_survive_a_slide_inserted_before(self):
        path = F.basic_deck(self.d)
        ids = lambda p: [r["p_id"] for s in P.read_deck(p)["sections"] for _, r in blocks(s, "p")]
        first = ids(path)
        self.assertEqual(first, ids(path))
        from pptx import Presentation
        prs = Presentation(str(path))
        new = prs.slides.add_slide(prs.slide_layouts[5])
        new.shapes.title.text = "Agenda"
        lst = prs.slides._sldIdLst
        lst.insert(0, lst[-1])                           # move the new slide to position 1
        prs.save(str(self.d / "moved.pptx"))
        moved = ids(self.d / "moved.pptx")
        self.assertEqual([i.split("-", 1)[1] for i in moved if i.startswith("s2-")],
                         [i.split("-", 1)[1] for i in first])
        self.assertEqual(len(moved), len(set(moved)))

    def test_u3_table_rows_cells_and_risk(self):
        sec = P.read_deck(F.table_deck(self.d))["sections"][0]
        rows = [r for _, r in blocks(sec, "r")]
        self.assertEqual([r["cells"] for r in rows], [["Metric", "Value"], ["Uptime", "99.9%"], ["Owner", "Integration team"]])
        self.assertEqual([r["r_id"] for r in rows], ["s1-sh3-r1", "s1-sh3-r2", "s1-sh3-r3"])
        self.assertEqual([r["s_id"] for r in rows], [r["r_id"] for r in rows])
        self.assertIn("label", rows[0]["risk"])
        self.assertIn("num", rows[1]["risk"])
        self.assertNotIn("num", rows[0]["risk"])

    def test_u4_nested_groups_are_walked_once_with_slide_positions(self):
        sec = P.read_deck(F.group_deck(self.d))["sections"][0]
        ps = [r for _, r in blocks(sec, "p")]
        self.assertEqual([r["text"] for r in ps], ["Northwind sends orders to the warehouse.",
                                                   "Southfield owns the billing feed."])
        self.assertEqual(len({r["p_id"] for r in ps}), 2)
        shift = (1 + F.GROUP_SHIFT_IN) * F.EMU_PER_INCH
        self.assertEqual([r["pos"]["x"] for r in ps], [shift, shift])        # both moved with the outer group
        self.assertEqual(ps[1]["pos"]["y"], 3 * F.EMU_PER_INCH)

    def test_u8_untitled_slide_and_empty_placeholders(self):
        sec = P.read_deck(F.untitled_deck(self.d))["sections"][0]
        self.assertEqual((sec["section"], sec["heading_path"], sec["heading_p_id"]), ("Slide 1", ["Slide 1"], None))
        self.assertEqual([r["text"] for _, r in blocks(sec)], ["Eastgate retired the mainframe in 2025."])

    def test_u12_table_cell_context(self):
        rows = [r for _, r in blocks(P.read_deck(F.table_deck(self.d))["sections"][0], "r")]
        self.assertNotIn("ctx", rows[0])
        self.assertEqual(rows[1]["ctx"], {"header": ["Metric", "Value"], "row_label": "Uptime"})

    def test_u13_sub_bullet_parent_is_in_the_same_shape(self):
        sec = P.read_deck(F.basic_deck(self.d))["sections"][0]
        by_text = {r["text"]: r for _, r in blocks(sec, "p")}
        self.assertEqual(by_text["Go-live is set for Q2 2027."]["ctx"], {"parent": "Northwind migration"})
        self.assertNotIn("ctx", by_text["Uptime held at 99.9% all year."])      # back at level 0: no parent
        self.assertNotIn("ctx", by_text["Northwind migration"])                 # the title is another shape

    def test_u14_footnote_and_document_date(self):
        deck = P.read_deck(F.footnote_deck(self.d))
        body = {r["text"]: r for _, r in blocks(deck["sections"][1], "p")}
        self.assertEqual(body["Uptime held at 99.9%."]["ctx"]["footnote"], F.FOOTNOTE)
        self.assertNotIn("ctx", body[F.FOOTNOTE])
        self.assertEqual(deck["doc_date"], {"value": "March 2026", "inferred": False})
        undated = P.read_deck(F.footnote_deck(self.d, dated_title=False))
        self.assertEqual(undated["doc_date"], {"value": "2026-04-01", "inferred": True})

    def test_a_quarter_on_a_content_slide_is_not_the_document_date(self):
        self.assertTrue(P.read_deck(F.basic_deck(self.d))["doc_date"]["inferred"])

    def test_residual_rule(self):
        walk = "Uptime by region"
        page = "Uptime by region Northwind Southfield Eastgate 99.1 97.4 95.2 Uptime %"
        res = P.residual_tokens(page, walk)
        self.assertEqual(res[:3], ["northwind", "southfield", "eastgate"])
        self.assertTrue(P.needs_vision(res))
        self.assertFalse(P.needs_vision(["northwind", "southfield"]))           # few words, nothing checkable
        self.assertTrue(P.needs_vision(["word"] * P.RESIDUAL_WORDS))


@unittest.skipUnless(HAVE_LIBS and HAVE_SOFFICE, "needs python-pptx, PyMuPDF and LibreOffice")
class RenderIntegration(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())

    def run_on(self, builder):
        path = builder(self.d)
        res = P.write_outputs(path, self.d / "out")
        secs = json.loads((self.d / "out" / "sections.json").read_text())
        stats = json.loads((self.d / "out" / "stats.json").read_text())
        return res, secs, stats

    def test_u5_image_only_slide_is_flagged_with_its_png(self):
        _, secs, stats = self.run_on(F.image_deck)
        self.assertEqual(secs[0]["paragraphs"], [])
        fig = secs[0]["figures"]
        self.assertEqual(len(fig), 1)
        self.assertEqual((fig[0]["figure"], fig[0]["p_id"], fig[0]["media"]), (1, "s1", "media/slide_01.png"))
        self.assertTrue((self.d / "out" / fig[0]["media"]).is_file())
        self.assertEqual(stats["flagged_slides"], [1])

    def test_u6_chart_values_force_the_vision_path(self):
        path = F.chart_deck(self.d)
        deck = P.read_deck(path)
        self.assertFalse(any(v in deck["walk_text"][1] for v in F.CHART_VALUES))   # the walk cannot see them
        res = render_pages.render(path, self.d / "r", rel_to=self.d)
        page_text = (self.d / "r" / res["pages"][0]["image"]).with_suffix(".txt").read_text()
        residual = P.residual_tokens(page_text, deck["walk_text"][1])
        self.assertTrue(set(F.CHART_VALUES) <= set(residual), residual)
        self.assertTrue(P.needs_vision(residual))
        _, secs, _ = self.run_on(F.chart_deck)
        self.assertEqual(secs[0]["figures"][0]["why"], "chart")

    def test_layout_footer_and_slide_number_are_not_unseen_text(self):
        # real decks render master/layout footers ("… 2026") and slide numbers on every slide: never a reason to flag
        path = F.footer_deck(self.d)
        res = render_pages.render(path, self.d / "r", rel_to=self.d)
        page_text = (self.d / "r" / res["pages"][0]["image"]).with_suffix(".txt").read_text()
        self.assertIn("2026", page_text)                       # the footer does render
        deck = P.read_deck(path)
        self.assertNotIn("confidential", " ".join(r["text"] for _, r in blocks(deck["sections"][0], "p")))
        self.assertEqual(P.residual_tokens(page_text, deck["walk_text"][1]), [])
        _, secs, stats = self.run_on(F.footer_deck)
        self.assertEqual(stats["flagged_slides"], [])

    def test_a_short_text_slide_is_not_sent_to_vision(self):
        _, secs, stats = self.run_on(F.basic_deck)            # ~100 characters: render_pages calls it thin-text
        self.assertEqual((secs[0]["figures"], stats["flagged_slides"]), ([], []))
        kinds = {"picture": False, "chart": False, "diagram": False}
        self.assertEqual(P._why({"flagged": True, "why": "thin-text"}, kinds, []), "")
        self.assertEqual(P._why({"flagged": True, "why": "dense-draw"}, kinds, []), "dense-draw")

    def test_r8_flags_unseen_text_even_without_a_chart_shape(self):
        kinds = {"picture": False, "chart": False, "diagram": False}
        self.assertEqual(P._why({"flagged": False}, kinds, ["99.1", "97.4"]), "text not in shapes")
        self.assertEqual(P._why({"flagged": False}, kinds, []), "")

    def test_u7_hidden_slide_is_skipped_and_images_map_by_visible_index(self):
        _, secs, stats = self.run_on(F.hidden_deck)
        self.assertEqual([s["section_id"] for s in secs], ["s01", "s03", "s04"])
        self.assertEqual(stats["hidden_slides"], [2])
        rendered = json.loads(next((self.d / "out" / "render").glob("*/pages.json")).read_text())
        self.assertEqual(len(rendered["pages"]), 3)
        txt = (self.d / "out" / "render" / rendered["pages"][1]["image"]).with_suffix(".txt").read_text()
        self.assertIn("Region 3", txt)                  # page 2 is slide 3

    def test_batches_carry_doc_date_and_rerun_is_deterministic(self):
        path = F.footnote_deck(self.d)
        P.write_outputs(path, self.d / "out")
        first = {p.name: p.read_bytes() for p in (self.d / "out").glob("*.json")}
        P.write_outputs(path, self.d / "out")
        second = {p.name: p.read_bytes() for p in (self.d / "out").glob("*.json")}
        self.assertEqual(first, second)
        batch = json.loads(first["batch_1.json"])
        self.assertEqual(batch["doc_date"]["value"], "March 2026")


@unittest.skipUnless(HAVE_LIBS, "needs python-pptx and PyMuPDF")
class Refusals(unittest.TestCase):
    """U9: a refusal names the reason, exits 2 and leaves no sections.json."""

    def setUp(self):
        self.d = Path(tempfile.mkdtemp())

    def cli(self, *args):
        return subprocess.run([sys.executable, str(DECK / "pptx_sections.py"), *map(str, args)],
                              capture_output=True, text=True, cwd="/")

    def assert_refused(self, src, needle):
        r = self.cli(src, "--out", self.d / "out")
        self.assertEqual(r.returncode, 2, r.stderr)
        self.assertIn(needle, r.stderr)
        self.assertFalse((self.d / "out" / "sections.json").exists())

    def test_ppt_is_refused(self):
        (self.d / "old.ppt").write_bytes(b"\xd0\xcf\x11\xe0legacy")
        self.assert_refused(self.d / "old.ppt", "reads .pptx only")

    def test_corrupt_package_is_refused(self):
        (self.d / "broken.pptx").write_bytes(b"not a zip")
        self.assert_refused(self.d / "broken.pptx", "not a readable .pptx package")

    def test_missing_file_is_refused(self):
        self.assert_refused(self.d / "nope.pptx", "no such file")

    def test_missing_soffice_is_refused_before_anything_is_written(self):
        path = F.basic_deck(self.d)
        with patch.object(P.render_pages, "soffice_bin", return_value=None):
            with self.assertRaises(P.Refusal) as e:
                P.write_outputs(path, self.d / "out")
        self.assertIn("LibreOffice", str(e.exception))
        self.assertFalse((self.d / "out").exists())

    def test_check_flag(self):
        r = self.cli("--check")
        if HAVE_SOFFICE:
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertRegex(r.stdout, r"^python \S+ python-pptx \S+ pymupdf \S+ soffice ")
        else:
            self.assertEqual(r.returncode, 2)

    def test_check_fails_cleanly_without_libraries(self):
        r = subprocess.run([sys.executable, "-S", "-I", str(DECK / "pptx_sections.py"), "--check"],
                           capture_output=True, text=True, cwd="/")
        self.assertEqual(r.returncode, 2)
        self.assertIn("deck-fact-check needs Python with python-pptx and PyMuPDF", r.stderr)


if __name__ == "__main__":
    unittest.main()
