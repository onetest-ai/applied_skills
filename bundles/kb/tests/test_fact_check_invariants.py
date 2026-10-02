"""Eval invariants for a /kb:fact-check run: what any run must satisfy whatever the Brain or corpus."""
from __future__ import annotations

import hashlib
import re
import tempfile
import unittest
from pathlib import Path

from test_fact_check_code import HAVE_DOCX, load_helpers

if HAVE_DOCX:
    from fact_check_invariants import check_run


def finding(id, verdict, severity, destination, quote="", **kw):
    return dict(id=id, verdict=verdict, severity=severity, destination=destination, quote=quote,
                section="1", evidence="Brain: x (2026-01-01, f.md)", fix="Fix it", source="f.md", **kw)


GOOD = [
    finding("C01", "Incorrect", "Major", "Word comment", "40 million in"),
    finding("C02", "Misleading", "Minor", "Word comment", "green fields"),
    finding("C03", "No Evidence", "Minor", "log only"),
    finding("C04", "Verified", "Minor", "count only"),
]


@unittest.skipUnless(HAVE_DOCX, "needs python-docx >= 1.2")
class TestRunInvariants(unittest.TestCase):
    def setUp(self):
        from docx import Document
        self.tmp = Path(tempfile.mkdtemp())
        d = Document()
        p = d.add_paragraph("Revenue was 40 million in 2025 across green fields.")
        p.add_run(" And more.")
        self.orig = self.tmp / "orig.docx"
        d.save(self.orig)
        self.sha = hashlib.sha256(self.orig.read_bytes()).hexdigest()
        self.out = self.tmp / "out.docx"

    def _annotate(self, findings):
        self.h = load_helpers()
        self.h["annotate"](str(self.orig), str(self.out), [f for f in findings if "comment" in f["destination"].lower()])

    def _check(self, findings, **kw):
        return check_run(findings, self.out, original_sha256=kw.pop("sha", self.sha),
                         original_path=kw.pop("original_path", self.orig), **kw)

    def test_a_correct_run_has_no_violations(self):
        self._annotate(GOOD)
        self.assertEqual(self._check(GOOD), [])

    def test_a_minor_defect_that_was_only_logged_is_a_violation(self):
        bad = [dict(f) for f in GOOD]
        bad[1]["destination"] = "log only"
        self._annotate(bad)
        self.assertTrue(any("C02" in v and "comment" in v for v in self._check(bad)))

    def test_no_evidence_must_be_minor_and_uncommented(self):
        bad = [dict(f) for f in GOOD]
        bad[2].update(severity="Major", destination="Word comment", quote="40 million in")
        self._annotate(bad)
        vs = self._check(bad)
        self.assertTrue(any("C03" in v and "No Evidence" in v for v in vs))

    def test_verified_claim_must_not_be_commented(self):
        bad = [dict(f) for f in GOOD]
        bad[3].update(destination="Word comment", quote="And more.")
        self._annotate(bad)
        self.assertTrue(any("C04" in v and "Verified" in v for v in self._check(bad)))

    def test_duplicate_claim_ids_are_a_violation(self):
        bad = GOOD + [finding("C01", "Verified", "Minor", "count only")]
        self._annotate(GOOD)
        self.assertTrue(any("duplicate" in v.lower() and "C01" in v for v in self._check(bad)))

    def test_unknown_verdict_or_severity_is_a_violation(self):
        bad = [dict(f) for f in GOOD]
        bad[0].update(verdict="Wrong", severity="Huge")
        self._annotate(GOOD)
        vs = self._check(bad)
        self.assertTrue(any("verdict" in v.lower() for v in vs))
        self.assertTrue(any("severity" in v.lower() for v in vs))

    def test_comment_must_be_anchored_to_the_quoted_words_only(self):
        self._annotate(GOOD)
        bad = [dict(f) for f in GOOD]
        bad[0]["quote"] = "Revenue was 40 million in 2025 across green fields."   # wider than what was highlighted
        self.assertTrue(any("C01" in v and "anchor" in v.lower() for v in self._check(bad)))

    def test_comment_count_must_match_the_approved_set(self):
        self._annotate(GOOD)
        bad = [dict(f) for f in GOOD]
        bad.append(finding("C05", "Incorrect", "Major", "Word comment", "And more."))
        self.assertTrue(any("C05" in v for v in self._check(bad)))

    def test_comment_longer_than_sixty_words_is_a_violation(self):
        fs = [dict(f) for f in GOOD]
        fs[0]["fix"] = "word " * 70
        self._annotate(fs)
        self.assertTrue(any("C01" in v and "60 words" in v for v in self._check(fs)))

    def test_original_document_must_be_untouched(self):
        self._annotate(GOOD)
        self.assertTrue(any("original" in v.lower() for v in self._check(GOOD, sha="0" * 64)))

    def test_output_must_not_overwrite_the_original(self):
        self._annotate(GOOD)
        self.assertTrue(any("overwrite" in v.lower() for v in
                            check_run(GOOD, self.orig, original_sha256=self.sha, original_path=self.orig)))

    def test_xml_special_characters_in_the_quote_are_not_false_violations(self):
        from docx import Document
        d = Document()
        d.add_paragraph("Green means < 3 Seconds & AT&T stays.")
        d.save(self.orig)
        self.sha = hashlib.sha256(self.orig.read_bytes()).hexdigest()
        fs = [finding("C01", "Misleading", "Major", "Word comment", "< 3 Seconds & AT&T")]
        self._annotate(fs)
        self.assertEqual(self._check(fs), [])

    def test_verified_claims_need_no_severity_but_defects_do(self):
        fs = [dict(f) for f in GOOD]
        fs[3]["severity"] = "-"
        self._annotate(fs)
        self.assertEqual(self._check(fs), [])
        fs[0]["severity"] = "-"
        self.assertTrue(any("C01" in v and "severity" in v for v in self._check(fs)))

    # --- review9 findings 1, 4, 7 ---
    def test_comment_destined_finding_without_id_is_a_violation_not_a_keyerror(self):
        self._annotate(GOOD)
        bad = [dict(f) for f in GOOD] + [finding("X", "Incorrect", "Major", "Word comment", "green fields")]
        del bad[-1]["id"]
        vs = self._check(bad)
        self.assertTrue(any("missing id" in v for v in vs), vs)

    def test_two_findings_without_id_are_each_missing_id_never_duplicates(self):
        self._annotate(GOOD)
        a, b = finding("X", "Verified", "Minor", "count only"), finding("Y", "Verified", "Minor", "count only")
        del a["id"], b["id"]
        vs = self._check(GOOD + [a, b])
        self.assertEqual(sum("missing id" in v for v in vs), 2, vs)
        self.assertFalse(any("duplicate" in v.lower() for v in vs), vs)

    def _rewrite_xml(self, fn):
        import zipfile
        tmp = self.tmp / "rw.docx"
        with zipfile.ZipFile(self.out) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
            for item in zi.infolist():
                data = zi.read(item.filename)
                if item.filename == "word/document.xml":
                    data = fn(data.decode("utf-8")).encode("utf-8")
                zo.writestr(item, data)
        tmp.replace(self.out)

    def test_range_texts_tolerates_extra_attributes_in_either_order(self):
        from fact_check_invariants import _range_texts
        self._annotate(GOOD)
        base = _range_texts(self.out)
        self.assertTrue(base)

        def extra(xml):
            xml = re.sub(r'<w:commentRangeStart w:id="(\d+)"/>',
                         r'<w:commentRangeStart w:displacedByCustomXml="next" w:id="\1"/>', xml)
            return re.sub(r'<w:commentRangeEnd w:id="(\d+)"/>',
                          r'<w:commentRangeEnd w:id="\1" w:displacedByCustomXml="next"/>', xml)
        self._rewrite_xml(extra)
        self.assertEqual(_range_texts(self.out), base)

    # RC-2: figures read but never claimed
    def _png(self):
        import struct, zlib
        def ch(t, d):
            c = struct.pack(">I", len(d)) + t + d
            return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
        raw = b"\x00\xff\x00\x00"
        return (b"\x89PNG\r\n\x1a\n" + ch(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
                + ch(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + ch(b"IEND", b""))

    def _with_image(self):
        from docx import Document
        png = self.tmp / "fig.png"
        png.write_bytes(self._png())
        d = Document()
        d.add_paragraph("Revenue was 40 million in 2025 across green fields.")
        d.add_picture(str(png))
        d.save(self.orig)
        self.sha = hashlib.sha256(self.orig.read_bytes()).hexdigest()

    def test_figures_present_but_no_figure_findings_is_a_violation(self):
        self._with_image()
        self._annotate(GOOD)
        self.assertIn("figures present but no I* findings", self._check(GOOD))

    def test_figure_finding_satisfies_the_figure_check(self):
        self._with_image()
        good = GOOD + [finding("I01", "Verified", "Minor", "count only", type="Figure")]
        self._annotate(good)
        self.assertNotIn("figures present but no I* findings", self._check(good))

    def test_no_images_means_no_figure_violation(self):
        self._annotate(GOOD)
        self.assertNotIn("figures present but no I* findings", self._check(GOOD))

    # RC-3: two-sided check must be recorded
    def test_verified_topo_without_checks_is_a_violation(self):
        fs = GOOD + [finding("C05", "Verified", "Minor", "count only", type="TOPO")]
        self._annotate(fs)
        self.assertTrue(any("C05" in v and "checks" in v for v in self._check(fs)))

    def test_verified_time_with_only_one_check_is_a_violation(self):
        fs = GOOD + [finding("C05", "Verified", "Minor", "count only", type="TIME",
                             checks=["search", "search latest_only=false"])]
        self._annotate(fs)
        self.assertTrue(any("C05" in v and "opposing" in v for v in self._check(fs)))

    def test_verified_topo_with_both_checks_is_clean(self):
        fs = GOOD + [finding("C05", "Verified", "Minor", "count only", type="TOPO",
                             checks=["search", "search latest_only=false", "opposing"])]
        self._annotate(fs)
        self.assertEqual(self._check(fs), [])

    def test_non_verified_or_num_findings_need_no_checks(self):
        fs = GOOD + [finding("C05", "No Evidence", "Minor", "log only", type="STATUS"),
                     finding("C06", "Verified", "Minor", "count only", type="NUM")]
        self._annotate(fs)
        self.assertEqual(self._check(fs), [])

    def test_invariants_cli_ships_in_the_skill_dir_and_runs(self):
        import subprocess
        import sys
        from test_plugin_structure import KB_ROOT
        cli = KB_ROOT / "skills" / "fact-check" / "fact_check_invariants.py"
        self.assertTrue(cli.exists())
        self.assertFalse((KB_ROOT / "tests" / "fact_check_invariants.py").exists())
        r = subprocess.run([sys.executable, str(cli), "--help"], capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("fact_check_invariants.py", r.stdout)


SKILL_MD = Path(__file__).resolve().parents[1] / "skills" / "fact-check" / "SKILL.md"



class TestTwoSidedCheckText(unittest.TestCase):
    """L1 text assertions: step 4 makes the agent retrieve the opposing side before Verified."""

    @classmethod
    def setUpClass(cls):
        cls.text = SKILL_MD.read_text(encoding="utf-8")
        a = cls.text.index("### 4.")
        b = cls.text.index("### 5.")
        c = cls.text.index("### 6.")
        cls.step4 = cls.text[a:b]
        cls.step5 = cls.text[b:c]

    def _bullet(self):
        lines = [l for l in self.step4.splitlines() if "Two-sided check" in l]
        self.assertTrue(lines, "step 4 has no 'Two-sided check' bullet")
        return lines[0]

    def test_step4_has_two_sided_check_label(self):
        self._bullet()

    def test_step4_names_the_opposing_search_and_explicit_none(self):
        self.assertIn("latest_only=false", self.step4)
        self.assertIn("opposing", self.step4)
        self.assertIn("no opposing span found", self.step4)

    def test_step5_controversial_cites_both_spans(self):
        self.assertIn("cite both spans", self.step5)

    def test_bullet_names_the_num_exclusion(self):
        b = self._bullet()
        self.assertIn("NUM", b)
        self.assertRegex(b, r"skipped for NUM")

    def test_description_still_opens_with_use_when(self):
        m = re.search(r"^description:\s*(.+)$", self.text, re.M)
        self.assertTrue(m and m.group(1).strip("\"' >").startswith("Use when "))


class TestBaselineReuseText(unittest.TestCase):
    """Text assertions: a rerun must not silently reuse a weak baseline."""

    @classmethod
    def setUpClass(cls):
        cls.text = SKILL_MD.read_text(encoding="utf-8")
        a = cls.text.index("## Inputs")
        b = cls.text.index("## Procedure")
        cls.inputs = cls.text[a:b]
        s = cls.text.index("### 10.")
        cls.step10 = cls.text[s:cls.text.index("## Rules")]

    def test_inputs_has_reusing_a_baseline_label(self):
        self.assertIn("Reusing a baseline", self.inputs)

    def test_reuse_conditions_name_version_checksum_and_claim_count(self):
        for term in ("knowledge_version", "checksum", "claim count"):
            self.assertIn(term, self.inputs)

    def test_reused_baseline_is_announced_in_first_line(self):
        self.assertIn("reused baseline of", self.inputs)

    def test_reused_verified_claims_need_the_two_sided_record(self):
        self.assertIn("two-sided record", self.inputs)

    def test_step10_baseline_records_checksum_version_and_count(self):
        line = next((l for l in self.step10.splitlines() if "baseline.md" in l), "")
        for term in ("checksum", "knowledge_version", "claim count"):
            self.assertIn(term, line)



if __name__ == "__main__":
    unittest.main()
