"""Evidence-backed comments, identity errors and idempotent anchors shared by PowerPoint and Word."""
from __future__ import annotations

import sys
import unittest

from test_plugin_structure import KB_ROOT

DOC = KB_ROOT / "skills" / "doc-fact-check"
DECK = KB_ROOT / "skills" / "deck-fact-check"
sys.path.insert(0, str(DOC))

import fact_check_invariants as I  # noqa: E402


def _f(**kw):
    f = {"id": "C01", "p_id": "s3-sh5-p1", "section": "Slide 3", "quote": "One ERP job failure raises a priority ticket",
         "verdict": "Incorrect", "severity": "Major", "evidence": "'two or three successive failures' raise a ticket",
         "fix": "Use two or three.", "source": "transcripts/Atlas Session - RouteDB-20260925_145826UTC-Meeting Recording.mp4.md",
         "conflicts": [], "why": "a priority ticket is logged only after two or three failures in a row",
         "suggested": "Two or three ERP job failures raise a priority ticket."}   # fits the slide box (DeckSuggestionFitsTheBox)
    f.update(kw)
    return f


class SourceLabel(unittest.TestCase):
    def test_meeting_recording(self):
        self.assertEqual(I.source_label("transcripts/Atlas Session - RouteDB-20260925_145826UTC-Meeting Recording.mp4.md"),
                         ("Atlas Session – RouteDB", "25 Sep 2026"))

    def test_dated_document_prefix(self):
        self.assertEqual(I.source_label("docs/team__Board Notes__260406 - Steering Pack v2.pptx.md"),
                         ("Steering Pack v2", "6 Apr 2026"))

    def test_iso_date_in_name(self):
        self.assertEqual(I.source_label("transcripts/From Kim__2026-09-17-routing-migration.txt.md"),
                         ("Routing migration", "17 Sep 2026"))

    def test_lower_case_slug_reads_as_a_name(self):
        # a file-name slug is lower case: first letter capitalised, a vowel-less short token is an acronym
        self.assertEqual(I.source_label("transcripts/From Kim__2026-09-15-ordering-crm-alignment.txt.md"),
                         ("Ordering CRM alignment", "15 Sep 2026"))

    def test_undated(self):
        self.assertEqual(I.source_label("docs/Platform Overview.pdf.md"), ("Platform Overview", "undated"))


class CommentShape(unittest.TestCase):
    def test_deck_comment(self):
        self.assertEqual(I.comment_text(_f()),
                         '[Incorrect]\n"One ERP job failure raises a priority ticket"\n\n'
                         "Why: Atlas Session – RouteDB (25 Sep 2026) says a priority ticket is logged only after two or "
                         "three failures in a row.\n\n"
                         'Suggested: "Two or three ERP job failures raise a priority ticket."')

    def test_word_comment_has_no_quote_line(self):
        t = I.comment_text(_f(p_id="p12"))
        self.assertTrue(t.startswith("[Incorrect]\nWhy: Atlas Session – RouteDB (25 Sep 2026) says"))
        self.assertNotIn('"One ERP', t)

    def test_no_ids_no_severity_in_any_comment(self):
        for f in (_f(), _f(p_id="p12")):
            t = I.comment_text(f)
            self.assertNotIn("Major", t); self.assertNotIn("C01", t); self.assertNotIn("·", t.split("\n")[0])

    def test_header_pattern_reads_the_category(self):
        self.assertEqual(I.HEADER.match(I.comment_text(_f()))["verdict"], "Incorrect")

    def test_quoted_body_lines_do_not_bypass_the_word_limit(self):
        f = _f(suggested='A correction.\n"' + "word " * 100 + '"')
        text = I.comment_text(f)
        self.assertGreater(I.comment_words(text, True), I.MAX_COMMENT_WORDS)

    def test_owner_found_by_its_text(self):
        fs = [_f(), _f(id="C02", quote="Another statement")]
        self.assertEqual(I.comment_owner(I.comment_text(fs[1]), fs), "C02")

    def test_owner_recognizes_every_supported_field_selection(self):
        from itertools import combinations
        f = _f()
        for n in range(1, len(I.COMMENT_FIELDS) + 1):
            for fields in combinations(I.COMMENT_FIELDS, n):
                with self.subTest(fields=fields):
                    self.assertEqual(I.comment_owner(I.comment_text(f, fields), [f]), "C01")

    def test_reordered_or_reduced_findings_do_not_duplicate_fanned_comments(self):
        import tempfile
        from pathlib import Path
        import deck_fixtures as F
        import pptx_annotate as A
        d = Path(tempfile.mkdtemp())
        source = F.full_deck(d)
        fs = [_f(id="C01", p_id="s1-sh3-p1", quote="First statement"),
              _f(id="C02", p_id="s1-sh3-p2", quote="Second statement")]
        out = d / "out.pptx"
        A.annotate(source, out, fs)
        for n, selected in enumerate((fs[::-1], fs[1:])):
            again = d / f"again{n}.pptx"
            A.annotate(out, again, selected)
            self.assertEqual(len(I.read_pptx_comments(again)), 2)

    def test_identical_comments_on_distinct_slides_and_shapes_are_not_suppressed(self):
        import tempfile
        from pathlib import Path
        from pptx import Presentation
        from pptx.util import Inches
        import pptx_annotate as A
        import annotate_cli

        d = Path(tempfile.mkdtemp())
        deck = Presentation()
        findings = []
        for slide_number in (1, 2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for y in (1, 3):
                shape = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(5), Inches(1))
                shape.text = _f()["quote"]
                findings.append(_f(id=f"C{len(findings) + 1:02d}", p_id=f"s{slide_number}-sh{shape.shape_id}-p1"))
        source, partial, out, again = (d / f"{name}.pptx" for name in ("source", "partial", "out", "again"))
        deck.save(source)
        A.annotate(source, partial, findings[:1])
        written, skipped = A.annotate(partial, out, findings)
        self.assertEqual((len(written), skipped), (4, []))
        self.assertEqual(len(I.read_pptx_comments(out)), 4)
        marked = annotate_cli.mark_destinations(findings, written)
        self.assertEqual(I.check_run(marked, out, run={"comments": "major"}), [])
        A.annotate(out, again, findings)
        self.assertEqual(len(I.read_pptx_comments(again)), 4)


class OnlyWithHardEvidence(unittest.TestCase):
    def test_a_supported_defect_is_commented(self):
        self.assertTrue(I.wants_comment(_f()))

    def test_disagreement_goes_to_the_findings_page(self):
        self.assertFalse(I.wants_comment(_f(verdict="Controversial")))

    def test_confirm_is_not_a_fix(self):
        for fix in ("Confirm the cutoff with operations.", "Check with the team.", "Verify the date with the owner."):
            self.assertFalse(I.wants_comment(_f(fix=fix, suggested="")), fix)

    def test_no_why_no_comment(self):
        self.assertFalse(I.wants_comment(_f(why="")))

    def test_minor_still_not_commented_under_major(self):
        self.assertFalse(I.wants_comment(_f(severity="Minor")))


class SupportedIdentityErrors(unittest.TestCase):
    def test_supported_identity_error_is_commented_without_a_known_provider(self):
        for pid in ("p12", "s3-sh5-p1"):
            f = _f(p_id=pid, quote="Rowan provides hosting", evidence="Rowan is an engagement team member.",
                   why="Rowan is a member of the engagement team, not the hosting company named by this statement",
                   fix="Confirm the actual hosting provider.",
                   suggested="Hosting provider: not established")
            self.assertTrue(I.wants_comment(f))
            self.assertIsNone(I.why_problem(f))
            self.assertIn("[Incorrect]", I.comment_text(f))
            self.assertNotIn("Controversial", I.comment_text(f))

    def test_unresolved_identity_disagreement_stays_on_the_findings_page(self):
        self.assertFalse(I.wants_comment(_f(verdict="Controversial")))

    def test_word_table_comments_use_the_readers_row_identifiers(self):
        import tempfile
        from pathlib import Path
        from docx import Document
        import annotate
        import annotate_cli
        import sections
        d = Path(tempfile.mkdtemp())
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        for row in table.rows:
            row.cells[0].text = "The same incorrect statement."
        source, out, again = (d / f"{name}.docx" for name in ("source", "out", "again"))
        doc.save(source)
        extracted = sections.extract_sections(source)
        rows = [rec for section in extracted for kind, rec, _ in section["blocks"] if kind == "r"]
        fs = [_f(id=f"C{n:02d}", p_id=row["r_id"], quote=row["cells"][0]) for n, row in enumerate(rows, 1)]
        written, skipped = annotate.annotate(source, out, fs)
        self.assertEqual((len(written), skipped), (2, []))
        marked = annotate_cli.mark_destinations(fs, written)
        self.assertEqual(I.check_run(marked, out, run={"comments": "major"}), [])
        annotate.annotate(out, again, fs)
        self.assertEqual(len(Document(again).comments), 2)

    def test_identical_text_comments_on_distinct_word_paragraphs_are_not_suppressed(self):
        import tempfile
        from pathlib import Path
        from docx import Document
        import annotate
        import annotate_cli
        d = Path(tempfile.mkdtemp())
        doc = Document()
        for n in (1, 2):
            doc.add_paragraph("The same incorrect statement.")
        source, out, again = (d / f"{name}.docx" for name in ("source", "out", "again"))
        doc.save(source)
        fs = [_f(id=f"C0{n}", p_id=f"p{n}", quote="The same incorrect statement.") for n in (1, 2)]
        annotate.annotate(source, out, fs[:1])
        written, _ = annotate.annotate(out, again, fs)
        self.assertEqual(len(Document(again).comments), 2)
        self.assertEqual(I.check_run(annotate_cli.mark_destinations(fs, written), again, run={"comments": "major"}), [])

    def test_drawing_comment_with_text_in_its_range_validates_and_reruns_once(self):
        import tempfile
        from pathlib import Path
        from docx import Document
        import deck_fixtures as F
        import annotate
        import annotate_cli
        d = Path(tempfile.mkdtemp())
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Before ").add_picture(str(F._png(d / "first.png")))
        p.add_run("Between ")
        p.add_run().add_picture(str(F._png(d / "second.png")))
        source, out, again = (d / f"{name}.docx" for name in ("source", "out", "again"))
        doc.save(source)
        f = _f(id="I01", p_id="p1", anchor="drawing", figure=1, quote="The pictures show an incorrect count")
        written, _ = annotate.annotate(source, out, [f])
        self.assertTrue(I._range_texts(out)["0"])
        self.assertEqual(I.check_run(annotate_cli.mark_destinations([f], written), out, run={"comments": "major"}), [])
        annotate.annotate(out, again, [f])
        self.assertEqual(len(Document(again).comments), 1)

    def test_drawing_reruns_match_the_figure_when_its_quote_is_not_text(self):
        import tempfile
        from pathlib import Path
        from docx import Document
        import deck_fixtures as F
        import annotate
        d = Path(tempfile.mkdtemp())
        doc = Document()
        for n in (1, 2):
            doc.add_picture(str(F._png(d / f"image{n}.png")))
        source, out, again = (d / f"{name}.docx" for name in ("source", "out", "again"))
        doc.save(source)
        fs = [_f(id=f"I0{n}", p_id=f"p{n}", anchor="drawing", figure=n, quote="A picture shows the wrong count") for n in (1, 2)]
        annotate.annotate(source, out, fs[:1])
        written, _ = annotate.annotate(out, again, fs)
        self.assertEqual(len(Document(again).comments), 2)
        import annotate_cli
        self.assertEqual(I.check_run(annotate_cli.mark_destinations(fs, written), again, run={"comments": "major"}), [])
        annotate.annotate(again, d / "final.docx", fs)
        self.assertEqual(len(Document(d / "final.docx").comments), 2)

    def test_identity_error_is_written_once_in_both_formats(self):
        import tempfile
        from pathlib import Path
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        import annotate
        import annotate_cli
        import pptx_annotate

        d = Path(tempfile.mkdtemp())
        quote = "Rowan provides hosting"
        doc = Document()
        doc.add_paragraph(quote)
        doc.save(d / "draft.docx")
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        shape.text = quote
        deck.save(d / "draft.pptx")
        for ext, pid, writer in (("docx", "p1", annotate.annotate),
                                 ("pptx", f"s1-sh{shape.shape_id}-p1", pptx_annotate.annotate)):
            f = _f(p_id=pid, quote=quote, evidence="Rowan is an engagement team member.",
                   why="Rowan is a team member, not the hosting company named here",
                   fix="Confirm the hosting provider.", source="Team Session-20261001.md",
                   suggested="Hosting provider: not established")
            source, out, again = (d / f"{name}.{ext}" for name in ("draft", "out", "again"))
            selected = [f] if I.wants_comment(f) else []
            written, skipped = writer(source, out, selected)
            self.assertEqual((written, skipped), (["C01"], []))
            marked = annotate_cli.mark_destinations([f], written)
            self.assertEqual(I.check_run(marked, out, run={"comments": "major"}), [])
            writer(out, again, selected)
            self.assertEqual(len(I.doc_view(again).comments()), 1)


class WhyChecks(unittest.TestCase):
    def test_a_good_why_passes(self):
        self.assertEqual(I.why_problem(_f()), None)

    def test_too_long(self):
        self.assertIn("45 words", I.why_problem(_f(why=" ".join(["word"] * 46))))

    def test_edit_instruction_is_not_a_suggestion(self):
        self.assertIn("full sentence", I.why_problem(_f(suggested="Change 'One' to 'two or three'")))

    def test_a_figure_in_why_must_be_in_the_evidence(self):
        self.assertIn("19,000", I.why_problem(_f(why="the team counts 19,000 calls a day")))

    def test_jargon_refused(self):
        self.assertIn("Brain", I.why_problem(_f(why="the Brain shows two or three failures")))


class Briefs(unittest.TestCase):
    def test_verify_brief_asks_for_why_and_suggested(self):
        v = (DOC / "verify.md").read_text(encoding="utf-8")
        for s in ("`why`", "`suggested`", "date unknown", "the same speaker revises it in the same exchange"):
            self.assertIn(s, v)

    def test_both_skills_describe_the_shape(self):
        for skill in (DOC, DECK):
            s = (skill / "SKILL.md").read_text(encoding="utf-8")
            with self.subTest(skill=skill.name):
                self.assertIn("[Incorrect]", s)
                self.assertIn("findings page only", s)


if __name__ == "__main__":
    unittest.main()


class MergeRefusesACommentWithoutWhy(unittest.TestCase):
    """A defect the comment policy would comment, written without a plain-language why, must be re-done by its
    worker — silently leaving it uncommented would turn a caught error into a miss."""

    def _errs(self, **kw):
        import merge_findings as MF
        f = _f(claim="c", type="NUM", confidence="High", trail="t", sources=[], **kw)
        errs = []
        MF._check_finding(f, "chunk 1", errs)
        return errs

    def test_missing_why_refused(self):
        self.assertTrue(any("why" in x for x in self._errs(why="", suggested="")))

    def test_exhaustive_minor_without_why_is_refused(self):
        import merge_findings as MF
        f = _f(claim="c", type="NUM", confidence="High", trail="t", sources=[], severity="Minor", why="", suggested="")
        errors = []
        MF._check_finding(f, "chunk 1", errors, comments="all")
        self.assertTrue(any("why" in e for e in errors))

    def test_good_why_accepted(self):
        self.assertEqual([x for x in self._errs() if "why" in x or "suggested" in x], [])

    def test_disagreement_needs_no_why(self):
        self.assertEqual([x for x in self._errs(verdict="Controversial", why="", suggested="",
                          conflicts=[{"span": "s", "source": "a.md", "date": "2026-01-01"}]) if "why" in x], [])


class DeckSuggestionFitsTheBox(unittest.TestCase):
    """A deck Suggested sentence replaces text inside a fixed-size slide box, so it keeps about the original length:
    at most 120% of the quoted statement (or 15 characters more, for a short label). The Why carries the explanation.
    User decision, 20 Oct: 'if 60 chars -> 60 chars, not 183'."""

    QUOTE = "IVA looks up the account by phone number and routes the call"          # 60 characters

    def _deck(self, suggested):
        return _f(p_id="s3-sh29-p1", quote=self.QUOTE, suggested=suggested)

    def test_a_suggestion_three_times_the_slide_text_is_refused(self):
        long = ("The IVA looks up the account by phone number; delivery changes can be completed through IVA "
                "self-service, with agent routing only when self-service fails or the customer requests one.")
        msg = I.why_problem(self._deck(long))
        self.assertIn("183 characters", msg)
        self.assertIn("75", msg)                                   # the limit for a 60-character statement

    def test_a_suggestion_of_about_the_same_length_passes(self):
        self.assertIsNone(I.why_problem(self._deck("IVA finds the account by phone; delivery changes self-serve")))

    def test_a_short_label_gets_room_for_the_correction(self):
        f = _f(p_id="s3-sh40-p1", quote="Help-desk volume: ~12K calls per month", suggested="Help-desk volume: ~12K calls per day (2025 YTD)")
        self.assertIsNone(I.why_problem(f))

    def test_a_word_suggestion_has_no_box_and_no_length_limit(self):
        f = _f(p_id="p12", quote=self.QUOTE, suggested="x " * 120)
        self.assertIsNone(I.why_problem(f))
