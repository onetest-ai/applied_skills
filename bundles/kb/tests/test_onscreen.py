"""Only what is on the slide is analysed (spec F7, eval 11). Text in a box outside the slide, text pushed past the
slide's edge and hidden-flag shapes are removed from the analysis entirely (blocks, context, statements, document
date) and reported as left out, never silently."""
from __future__ import annotations

import json
import re
import tempfile
import unittest
from pathlib import Path

from test_plugin_structure import KB_ROOT

import sys
DECK = KB_ROOT / "skills" / "deck-fact-check"
sys.path.insert(0, str(DECK))

try:
    import pptx  # noqa: F401
    import pymupdf  # noqa: F401
    import deck_fixtures as F
    import pptx_sections as P
    from render_pages import soffice_bin
    HAVE = soffice_bin() is not None
    HAVE_PPTX = True
except ImportError:
    HAVE = HAVE_PPTX = False

sq = lambda s: re.sub(r"\W+", "", str(s).casefold())


def everything_seen(sections: list[dict]) -> str:
    """All text the analysis can see: paragraphs, row cells and every ctx value."""
    parts = []
    for s in sections:
        for p in s.get("paragraphs", []):
            parts += [p["text"], json.dumps(p.get("ctx", {}), ensure_ascii=False)]
        for t in s.get("tables", []):
            for r in t["rows"]:
                parts += [" ".join(r["cells"]), json.dumps(r.get("ctx", {}), ensure_ascii=False)]
    return sq(" ".join(parts))


@unittest.skipUnless(HAVE_PPTX, "needs python-pptx")
class GeometryOnly(unittest.TestCase):
    """Without a render, the shape walk still drops what is provably not on the slide."""

    def test_rejected_title_and_lede_paragraphs_never_reenter_context(self):
        from pptx import Presentation
        from pptx.util import Inches
        d = Path(tempfile.mkdtemp())
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(0.5))
        title.text = "Visible delivery scenario"
        title.text_frame.add_paragraph().text = "Invisible alternate scenario"
        lede = slide.shapes.add_textbox(Inches(1), Inches(1.1), Inches(6), Inches(1))
        lede.text = "A customer calls an agent to change tomorrow's delivery."
        lede.text_frame.add_paragraph().text = "Hidden renewal process changes the meaning of the card."
        card = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
        card.text = "The portal owns the change."
        path = d / "partial.pptx"
        prs.save(path)
        rendered = " ".join((title.text_frame.paragraphs[0].text, lede.text_frame.paragraphs[0].text, card.text))
        result = P.read_deck(path, page_text={1: rendered})["sections"][0]
        self.assertEqual(result["slide_context"]["title"], "Visible delivery scenario")
        self.assertEqual(result["slide_context"]["lede"], lede.text_frame.paragraphs[0].text)
        self.assertNotIn("Invisible", json.dumps(result))
        self.assertNotIn("Hidden renewal", json.dumps(result))

    def test_boxes_outside_the_slide_and_hidden_shapes_are_not_blocks(self):
        deck = P.read_deck(F.offscreen_deck(Path(tempfile.mkdtemp())))
        seen = everything_seen([{"paragraphs": [r for k, r, _ in s["blocks"] if k == "p"],
                                 "tables": [{"rows": [r for k, r, _ in s["blocks"] if k == "r"]}]} for s in deck["sections"]])
        for gone in F.OFFSCREEN_DROP[:3]:          # right of the slide, left of the slide, hidden="1"
            self.assertNotIn(sq(gone), seen, gone)
        reasons = {o["reason"] for o in deck["offscreen"]}
        self.assertEqual(reasons, {"outside the slide", "hidden shape"})
        self.assertEqual(sum(o["blocks"] for o in deck["offscreen"]), 3)

    def test_partial_paragraph_or_short_value_cannot_pass_by_majority(self):
        self.assertFalse(P.on_page("All orders reach the driver except failed syncs", sq("All orders reach the driver")))
        self.assertFalse(P.on_page("9", sq("Metric")))
        self.assertFalse(P.row_on_page(["Tickets", "9"], sq("Tickets")))


@unittest.skipUnless(HAVE, "needs python-pptx and PyMuPDF")
class PaintedText(unittest.TestCase):
    def setUp(self):
        import pymupdf
        from pptx_visibility import RenderedText
        self.doc = pymupdf.open()
        self.page = self.doc.new_page(width=720, height=540)
        self.reader = RenderedText

    def tearDown(self):
        self.doc.close()

    def test_later_opaque_rectangle_hides_text_but_earlier_background_does_not(self):
        self.page.draw_rect((20, 20, 600, 150), fill=(1, 1, 1))
        self.page.insert_text((40, 60), "Correct visible statement")
        self.page.insert_text((40, 100), "Invisible incorrect statement")
        self.page.draw_rect((30, 80, 400, 120), fill=(1, 1, 1), overlay=True)
        text = self.reader(self.page)
        self.assertTrue(text.locations("Correct visible statement"))
        self.assertFalse(text.locations("Invisible incorrect statement"))
        self.assertNotIn("Invisible", text.text)

    def test_outline_and_translucent_shapes_do_not_erase_visible_text(self):
        self.page.insert_text((40, 60), "Visible through an outline")
        self.page.draw_rect((30, 40, 400, 80))
        self.page.draw_rect((30, 40, 400, 80), fill=(1, 1, 1), fill_opacity=0.2)
        self.assertTrue(self.reader(self.page).locations("Visible through an outline"))

    def test_curved_cover_and_white_on_white_text_are_excluded(self):
        self.page.insert_text((40, 60), "Invisible under curved cover")
        self.page.draw_oval((5, 20, 400, 100), fill=(1, 1, 1), color=None)
        self.page.insert_text((40, 140), "Invisible white text", color=(1, 1, 1))
        text = self.reader(self.page)
        self.assertFalse(text.locations("Invisible under curved cover"))
        self.assertFalse(text.locations("Invisible white text"))

    def test_cover_touching_only_glyph_padding_preserves_visible_text(self):
        self.page.insert_text((40, 60), "Readable statement", fontsize=20)
        self.page.draw_rect((39, 44, 41, 48), fill=(1, 1, 1), color=None)
        self.assertTrue(self.reader(self.page).locations("Readable statement"))

    def test_hidden_duplicate_cannot_borrow_visible_text_at_another_anchor(self):
        self.page.insert_text((40, 60), "Same statement in two places")
        self.page.insert_text((40, 200), "Same statement in two places")
        self.page.draw_rect((30, 180, 400, 220), fill=(1, 1, 1))
        text = self.reader(self.page)
        self.assertTrue(text.contains("Same statement in two places", {"x": 30, "y": 40, "w": 400, "h": 40}, 720))
        scale = 12700
        hidden_box = {k: v * scale for k, v in {"x": 30, "y": 180, "w": 400, "h": 40}.items()}
        self.assertFalse(text.contains("Same statement in two places", hidden_box, 720 * scale))

    def test_partial_sentence_and_value_at_page_edge_are_not_whole_claims(self):
        self.page.insert_text((650, 100), "Visible start with an invisible qualification", fontsize=18)
        self.page.insert_text((40, 539), "9", fontsize=20)
        text = self.reader(self.page)
        self.assertFalse(text.locations("Visible start with an invisible qualification"))
        self.assertFalse(text.locations("9"))

    def test_fully_visible_text_can_overflow_the_textbox(self):
        self.page.insert_text((40, 60), "A fully visible sentence beyond a narrow text box")
        scale = 12700
        box = {k: v * scale for k, v in {"x": 35, "y": 40, "w": 20, "h": 40}.items()}
        self.assertTrue(self.reader(self.page).contains("A fully visible sentence beyond a narrow text box", box,
                                                       720 * scale))

    def test_covered_percent_sign_cannot_turn_into_a_visible_percentage(self):
        self.page.insert_text((40, 60), "19%", fontsize=20)
        self.page.draw_rect((62.5, 40, 100, 70), fill=(1, 1, 1))
        self.assertFalse(self.reader(self.page).locations("19%"))

    def test_opaque_picture_hides_text_but_a_transparent_picture_does_not(self):
        import pymupdf
        self.page.insert_text((40, 60), "Text under an opaque picture")
        opaque = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 2, 2), False)
        opaque.clear_with(255)
        self.page.insert_image((30, 40, 400, 80), pixmap=opaque, keep_proportion=False)
        self.page.insert_text((40, 140), "Text under a transparent picture")
        transparent = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 2, 2), True)
        transparent.clear_with()
        self.page.insert_image((30, 120, 400, 160), pixmap=transparent, keep_proportion=False)
        text = self.reader(self.page)
        self.assertFalse(text.locations("Text under an opaque picture"))
        self.assertTrue(text.locations("Text under a transparent picture"))


@unittest.skipUnless(HAVE, "needs python-pptx, PyMuPDF and LibreOffice")
class OnScreenOnly(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = Path(tempfile.mkdtemp())
        cls.res = P.write_outputs(F.offscreen_deck(cls.d), cls.d / "out")
        cls.sections = json.loads((cls.d / "out" / "sections.json").read_text())
        cls.stats = json.loads((cls.d / "out" / "stats.json").read_text())
        cls.batch = (cls.d / "out" / "batch_1.json").read_text()
        cls.seen = everything_seen(cls.sections)

    def test_nothing_off_the_slide_reaches_the_analysis(self):
        leaked = [t for t in F.OFFSCREEN_DROP if sq(t) in self.seen]
        self.assertEqual(leaked, [])
        self.assertFalse(any(sq(t) in sq(self.batch) for t in F.OFFSCREEN_DROP), "an off-screen text reached a batch file")

    def test_everything_on_the_slide_is_kept(self):
        # guards: a letter-spaced label, a table row, text in a group, overflow that stays on the slide
        lost = [t for t in F.OFFSCREEN_KEEP if sq(t) not in self.seen]
        self.assertEqual(lost, [])

    def test_speaker_notes_never_enter(self):
        self.assertNotIn(sq(F.OFFSCREEN_NOTES), self.seen)

    def test_left_out_text_is_reported(self):
        off = self.stats["offscreen"]
        self.assertEqual(off["blocks"], len(F.OFFSCREEN_DROP))
        self.assertEqual(off["slides"], [1])
        self.assertEqual(off["by_reason"], {"outside the slide": 2, "hidden shape": 1, "not shown on the slide": 5})

    def test_statement_counts_exclude_left_out_text(self):
        statements = sum(len(p["sentences"]) for s in self.sections for p in s["paragraphs"]) + \
            sum(len(t["rows"]) for s in self.sections for t in s["tables"])
        self.assertEqual(self.stats["statements_total"], statements)

    def test_coverage_line_names_what_was_left_out(self):
        import findings_report as R
        line = R.coverage_line({"scope": "all", "sections_total": 1, "claims_verified": 3, "offscreen": self.stats["offscreen"]})
        self.assertIn("Not checked: 8 text blocks not on the slide (slide 1).", line)
        self.assertNotIn("not on the slide", R.coverage_line({"scope": "all", "sections_total": 1, "claims_verified": 3}))

    def test_merge_carries_offscreen_into_run_json(self):
        src = (KB_ROOT / "skills" / "doc-fact-check" / "merge_findings.py").read_text()
        self.assertIn('run["offscreen"] = stats["offscreen"]', src)


@unittest.skipUnless(HAVE, "needs python-pptx, PyMuPDF and LibreOffice")
class OffscreenDocumentDate(unittest.TestCase):
    def test_a_date_parked_beside_the_title_slide_is_not_the_document_date(self):
        from pptx.util import Inches
        d = Path(tempfile.mkdtemp())
        prs = F._prs()
        t = prs.slides.add_slide(prs.slide_layouts[0])
        t.shapes.title.text = "Quarterly review"
        t.placeholders[1].text = "Operations"
        t.shapes.add_textbox(prs.slide_width + Inches(1), Inches(1), Inches(3), Inches(0.5)).text_frame.text = "March 2019"
        res = P.write_outputs(F._save(prs, d / "parked.pptx"), d / "out")
        stats = json.loads((d / "out" / "stats.json").read_text())
        self.assertNotEqual((stats.get("doc_date") or {}).get("value"), "March 2019")


@unittest.skipUnless(HAVE, "needs python-pptx, PyMuPDF and LibreOffice")
class CoveredAndClippedContent(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
        cls.d = Path(tempfile.mkdtemp())
        prs = F._prs()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        def text(value, x, y, w, h):
            sh = F._tb(slide.shapes, Inches(x), Inches(y), Inches(w), Inches(h), value)
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(18)
            return sh

        def cover(x, y, w, h):
            sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor(255, 255, 255)
            sh.line.fill.background()

        text("Visible orders use the warehouse feed.", 0.5, 0.4, 8, 0.6)
        text("Covered title invents a different scenario.", 0.5, 1.3, 8, 0.6)
        cover(0.4, 1.2, 8.2, 0.8)
        text("Most of this paragraph is visible but the final qualifier is outside the slide.", 8, 2.2, 7, 0.5)
        table = slide.shapes.add_table(2, 2, Inches(0.5), Inches(3), Inches(5), Inches(1)).table
        for r, values in enumerate([("Metric", "Daily count"), ("Orders", "9")]):
            for c, value in enumerate(values):
                table.cell(r, c).text = value
        cover(3, 3.5, 2.5, 0.5)
        second = slide.shapes.add_table(2, 2, Inches(0.5), Inches(4.5), Inches(5), Inches(1)).table
        for r, values in enumerate([("Measure", "Secret weekly period"), ("Shipments", "88")]):
            for c, value in enumerate(values):
                second.cell(r, c).text = value
        cover(3, 4.5, 2.5, 0.5)
        path = cls.d / "visibility.pptx"
        prs.save(path)
        P.write_outputs(path, cls.d / "out")
        cls.sections = json.loads((cls.d / "out/sections.json").read_text())
        cls.batch = (cls.d / "out/batch_1.json").read_text()

    def test_covered_title_and_partial_paragraph_never_reach_claim_inputs_or_context(self):
        text = json.dumps(self.sections) + self.batch
        self.assertIn("Visible orders use the warehouse feed", text)
        self.assertNotIn("Covered title invents", text)
        self.assertNotIn("final qualifier is outside", text)

    def test_hidden_value_or_header_cannot_pass_because_another_cell_is_visible(self):
        rows = [r for s in self.sections for t in s["tables"] for r in t["rows"]]
        self.assertFalse(any("Orders" in r["cells"] or "Shipments" in r["cells"] for r in rows))
        self.assertNotIn("Secret weekly period", self.batch)


if __name__ == "__main__":
    unittest.main()
