"""sections.py: deterministic heading sections and word-bounded batches from a .docx (synthetic fixture)."""
from __future__ import annotations

import hashlib
import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from test_fact_check_code import HAVE_DOCX, _tiny_png
from test_plugin_structure import KB_ROOT

SCRIPT = KB_ROOT / "skills" / "doc-fact-check" / "sections.py"
MAX = 100
LONG_PARAS = 8          # 8 x 20 words = 160 words > MAX: section "2 Roadmap" must split


def build_fixture(path: Path) -> None:
    from docx import Document
    from docx.shared import Pt
    d = Document()
    d.add_paragraph("Integration overview", style="Title")
    d.add_heading("1 Scope", level=1)
    d.add_paragraph("The platform runs 12 services in two regions.")
    d.add_paragraph("")                                         # empty paragraph: no text, no record
    d.add_paragraph("System A sends orders to system B every night.")
    d.add_heading("1.1 Systems", level=2)
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("System", "Owner"), ("Alpha", "Team one"), ("Beta", "Team two")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.add_paragraph("Figure 1 shows the flow.").add_run().add_picture(io.BytesIO(_tiny_png()), width=Pt(10))
    d.add_paragraph("Beta is retired in 2027.")
    d.add_heading("2 Roadmap", level=1)
    for i in range(LONG_PARAS):
        d.add_paragraph(" ".join([f"milestone{i}"] + ["word"] * 19))
    d.save(path)


def run_cli(docx: Path, out: Path, *extra: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(SCRIPT), str(docx), "--out", str(out), *extra],
                          capture_output=True, text=True)


@unittest.skipUnless(HAVE_DOCX, "needs python-docx >= 1.2")
class TestSections(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        cls.docx = cls.tmp / "draft.docx"
        build_fixture(cls.docx)
        cls.out = cls.tmp / "out"
        r = run_cli(cls.docx, cls.out, "--max-words", str(MAX))
        assert r.returncode == 0, r.stderr
        cls.sections = json.loads((cls.out / "sections.json").read_text())
        cls.batches = json.loads((cls.out / "batches.json").read_text())
        cls.batch_files = {b["batch"]: json.loads((cls.out / b["file"]).read_text()) for b in cls.batches}

    # -- sections.json --------------------------------------------------------------------
    def test_sections_snapshot(self):
        self.assertEqual([s["heading_path"] for s in self.sections], [
            ["Integration overview"],
            ["Integration overview", "1 Scope"],
            ["Integration overview", "1 Scope", "1.1 Systems"],
            ["Integration overview", "2 Roadmap"],
        ])
        self.assertEqual([s["section_id"] for s in self.sections], ["s01", "s02", "s03", "s04"])
        self.assertEqual(self.sections[1]["section"], "Integration overview > 1 Scope")
        scope = self.sections[1]
        self.assertEqual([p["text"] for p in scope["paragraphs"]],
                         ["The platform runs 12 services in two regions.",
                          "System A sends orders to system B every night."])
        self.assertEqual([p["p_id"] for p in scope["paragraphs"]], ["p3", "p5"])   # p4 is the empty one
        systems = self.sections[2]
        self.assertEqual(systems["tables"], [{"t_id": "t1", "rows": [
            {"r_id": "t1r1", "cells": ["System", "Owner"], "s_id": "t1r1", "risk": ["ownership"]},
            {"r_id": "t1r2", "cells": ["Alpha", "Team one"], "s_id": "t1r2", "risk": ["num"]},
            {"r_id": "t1r3", "cells": ["Beta", "Team two"], "s_id": "t1r3", "risk": ["num"]}]}])
        self.assertEqual([p["text"] for p in systems["paragraphs"]],
                         ["Figure 1 shows the flow.", "Beta is retired in 2027."])
        self.assertEqual(len(self.sections[3]["paragraphs"]), LONG_PARAS)

    def test_heading_levels_come_from_heading_styles(self):
        self.assertEqual([s["level"] for s in self.sections], [0, 1, 2, 1])
        self.assertEqual([s["heading_p_id"] for s in self.sections], ["p1", "p2", "p6", "p9"])

    def test_image_is_a_figure_with_its_paragraph_id(self):
        figs = [f for s in self.sections for f in s["figures"]]
        self.assertEqual(figs, [{"figure": 1, "p_id": "p7", "image": "word/media/image1.png",
                                "media": "media/image1.png"}])
        self.assertEqual(self.sections[2]["figures"], figs)

    def test_words_count_heading_paragraphs_and_cells(self):
        s = self.sections[3]
        self.assertEqual(s["words"], 2 + LONG_PARAS * 20)

    # -- batches.json ---------------------------------------------------------------------
    def _units(self):
        return [u for b in self.batches for u in self.batch_files[b["batch"]]["sections"]]

    def test_every_paragraph_and_row_in_exactly_one_batch(self):
        want_p = [p["p_id"] for s in self.sections for p in s["paragraphs"]]
        want_r = [r["r_id"] for s in self.sections for t in s["tables"] for r in t["rows"]]
        got_p = [p["p_id"] for u in self._units() for p in u["paragraphs"]]
        got_r = [r["r_id"] for u in self._units() for t in u["tables"] for r in t["rows"]]
        self.assertEqual(got_p, want_p)                     # same order, no duplicates, none missing
        self.assertEqual(got_r, want_r)
        self.assertEqual(len(got_p), len(set(got_p)))

    def test_small_sections_are_grouped_and_never_split(self):
        b1 = self.batch_files[1]["sections"]
        self.assertEqual([u["section_id"] for u in b1], ["s01", "s02", "s03"])
        for u in b1:
            self.assertNotIn("part", u)
        for b in self.batches:
            ids = [u["section_id"] for u in self.batch_files[b["batch"]]["sections"]]
            self.assertTrue(set(ids) <= {"s01", "s02", "s03"} or set(ids) == {"s04"})

    def test_oversized_section_is_split_at_paragraph_boundaries(self):
        parts = [u for u in self._units() if u["section_id"] == "s04"]
        self.assertGreater(len(parts), 1)
        self.assertEqual([u["part"] for u in parts], list(range(1, len(parts) + 1)))
        self.assertTrue(all(u["parts"] == len(parts) for u in parts))
        for u in parts:
            self.assertLessEqual(u["words"], MAX)
            self.assertEqual(u["section"], "Integration overview > 2 Roadmap")
        for b in self.batches:
            self.assertLessEqual(b["words"], MAX)

    def test_batches_index_matches_batch_files(self):
        self.assertEqual([b["batch"] for b in self.batches], list(range(1, len(self.batches) + 1)))
        for b in self.batches:
            f = self.batch_files[b["batch"]]
            self.assertEqual(b["file"], f"batch_{b['batch']}.json")
            self.assertEqual(f["batch"], b["batch"])
            self.assertEqual(b["sections"], [{"section_id": u["section_id"], "section": u["section"],
                                              **({"part": u["part"]} if "part" in u else {})}
                                             for u in f["sections"]])
            self.assertEqual(b["words"], sum(u["words"] for u in f["sections"]))

    def test_determinism_identical_bytes(self):
        again = self.tmp / "again"
        r = run_cli(self.docx, again, "--max-words", str(MAX))
        self.assertEqual(r.returncode, 0, r.stderr)
        names = sorted(str(p.relative_to(self.out)) for p in self.out.rglob("*") if p.is_file())
        self.assertEqual(names, sorted(str(p.relative_to(again)) for p in again.rglob("*") if p.is_file()))
        for n in names:
            self.assertEqual((self.out / n).read_bytes(), (again / n).read_bytes(), n)

    def test_default_limit_keeps_the_small_document_in_one_batch(self):
        out = self.tmp / "default"
        r = run_cli(self.docx, out)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(len(json.loads((out / "batches.json").read_text())), 1)

    def test_missing_input_fails_loudly(self):
        r = run_cli(self.tmp / "nope.docx", self.tmp / "x")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("nope.docx", r.stderr)


@unittest.skipUnless(HAVE_DOCX, "needs python-docx >= 1.2")
class TestSectionsEdgeCases(unittest.TestCase):
    def test_text_before_the_first_heading_is_its_own_section(self):
        from docx import Document
        tmp = Path(tempfile.mkdtemp())
        d = Document()
        d.add_paragraph("Draft prepared for review.")
        d.add_heading("Only heading", level=1)
        d.add_paragraph("Body.")
        d.save(tmp / "d.docx")
        r = run_cli(tmp / "d.docx", tmp / "o")
        self.assertEqual(r.returncode, 0, r.stderr)
        secs = json.loads((tmp / "o" / "sections.json").read_text())
        self.assertEqual(secs[0]["heading_path"], ["(before first heading)"])
        self.assertIsNone(secs[0]["heading_p_id"])
        self.assertEqual(secs[1]["heading_path"], ["Only heading"])


import sections as S


class RiskTagTests(unittest.TestCase):
    def tags(self, t):
        return S.tag_risk(t)

    def test_num_positive(self):
        for t in ("about ~2.8M customers", ">1 million API calls per day", "retained for 60 days",
                  "runs on four schedules", "45% of traffic"):
            self.assertIn("num", self.tags(t), t)

    def test_num_ignores_numbering(self):
        for t in ("see Section 6.12 for detail", "as shown in \u00a74.8", "Figure 2 shows the flow",
                  "Phase 1 covers discovery", "release v1.3 notes"):
            self.assertNotIn("num", self.tags(t), t)

    def test_date_positive(self):
        for t in ("go-live in December 2026", "planned for Q3", "as of 12/31/2025", "currently in pilot"):
            self.assertIn("date", self.tags(t), t)

    def test_date_negative(self):
        for t in ("the team may revisit this", "version two of the design"):
            self.assertNotIn("date", self.tags(t), t)

    def test_absolute_positive(self):
        for t in ("has never attended a session", "consistent across all four databases",
                  "the sole developer of the app", "only one integration exists"):
            self.assertIn("absolute", self.tags(t), t)

    def test_absolute_negative(self):
        for t in ("the team reviews the backlog", "integration is planned"):
            self.assertNotIn("absolute", self.tags(t), t)

    def test_ownership_positive(self):
        for t in ("the vendor owns the platform", "operated by the central team", "she attended the review",
                  "the board approved the budget", "managed by the partner"):
            self.assertIn("ownership", self.tags(t), t)

    def test_ownership_negative(self):
        for t in ("the system sends a message", "data flows nightly"):
            self.assertNotIn("ownership", self.tags(t), t)

    def test_plain_statement_has_no_tags(self):
        self.assertEqual(self.tags("The portal displays the order history."), [])

    def test_split_sentences(self):
        self.assertEqual(S.split_sentences("First claim. Second one; third\nfourth"),
                         ["First claim.", "Second one; third", "fourth"])

    def test_estimate_minutes(self):
        self.assertEqual(S.estimate_minutes(0, 0), {"fast": 5, "deep": 5})
        self.assertEqual(S.estimate_minutes(80, 16), {"fast": 15, "deep": 15})
        self.assertEqual(S.estimate_minutes(800, 160), {"fast": 25, "deep": 105})


@unittest.skipUnless(HAVE_DOCX, "needs python-docx >= 1.2")
class StatsOutputTests(unittest.TestCase):
    def test_batch_records_carry_sentences_and_risk_and_stats_written(self):
        from docx import Document
        d = Path(tempfile.mkdtemp())
        doc = Document()
        doc.add_heading("Scope", 1)
        doc.add_paragraph("The platform runs 12 services. The portal shows history.")
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Owner", "managed by the partner"
        src = d / "in.docx"
        doc.save(str(src))
        S.write_outputs(src, d / "out")
        batch = json.loads((d / "out" / "batch_1.json").read_text())
        sec = [s for s in batch["sections"] if s["section"] == "Scope"][0]
        sents = sec["paragraphs"][0]["sentences"]
        self.assertEqual([s["s_id"] for s in sents], ["p2s1", "p2s2"])
        self.assertEqual(sents[0]["risk"], ["num"])
        self.assertEqual(sents[1]["risk"], [])
        row = sec["tables"][0]["rows"][0]
        self.assertEqual(row["s_id"], row["r_id"])
        self.assertIn("ownership", row["risk"])
        stats = json.loads((d / "out" / "stats.json").read_text())
        self.assertEqual(stats["statements_total"], 3)
        self.assertEqual(stats["statements_risk"], 2)
        self.assertEqual(stats["estimate_minutes"], S.estimate_minutes(3, 2))
        self.assertEqual(stats["source_sha256"], hashlib.sha256(Path(src).read_bytes()).hexdigest())



@unittest.skipUnless(HAVE_DOCX, "needs python-docx >= 1.2")
class MediaAndSummaryTests(unittest.TestCase):
    def _doc_with_figure(self, d):
        from docx import Document
        png = d / "fig.png"; png.write_bytes(_tiny_png())
        doc = Document(); doc.add_heading("Scope", 1)
        doc.add_paragraph("Figure 1 shows the flow."); doc.add_picture(str(png))
        src = d / "in.docx"; doc.save(str(src))
        return src

    def test_referenced_media_is_extracted_and_linked_from_the_figure(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        S.write_outputs(src, d / "out")
        secs = json.loads((d / "out" / "sections.json").read_text())
        fig = [f for s in secs for f in s["figures"]][0]
        self.assertTrue(fig["media"].startswith("media/"))
        self.assertEqual((d / "out" / fig["media"]).read_bytes(), _tiny_png())

    def test_orphaned_package_media_is_not_extracted(self):
        import zipfile
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        orphan = d / "orphan.docx"
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(orphan, "w") as zout:
            for i in zin.infolist():
                zout.writestr(i, zin.read(i.filename))
            zout.writestr("word/media/stray.png", _tiny_png())
        S.write_outputs(orphan, d / "out")
        secs = json.loads((d / "out" / "sections.json").read_text())
        fig = [f for s in secs for f in s["figures"]][0]
        files = [p.name for p in (d / "out" / "media").iterdir()]
        self.assertEqual(files, [Path(fig["media"]).name])
        self.assertNotIn("stray.png", files)

    def test_cli_prints_the_summary_the_agent_needs(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        r = subprocess.run([sys.executable, str(SCRIPT), str(src), "--out", str(d / "out")],
                           capture_output=True, text=True)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertRegex(r.stdout, r"statements \d+, high-risk \d+, estimate fast ~\d+ min, deep ~\d+ min")
        self.assertRegex(r.stdout, r"figures: 1 \(figure 1 in p\d+ -> media/image1\.png\)")
        self.assertIn("embedded objects: 0", r.stdout)

    def test_one_image_used_by_two_figures_is_written_once_and_linked_from_both(self):
        from docx import Document
        d = Path(tempfile.mkdtemp()); png = d / "fig.png"; png.write_bytes(_tiny_png())
        doc = Document(); doc.add_paragraph("A"); doc.add_picture(str(png)); doc.add_paragraph("B"); doc.add_picture(str(png))
        src = d / "in.docx"; doc.save(str(src))
        S.write_outputs(src, d / "out")
        figs = [f for s in json.loads((d / "out" / "sections.json").read_text()) for f in s["figures"]]
        self.assertEqual(len(figs), 2)
        self.assertEqual({f["media"] for f in figs}, {figs[0]["media"]})
        self.assertEqual(len(list((d / "out" / "media").iterdir())), 1)

    def test_external_linked_image_is_skipped_not_fatal(self):
        d = Path(tempfile.mkdtemp())
        secs = [{"blocks": [("f", {"figure": 1, "p_id": "p2", "image": "http://example.invalid/x.png"}, 0)]}]
        src = self._doc_with_figure(d)
        self.assertEqual(S.extract_media(src, d / "out", secs), [])
        self.assertNotIn("media", secs[0]["blocks"][0][1])

    def _with_members(self, src, d, members):
        import zipfile
        out = d / "withobj.docx"
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w") as zout:
            for i in zin.infolist():
                zout.writestr(i, zin.read(i.filename))
            for name, data in members.items():
                zout.writestr(name, data)
        return out

    def _xlsx_bytes(self, extra=None):
        import zipfile
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("[Content_Types].xml", "<Types/>")
            z.writestr("xl/sharedStrings.xml", "<sst><si><t>Total 12</t></si></sst>")
            z.writestr("xl/worksheets/sheet1.xml", "<worksheet/>")
            z.writestr("xl/media/image1.png", _tiny_png())          # not XML: not expanded
            for name, data in (extra or {}).items():
                z.writestr(name, data)
        return buf.getvalue()

    def _run_cli(self, docx, out):
        return subprocess.run([sys.executable, str(SCRIPT), str(docx), "--out", str(out)],
                              capture_output=True, text=True)

    def test_embedded_ooxml_package_is_expanded_into_its_xml_parts(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        docx = self._with_members(src, d, {"word/embeddings/Microsoft_Excel_Worksheet1.xlsx": self._xlsx_bytes()})
        r = self._run_cli(docx, d / "out")
        self.assertEqual(r.returncode, 0, r.stderr)
        obj = d / "out" / "embedded" / "Microsoft_Excel_Worksheet1"
        self.assertIn(b"Total 12", (obj / "xl" / "sharedStrings.xml").read_bytes())
        self.assertTrue((obj / "xl" / "worksheets" / "sheet1.xml").is_file())
        self.assertTrue((obj / "[Content_Types].xml").is_file())
        self.assertFalse((obj / "xl" / "media").exists())
        self.assertIn("embedded objects: 1 (embedded/Microsoft_Excel_Worksheet1/: 3 XML parts)", r.stdout)

    def test_expanded_parts_never_escape_the_out_dir(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        evil = self._xlsx_bytes({"../../escaped.xml": "<x/>", "/abs.xml": "<x/>", "xl/../../up.xml": "<x/>"})
        docx = self._with_members(src, d, {"word/embeddings/Book1.xlsx": evil})
        r = self._run_cli(docx, d / "out")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertFalse((d / "escaped.xml").exists())
        self.assertFalse((d / "out" / "escaped.xml").exists())
        self.assertFalse((d / "out" / "embedded" / "up.xml").exists())
        found = sorted(p.relative_to(d / "out" / "embedded" / "Book1").as_posix()
                       for p in (d / "out" / "embedded" / "Book1").rglob("*.xml"))
        self.assertEqual(found, ["[Content_Types].xml", "abs.xml", "xl/sharedStrings.xml", "xl/worksheets/sheet1.xml"])

    def test_ole_binary_is_copied_and_reported_not_readable(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        docx = self._with_members(src, d, {"word/embeddings/oleObject1.bin": b"OLE",
                                           "word/diagrams/data1.xml": b"<x/>"})
        r = self._run_cli(docx, d / "out")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual((d / "out" / "embedded" / "oleObject1" / "oleObject1.bin").read_bytes(), b"OLE")
        self.assertEqual((d / "out" / "embedded" / "diagrams" / "data1.xml").read_bytes(), b"<x/>")
        self.assertIn("embedded objects: 1 (embedded/oleObject1/oleObject1.bin: not readable), "
                      "diagrams: 1 (embedded/diagrams/data1.xml)", r.stdout)

    def test_chart_parts_are_extracted_without_their_rels(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        docx = self._with_members(src, d, {"word/charts/chart1.xml": b"<c:chartSpace/>",
                                           "word/charts/_rels/chart1.xml.rels": b"<Relationships/>",
                                           "word/charts/embeddings/Microsoft_Excel_Worksheet.xlsx": self._xlsx_bytes()})
        r = self._run_cli(docx, d / "out")
        self.assertEqual(r.returncode, 0, r.stderr)
        charts = d / "out" / "embedded" / "charts"
        self.assertEqual((charts / "chart1.xml").read_bytes(), b"<c:chartSpace/>")
        self.assertEqual(sorted(p.name for p in charts.rglob("*")), ["chart1.xml"])
        self.assertIn("charts: 1 (embedded/charts/chart1.xml)", r.stdout)

    def test_two_embeddings_with_one_stem_get_separate_folders(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        docx = self._with_members(src, d, {"word/embeddings/obj.bin": b"OLE", "word/embeddings/obj.xlsx": self._xlsx_bytes()})
        r = self._run_cli(docx, d / "out")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual((d / "out" / "embedded" / "obj" / "obj.bin").read_bytes(), b"OLE")
        self.assertTrue((d / "out" / "embedded" / "obj-xlsx" / "xl" / "sharedStrings.xml").is_file())

    def _doc_with_chart_figure(self, d, smartart=False):
        """A drawing whose graphicData is a c:chart pointing at word/charts/chart1.xml (no image), or with
        smartart=True a dgm:relIds whose r:dm points at word/diagrams/data1.xml."""
        from docx import Document
        from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part
        from docx.oxml.ns import qn
        from lxml import etree
        png = d / "fig.png"; png.write_bytes(_tiny_png())
        doc = Document(); doc.add_heading("Scope", 1)
        doc.add_paragraph("Figure 1 shows the volumes."); doc.add_picture(str(png))
        name, ct, rt, tag, attr = (("/word/diagrams/data1.xml", CT.DML_DIAGRAM_DATA, RT.DIAGRAM_DATA, "dgm:relIds", "r:dm")
                                   if smartart else ("/word/charts/chart1.xml", CT.DML_CHART, RT.CHART, "c:chart", "r:id"))
        rid = doc.part.relate_to(Part(PackURI(name), ct, b"<x/>", doc.part.package), rt)
        gd = next(doc.element.body.iter(qn("a:graphicData")))
        for child in list(gd):
            gd.remove(child)
        etree.SubElement(gd, qn(tag)).set(qn(attr), rid)
        src = d / "chart.docx"; doc.save(str(src))
        return src

    def test_smartart_figure_names_its_diagram_data_part(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_chart_figure(d, smartart=True)
        r = self._run_cli(src, d / "out")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertRegex(r.stdout, r"figures: 1 \(figure 1 in p\d+ -> embedded/diagrams/data1\.xml\)")

    def test_chart_figure_names_the_part_it_points_to(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_chart_figure(d)
        r = self._run_cli(src, d / "out")
        self.assertEqual(r.returncode, 0, r.stderr)
        fig = [f for s in json.loads((d / "out" / "sections.json").read_text()) for f in s["figures"]][0]
        self.assertIsNone(fig["image"])
        self.assertEqual(fig["part"], "embedded/charts/chart1.xml")
        self.assertTrue((d / "out" / fig["part"]).is_file())
        self.assertRegex(r.stdout, r"figures: 1 \(figure 1 in p\d+ -> embedded/charts/chart1\.xml\)")
        self.assertNotIn("not extracted", r.stdout)

    def test_rerun_into_the_same_out_dir_drops_stale_media_and_embedded(self):
        from docx import Document
        d = Path(tempfile.mkdtemp())
        src = self._with_members(self._doc_with_figure(d), d, {"word/embeddings/oleObject1.bin": b"OLE"})
        S.write_outputs(src, d / "out")
        self.assertTrue(any((d / "out" / "media").iterdir()))
        plain = d / "plain.docx"; doc = Document(); doc.add_paragraph("No figure."); doc.save(str(plain))
        S.write_outputs(plain, d / "out")
        self.assertEqual(list((d / "out").glob("media/*")), [])
        self.assertEqual(list((d / "out").glob("embedded/*")), [])

    def test_refuses_to_clear_media_in_a_folder_it_did_not_create(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        for sub in ("media", "embedded"):
            with self.subTest(sub=sub):
                out = d / f"foreign_{sub}"
                (out / sub).mkdir(parents=True)
                keep = out / sub / "precious.txt"; keep.write_text("keep")
                r = self._run_cli(src, out)
                self.assertEqual(r.returncode, 2, r.stderr)
                self.assertIn(f"refusing to clear {sub}/ in a folder sections.py did not create", r.stderr)
                self.assertEqual(keep.read_text(), "keep")
                self.assertFalse((out / "sections.json").exists())

    def test_refuses_to_clear_a_symlinked_media_dir(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        out = d / "out"
        S.write_outputs(src, out)
        elsewhere = d / "elsewhere"; elsewhere.mkdir()
        keep = elsewhere / "precious.txt"; keep.write_text("keep")
        import shutil
        shutil.rmtree(out / "media")
        (out / "media").symlink_to(elsewhere, target_is_directory=True)
        r = self._run_cli(src, out)
        self.assertEqual(r.returncode, 2, r.stderr)
        self.assertIn("refusing to clear media/", r.stderr)
        self.assertEqual(keep.read_text(), "keep")

    def test_batch_files_carry_the_media_link(self):
        d = Path(tempfile.mkdtemp()); src = self._doc_with_figure(d)
        S.write_outputs(src, d / "out")
        figs = [f for b in (d / "out").glob("batch_*.json")
                for sec in json.loads(b.read_text())["sections"] for f in sec["figures"]]
        self.assertEqual(len(figs), 1)
        self.assertTrue(figs[0]["media"].startswith("media/"))

    def test_skill_step_1_no_longer_says_unzip(self):
        skill = (KB_ROOT / "skills" / "doc-fact-check" / "SKILL.md").read_text(encoding="utf-8")
        self.assertNotIn("unzip `word/media/*`", skill)
        self.assertIn("`<work dir>/media/`", skill)

if __name__ == "__main__":
    unittest.main()
