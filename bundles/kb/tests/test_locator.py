"""A deck comment says which statement it is about (spec F1, evals 1 and 2): it quotes the statement, its header
names the slide only, and it is pinned where the statement renders, not at its shape's corner. Word comments are
unchanged (they are anchored on the exact words)."""
from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from test_plugin_structure import KB_ROOT

import sys
DECK = KB_ROOT / "skills" / "deck-fact-check"
sys.path.insert(0, str(DECK))

import fact_check_invariants as I  # noqa: E402

try:
    import pptx  # noqa: F401
    import pymupdf  # noqa: F401
    import deck_fixtures as F
    import pptx_annotate as A
    import pptx_sections as P
    from render_pages import soffice_bin
    HAVE = soffice_bin() is not None
except ImportError:
    HAVE = False

BASE = {"verdict": "Incorrect", "severity": "Major", "evidence": "The source states 99.1%.", "fix": "Use 99.1%.",
        "why": "the source states 99.1%", "suggested": "Uptime was 99.1%.", "source": "ops.md"}
LONG = "Uptime held at 99.9% all year across every region including the northern depots and the two new hubs"


class CommentText(unittest.TestCase):
    def test_deck_comment_quotes_the_statement_and_explains_the_correction(self):
        f = {**BASE, "id": "C07", "p_id": "s12-sh8-p3", "section": "Slide 12: Traffic mix", "quote": LONG}
        text = I.comment_text(f)
        self.assertTrue(text.startswith('[Incorrect]\n"Uptime held at 99.9% all year across every region including the northern…"\n'))
        self.assertIn("Why: Ops (undated) says the source states 99.1%.", text)
        self.assertIn('Suggested: "Uptime was 99.1%."', text)
        self.assertNotIn("C07", text)
        full = I.comment_text(f, I.COMMENT_FIELDS)
        self.assertIn("Brain: The source states 99.1%.", full)
        self.assertIn("Source: ops.md", full)

    def test_short_quote_is_not_cut(self):
        f = {**BASE, "id": "C01", "p_id": "s2-sh3-r2", "section": "Slide 2: North", "quote": "Uptime | 99.9%"}
        self.assertIn('\n"Uptime | 99.9%"\n', I.comment_text(f))

    def test_deck_figure_comment_quotes_the_visual_statement(self):
        f = {**BASE, "id": "I01", "p_id": "s3", "section": "Figure", "quote": "bar Southfield = 97.4%"}
        self.assertTrue(I.comment_text(f).startswith('[Incorrect]\n"bar Southfield = 97.4%"\n'))

    def test_word_comment_explains_the_correction_without_repeating_its_anchor(self):
        f = {**BASE, "id": "C10", "p_id": "p14", "section": "4 Integration > 4.8 Billing", "quote": LONG}
        self.assertEqual(I.comment_text(f), '[Incorrect]\nWhy: Ops (undated) says the source states 99.1%.\nSuggested: "Uptime was 99.1%."')
        self.assertNotIn(LONG, I.comment_text(f))

    def test_word_cap_excludes_category_and_deck_statement(self):
        deck = {**BASE, "id": "C07", "p_id": "s12-sh8-p3", "section": "Slide 12", "quote": LONG}
        body = 'Why: Ops (undated) says the source states 99.1%. Suggested: "Uptime was 99.1%."'
        self.assertEqual(I.comment_words(I.comment_text(deck), True), len(body.split()))
        word = {**deck, "p_id": "p3"}
        self.assertEqual(I.comment_words(I.comment_text(word), False), len(body.split()))


@unittest.skipUnless(HAVE, "needs python-pptx, PyMuPDF and LibreOffice")
class StatementPin(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = Path(tempfile.mkdtemp())
        cls.deck = F.pins_deck(cls.d)
        P.write_outputs(cls.deck, cls.d / "work")
        cls.sections = json.loads((cls.d / "work" / "sections.json").read_text())
        cls.paras = cls.sections[0]["paragraphs"]

    def test_reader_records_where_each_statement_renders(self):
        ys = [p["at"]["y"] for p in self.paras]
        self.assertEqual(ys, sorted(ys))
        self.assertEqual(len(set(ys)), 6)
        x, y, w, h = F.PINS_BOX
        line = F.PINS_PT * 1.2 * 12700
        for k, p in enumerate(self.paras):
            self.assertTrue(x <= p["at"]["x"] <= x + w and y <= p["at"]["y"] <= y + h)
            self.assertLessEqual(abs(p["at"]["y"] - (y + 45720 + k * line)) / line, 0.5, p["p_id"])

    def _findings(self, pids):
        return [{**BASE, "id": f"C{i:02d}", "p_id": pid, "section": "Slide 1", "quote": t}
                for i, (pid, t) in enumerate(pids, 1)]

    def test_two_statements_in_one_box_get_two_pins_at_their_lines(self):
        fs = self._findings([(p["p_id"], p["text"]) for p in self.paras[:2]])
        out = self.d / "two.pptx"
        written, _ = A.annotate(self.deck, out, fs, sections=self.d / "work" / "sections.json")
        pins = {I.comment_owner(c["text"], fs): c for c in I.read_pptx_comments(out)}
        self.assertEqual(len(written), 2)
        self.assertNotEqual(pins["C01"]["y"], pins["C02"]["y"])
        for f in fs:
            c = pins[f["id"]]
            at = next(p["at"] for p in self.paras if p["p_id"] == f["p_id"])
            self.assertEqual((c["x"], c["y"]), (A.to_pos(at["x"]), A.to_pos(at["y"])))
            self.assertIn(f'"{" ".join(f["quote"].split()[:12])}', c["text"])
        view = I.doc_view(out)
        by = {I.comment_owner(c.text, fs): c for c in view.comments()}
        self.assertEqual([view.anchor_problem(by[f["id"]], f) for f in fs], [None, None])

    def test_without_sections_the_pin_falls_back_to_the_shape_corner(self):
        fs = self._findings([(self.paras[3]["p_id"], self.paras[3]["text"])])
        out = self.d / "corner.pptx"
        A.annotate(self.deck, out, fs)
        c = I.read_pptx_comments(out)[0]
        x, y, *_ = F.PINS_BOX
        self.assertEqual((c["x"], c["y"]), (A.to_pos(x), A.to_pos(y)))

    def test_a_short_statement_is_located_inside_its_own_box(self):
        from pptx.util import Inches
        d = Path(tempfile.mkdtemp())
        prs = F._prs()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for k, y in ((1, 1), (2, 2), (3, 3)):                 # agenda numbers: "01", "02", "03" in their own boxes
            F._tb(s.shapes, Inches(1), Inches(y), Inches(1), Inches(0.5), f"0{k}")
        P.write_outputs(F._save(prs, d / "short.pptx"), d / "work")
        paras = json.loads((d / "work" / "sections.json").read_text())[0]["paragraphs"]
        self.assertEqual([("at" in p) for p in paras], [True, True, True])
        for p in paras:
            self.assertTrue(p["pos"]["y"] <= p["at"]["y"] <= p["pos"]["y"] + p["pos"]["h"], p["p_id"])

    def test_table_rows_are_located_too(self):
        d = Path(tempfile.mkdtemp())
        P.write_outputs(F.table_deck(d), d / "work")
        rows = json.loads((d / "work" / "sections.json").read_text())[0]["tables"][0]["rows"]
        ys = [r["at"]["y"] for r in rows if "at" in r]
        self.assertEqual(len(ys), 3)
        self.assertEqual(ys, sorted(ys))

    def test_cli_takes_sections(self):
        fj = self.d / "f.json"
        fj.write_text(json.dumps(self._findings([(self.paras[4]["p_id"], self.paras[4]["text"])])))
        import subprocess
        r = subprocess.run([sys.executable, "-I", str(DECK / "pptx_annotate.py"), str(self.deck), str(self.d / "cli.pptx"),
                            "--findings", str(fj), "--select", "major", "--sections", str(self.d / "work" / "sections.json")],
                           capture_output=True, text=True, cwd="/")
        self.assertEqual(r.returncode, 0, r.stderr)
        c = I.read_pptx_comments(self.d / "cli.pptx")[0]
        self.assertEqual(c["y"], A.to_pos(self.paras[4]["at"]["y"]))

    def test_word_annotate_refuses_sections(self):
        import subprocess
        doc = KB_ROOT / "skills" / "doc-fact-check" / "annotate.py"
        fj = self.d / "w.json"
        fj.write_text("[]")
        r = subprocess.run([sys.executable, "-I", str(doc), str(self.deck), str(self.d / "x.docx"), "--findings", str(fj),
                            "--select", "major", "--sections", str(fj)], capture_output=True, text=True, cwd="/")
        self.assertEqual(r.returncode, 2)
        self.assertIn("--sections is only used by deck-fact-check", r.stderr)


if __name__ == "__main__":
    unittest.main()
