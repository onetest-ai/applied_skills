"""Claim context (spec §4.5, R10): every claim carries a self-contained `claim` built only from the draft, the
verbatim `quote` stays the anchor, and invented context is refused before stage V."""
from __future__ import annotations

import json
import re
import tempfile
import unittest
from pathlib import Path

from test_fact_check_code import HAVE_DOCX
from test_plugin_structure import KB_ROOT

import sys
sys.path.insert(0, str(KB_ROOT / "skills" / "doc-fact-check"))
import chunk_claims as C  # noqa: E402
import findings_report as R  # noqa: E402

D = KB_ROOT / "skills" / "doc-fact-check"

ROW = {"r_id": "t1r2", "s_id": "t1r2", "cells": ["Uptime", "99.9%"], "risk": ["num"],
       "ctx": {"header": ["Metric", "Value"], "row_label": "Uptime"}}
BATCH = {"batch": 1, "doc_date": {"value": "2026-03-02", "inferred": True}, "sections": [
    {"section_id": "s01", "section": "North region — FY2025", "heading_path": ["North region — FY2025"],
     "paragraphs": [{"p_id": "p2", "text": "Go-live is planned.",
                     "sentences": [{"s_id": "p2s1", "text": "Go-live is planned.", "risk": ["date"]}]}],
     "tables": [{"t_id": "t1", "rows": [ROW]}]}]}


def claim(**kw):
    base = {"claim_id": "B1-C01", "p_id": "t1r2", "s_id": "t1r2", "section": "North region — FY2025",
            "quote": "99.9%", "claim": "Uptime in the North region was 99.9% in FY2025.", "type": "NUM",
            "questions": ["what was uptime in the north region in fy2025"]}
    return {**base, **kw}


class ClaimProblem(unittest.TestCase):
    def setUp(self):
        self.known = C.known_text_by_s_id(BATCH)

    def test_context_from_heading_and_row_makes_a_valid_claim(self):
        self.assertIsNone(C.claim_problem(claim(), self.known["t1r2"]))

    def test_missing_or_blank_claim_is_refused(self):
        for bad in (None, "", "   "):
            with self.subTest(claim=bad):
                c = claim()
                c.pop("claim") if bad is None else c.update(claim=bad)
                self.assertIn("self-contained", C.claim_problem(c, self.known["t1r2"]))

    def test_over_40_words_is_refused(self):
        why = C.claim_problem(claim(claim=" ".join(["word"] * 41)), self.known["t1r2"])
        self.assertIn("41 words", why)

    def test_unresolved_must_name_known_dimensions(self):
        self.assertIsNone(C.claim_problem(claim(unresolved=["period", "scope"]), self.known["t1r2"]))
        self.assertIn("unresolved", C.claim_problem(claim(unresolved=["vibe"]), self.known["t1r2"]))
        self.assertIn("unresolved", C.claim_problem(claim(unresolved="period"), self.known["t1r2"]))

    def test_invented_year_is_refused_and_named(self):
        why = C.claim_problem(claim(claim="Uptime in the North region was 99.9% in FY2024."), self.known["t1r2"])
        self.assertIn("2024", why)
        self.assertIn("unresolved", why)

    def test_the_same_year_is_accepted_when_the_context_holds_it(self):
        known = C.known_text_by_s_id({**BATCH, "sections": [{**BATCH["sections"][0], "section": "North — FY2024",
                                                             "heading_path": ["North — FY2024"]}]})
        self.assertIsNone(C.claim_problem(claim(claim="Uptime in the North region was 99.9% in FY2024."), known["t1r2"]))

    def test_doc_date_is_context(self):
        c = claim(s_id="p2s1", quote="Go-live is planned.", claim="As of 2026-03-02, go-live is planned.")
        self.assertIsNone(C.claim_problem(c, self.known["p2s1"]))

    def test_thousands_separators_do_not_count_as_invention(self):
        known = C.known_text_by_s_id({"sections": [{"section": "Users", "paragraphs": [{"sentences": [
            {"s_id": "p1s1", "text": "We serve 1,200 users."}]}]}]})
        self.assertIsNone(C.claim_problem(claim(s_id="p1s1", quote="1,200 users", claim="The service has 1200 users."),
                                          known["p1s1"]))

    def test_figure_claims_need_a_claim_but_skip_the_number_check(self):
        self.assertIsNone(C.claim_problem({"claim": "The chart shows 97.4% uptime for Southfield."}))
        self.assertIsNotNone(C.claim_problem({"quote": "edge A to B"}))


class ChunkRefusesBeforeStageV(unittest.TestCase):
    def setUp(self):
        self.d = Path(tempfile.mkdtemp())
        (self.d / "batches.json").write_text(json.dumps([{"batch": 1, "file": "batch_1.json", "words": 9,
                                                          "sections": [{"section_id": "s01", "section": "North region — FY2025"}]}]))
        (self.d / "batch_1.json").write_text(json.dumps(BATCH))

    def run_chunk(self, claims):
        (self.d / "claims_batch_1.json").write_text(json.dumps(claims))
        return C.chunk(self.d, "all")

    def test_valid_claim_travels_into_the_chunk_with_claim_and_unresolved(self):
        out, errors = self.run_chunk([claim(unresolved=["scope"])])
        self.assertEqual(errors, [])
        got = json.loads((self.d / "chunk_1.json").read_text())["claims"][0]
        self.assertEqual((got["claim"], got["unresolved"]), (claim()["claim"], ["scope"]))

    def test_invented_number_fails_the_chunk_naming_the_claim(self):
        out, errors = self.run_chunk([claim(claim="Uptime in the North region was 99.9% across 14 sites.")])
        self.assertEqual(out, [])
        self.assertTrue(any("B1-C01" in e and "14" in e for e in errors), errors)
        self.assertFalse((self.d / "chunk_1.json").exists())


class CheckedAs(unittest.TestCase):
    """A9: the reviewer sees how a fragment was read."""

    def page(self, **kw):
        f = {"id": "C01", "p_id": "t1r2", "section": "North region — FY2025", "quote": "99.9%",
             "claim": "Uptime in the North region was 99.9% in FY2025.", "type": "NUM", "verdict": "Incorrect",
             "severity": "Major", "confidence": "High", "evidence": "The source states 99.1%.", "trail": "t",
             "conflicts": [], "fix": "Use 99.1%.", "source": "ops.md", "sources": [], "destination": "Word comment"}
        return R.render_html([{**f, **kw}], document="d.docx")

    def test_ledger_and_majors_panel_show_checked_as(self):
        html = self.page()
        self.assertEqual(html.count("Checked as: Uptime in the North region was 99.9% in FY2025."), 2)

    def test_no_checked_as_when_the_claim_is_the_quote_or_absent(self):
        self.assertNotIn("Checked as:", self.page(claim="99.9%"))
        self.assertNotIn("Checked as:", self.page(claim=None))


class SchemaAndBriefs(unittest.TestCase):
    def test_claim_is_a_finding_key_copied_from_the_claim(self):
        from fact_check_invariants import COPIED_FROM_CLAIM, FINDING_KEYS, CLAIM_ONLY_KEYS
        self.assertEqual(FINDING_KEYS[FINDING_KEYS.index("quote") + 1], "claim")
        self.assertIn("claim", COPIED_FROM_CLAIM)
        self.assertIn("unresolved", CLAIM_ONLY_KEYS)

    def test_briefs_are_format_neutral(self):
        for name in ("extract.md", "verify.md"):
            t = (D / name).read_text(encoding="utf-8")
            with self.subTest(brief=name):
                self.assertNotIn("doc-fact-check", t)
                self.assertNotRegex(t, r"\bWord\b")
                self.assertNotIn(".docx", t)

    def test_extract_brief_carries_the_claim_rules(self):
        t = re.sub(r"\s+", " ", (D / "extract.md").read_text(encoding="utf-8"))
        for ph in ("`claim` is what stage V checks", "using **only** the quote", "list it in `unresolved`",
                   "`chunk_claims.py` refuses a `claim` holding a number", "written from its `claim`"):
            self.assertIn(ph, t)
        out = t[t.index("## Output"):]
        self.assertIn("quote, claim, type, questions", out)

    def test_verify_brief_checks_the_claim_and_caps_unresolved(self):
        t = re.sub(r"\s+", " ", (D / "verify.md").read_text(encoding="utf-8"))
        for ph in ("**Check the `claim`, quote the `quote`.**", "the claim is never Verified",
                   "the `fix` names the missing context"):
            self.assertIn(ph, t)


@unittest.skipUnless(HAVE_DOCX, "needs python-docx >= 1.2")
class DocxContext(unittest.TestCase):
    def build(self, path: Path):
        import datetime
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc = Document()
        doc.add_heading("Plan", level=1)

        def item(text, level):
            p = doc.add_paragraph(text)
            num = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), str(level))
            nid = OxmlElement("w:numId"); nid.set(qn("w:val"), "1")
            num.append(ilvl); num.append(nid)
            p._p.get_or_add_pPr().append(num)
        item("Northwind migration", 0)
        item("Cut-over in Q2 2027", 1)
        item("Southfield migration", 0)
        item("Cut-over in Q4 2027", 1)
        doc.core_properties.modified = datetime.datetime(2026, 3, 2, 9, 0)
        doc.save(str(path))

    def test_nested_list_item_gets_its_parent_and_batches_carry_the_doc_date(self):
        import sections
        d = Path(tempfile.mkdtemp())
        self.build(d / "plan.docx")
        sections.write_outputs(d / "plan.docx", d / "out")
        paras = json.loads((d / "out" / "sections.json").read_text())[0]["paragraphs"]
        ctx = {p["text"]: p.get("ctx") for p in paras}
        self.assertEqual(ctx["Cut-over in Q2 2027"], {"parent": "Northwind migration"})
        self.assertEqual(ctx["Cut-over in Q4 2027"], {"parent": "Southfield migration"})
        self.assertIsNone(ctx["Northwind migration"])
        batch = json.loads((d / "out" / "batch_1.json").read_text())
        self.assertEqual(batch["doc_date"], {"value": "2026-03-02", "inferred": True})


if __name__ == "__main__":
    unittest.main()
