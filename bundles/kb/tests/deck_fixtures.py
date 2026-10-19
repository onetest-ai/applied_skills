"""Generated .pptx fixtures for deck-fact-check tests (no binary is committed; names and figures are generic).

Each builder writes a deck into a directory and returns its path; the expectations the tests need are module
constants beside the builders, so the input and what it should produce are read together."""
from __future__ import annotations

import datetime
import struct
import zlib
from pathlib import Path

EMU_PER_INCH = 914400
CHART_VALUES = ("99.1", "97.4", "95.2")
FOOTNOTE = "Source: ops report, Mar 2026"


def _png(path: Path, w: int = 64, h: int = 48) -> Path:
    """A tiny solid PNG, written with the stdlib only."""
    raw = b"".join(b"\x00" + b"\x80\x90\xa0" * w for _ in range(h))
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    return path


def _prs():
    from pptx import Presentation
    return Presentation()          # 10 x 7.5 in, default layouts: 0 title, 1 title+content, 5 title only, 6 blank


def _bullets(slide, items):
    """items: [(text, level)] into the body placeholder of a title+content slide."""
    body = slide.placeholders[1].text_frame
    body.text = items[0][0]
    body.paragraphs[0].level = items[0][1]
    for text, level in items[1:]:
        p = body.add_paragraph()
        p.text, p.level = text, level
    return slide.placeholders[1]


def basic_deck(d: Path) -> Path:
    """Slide 1: title + body with a two-level bullet list."""
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Delivery plan"
    _bullets(s, [("Northwind migration", 0), ("Go-live is set for Q2 2027.", 1),
                 ("Uptime held at 99.9% all year.", 0)])
    return _save(prs, d / "basic.pptx")


def table_deck(d: Path) -> Path:
    from pptx.util import Inches
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "North region — FY2025"
    tbl = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for r, row in enumerate([("Metric", "Value"), ("Uptime", "99.9%"), ("Owner", "Integration team")]):
        for c, text in enumerate(row):
            tbl.cell(r, c).text = text
    return _save(prs, d / "table.pptx")


GROUP_SHIFT_IN = 2        # the outer group is moved 2 in right after its children were placed


def group_deck(d: Path) -> Path:
    """A group inside a group; the outer group is then moved, so a child's slide position needs the group transform."""
    from pptx.util import Inches
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    outer = s.shapes.add_group_shape()
    a = outer.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    a.text_frame.text = "Northwind sends orders to the warehouse."
    inner = outer.shapes.add_group_shape()
    b = inner.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1))
    b.text_frame.text = "Southfield owns the billing feed."
    outer.left = outer.left + Inches(GROUP_SHIFT_IN)
    return _save(prs, d / "group.pptx")


def image_deck(d: Path) -> Path:
    from pptx.util import Inches
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(_png(d / "pic.png")), Inches(1), Inches(1), Inches(8), Inches(5.5))
    return _save(prs, d / "image.pptx")


def chart_deck(d: Path) -> Path:
    """Chart values exist only in the rendered PDF text layer, never in the shape walk (R8)."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Uptime by region"
    data = CategoryChartData()
    data.categories = ["Northwind", "Southfield", "Eastgate"]
    data.add_series("Uptime %", tuple(float(v) for v in CHART_VALUES))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), data)
    gf.chart.plots[0].has_data_labels = True
    return _save(prs, d / "chart.pptx")


def hidden_deck(d: Path) -> Path:
    """Four title+body slides; slide 2 is hidden, so the PDF has 3 pages."""
    prs = _prs()
    for n in range(1, 5):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Region {n} review"
        s.placeholders[1].text = f"Region {n} closed {n * 10} tickets in 2026."
    prs.slides[1]._element.set("show", "0")
    return _save(prs, d / "hidden.pptx")


def untitled_deck(d: Path) -> Path:
    """A slide with no title placeholder text and an empty body placeholder, plus one text box."""
    from pptx.util import Inches
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[1])        # title and body placeholders left empty
    s.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1)).text_frame.text = "Eastgate retired the mainframe in 2025."
    return _save(prs, d / "untitled.pptx")


def footnote_deck(d: Path, *, dated_title: bool = True) -> Path:
    """Slide 1 is a title slide (dated or not); slide 2 has a body and a small source line at the bottom."""
    from pptx.util import Inches, Pt
    prs = _prs()
    t = prs.slides.add_slide(prs.slide_layouts[0])
    t.shapes.title.text = "Quarterly review"
    t.placeholders[1].text = "Operations, March 2026" if dated_title else "Operations"
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Service levels"
    s.placeholders[1].text = "Uptime held at 99.9%."
    note = s.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(6), Inches(0.4)).text_frame
    note.text = FOOTNOTE
    note.paragraphs[0].runs[0].font.size = Pt(9)
    prs.core_properties.modified = datetime.datetime(2026, 4, 1, 9, 0)
    return _save(prs, d / ("footnote.pptx" if dated_title else "footnote_undated.pptx"))


def footer_deck(d: Path) -> Path:
    """A layout carrying footer text with a year and a number (rendered on every slide, held by no slide)."""
    import copy
    from pptx.util import Inches
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Service levels"
    s.placeholders[1].text = "Uptime held at 99.9%."
    box = s.shapes.add_textbox(Inches(0.3), Inches(7.0), Inches(5), Inches(0.3))
    box.text_frame.text = "Company confidential 2026 · Operations deck v3"
    layout_tree = s.slide_layout.shapes._spTree
    layout_tree.append(copy.deepcopy(box._element))
    box._element.getparent().remove(box._element)
    return _save(prs, d / "footer.pptx")


def picture_placeholder_deck(d: Path) -> Path:
    """'Picture with Caption' layout: the picture fills a placeholder, so its shape_type is PLACEHOLDER."""
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[8])
    s.shapes.title.text = "Target architecture"
    s.placeholders[1].insert_picture(str(_png(d / "arch.png", 400, 300)))
    return _save(prs, d / "picph.pptx")


def footer_placeholder_deck(d: Path) -> Path:
    """Header & Footer switched on: date, footer and slide-number placeholders cloned from the layout onto the slide."""
    import copy
    from pptx.enum.shapes import PP_PLACEHOLDER
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Service levels"
    s.placeholders[1].text = "Uptime held at 99.9%."
    texts = {PP_PLACEHOLDER.DATE: "10/12/2026", PP_PLACEHOLDER.FOOTER: "Confidential", PP_PLACEHOLDER.SLIDE_NUMBER: "3"}
    for ph in s.slide_layout.placeholders:
        if ph.placeholder_format.type in texts:
            el = copy.deepcopy(ph._element)
            s.shapes._spTree.append(el)
            shape = next(x for x in s.shapes if x._element is el)
            shape.text_frame.text = texts[ph.placeholder_format.type]
    return _save(prs, d / "footerph.pptx")


# --- off-screen deck (spec F7): only what is on the slide may be analysed -----------------------------------------
OFFSCREEN_KEEP = (
    "Uptime held at 99.9% all year.",                 # ordinary on-slide text box
    "CURRENT STATE READOUT",                          # letter-spaced label: renders as spaced letters, still on screen
    "Tickets",                                        # a table row (Tickets | 1,200)
    "Northwind sends orders to the warehouse.",       # text inside a group
    "Bottom line one: 5 regions report daily.",       # first paragraph of a box straddling the bottom edge
    *(f"Overflow item {k}: {k}0 units." for k in range(1, 10)),   # overflow that still sits on the slide
)
OFFSCREEN_DROP = (
    "Working note: Southfield closes 40 tickets daily.",          # box right of the slide
    "Parked quote: Eastgate retired 3 mainframes in 2024.",       # box left of the slide
    "Hidden draft: Westbrook owns 7 billing feeds.",              # shape with hidden="1"
    "Bottom line two: 9 regions report weekly.",                  # pushed below the slide's bottom edge
    *(f"Overflow item {k}: {k}0 units." for k in range(13, 17)),  # overflow past the slide's bottom edge
)
OFFSCREEN_NOTES = "Speaker note: Northwind revenue reached 12M in 2025."


def offscreen_deck(d: Path) -> Path:
    """One 10 x 7.5 in slide mixing on-screen text with text that is not on the slide. Overflow items 10-12 sit on
    the edge and are deliberately in neither list."""
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
    prs = _prs()
    W, H = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = lambda x, y, w, h, text: _tb(s.shapes, x, y, w, h, text)
    box(Inches(1), Inches(1.5), Inches(8), Inches(0.6), OFFSCREEN_KEEP[0])
    label = box(Inches(1), Inches(0.3), Inches(8), Inches(0.5), OFFSCREEN_KEEP[1])
    for r in label.text_frame.paragraphs[0].runs:
        r._r.get_or_add_rPr().set("spc", "800")                     # character spacing: "C U R R E N T ..."
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(2.3), Inches(4), Inches(0.8)).table
    for r, row in enumerate([("Metric", "Value"), ("Tickets", "1,200")]):
        for c, text in enumerate(row):
            tbl.cell(r, c).text = text
    grp = s.shapes.add_group_shape()
    _tb(grp.shapes, Inches(1), Inches(3.4), Inches(6), Inches(0.5), OFFSCREEN_KEEP[3])
    box(W + Inches(1), Inches(1), Inches(4), Inches(0.6), OFFSCREEN_DROP[0])
    box(-Inches(5), Inches(2), Inches(4), Inches(0.6), OFFSCREEN_DROP[1])
    hid = box(Inches(1), Inches(4.2), Inches(6), Inches(0.5), OFFSCREEN_DROP[2])
    hid._element.nvSpPr.cNvPr.set("hidden", "1")
    edge = box(Inches(1), H - Inches(0.6), Inches(4), Inches(1.0), OFFSCREEN_KEEP[4])
    p2 = edge.text_frame.add_paragraph()
    p2.text, p2.space_before = OFFSCREEN_DROP[3], Pt(60)
    over = box(Inches(5.5), Inches(4.6), Inches(4), Inches(0.4), "Overflow item 1: 10 units.")
    for k in range(2, 17):
        over.text_frame.add_paragraph().text = f"Overflow item {k}: {k}0 units."
    for p in over.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(14)
    over.text_frame._txBody.bodyPr.set("wrap", "square")
    s.notes_slide.notes_text_frame.text = OFFSCREEN_NOTES
    return _save(prs, d / "offscreen.pptx")


def _tb(shapes, x, y, w, h, text):
    from pptx.enum.text import MSO_AUTO_SIZE
    tb = shapes.add_textbox(x, y, w, h)
    tb.text_frame.text = text
    tb.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    return tb


# --- pin geometry deck (eval 2): six non-wrapping 18 pt paragraphs in one box at a known position -----------------
PINS_BOX = (914400, 914400, 8 * 914400, 4 * 914400)      # x, y, w, h in EMU
PINS_PT = 18
PINS_TEXT = tuple(f"Claim {k}: Region {k} closed {k}40 tickets in 2026." for k in range(1, 7))


def pins_deck(d: Path) -> Path:
    from pptx.util import Pt
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = _tb(s.shapes, *PINS_BOX, PINS_TEXT[0])
    tb.text_frame.word_wrap = False
    for t in PINS_TEXT[1:]:
        tb.text_frame.add_paragraph().text = t
    for p in tb.text_frame.paragraphs:
        p.space_before = p.space_after = Pt(0)
        for r in p.runs:
            r.font.size = Pt(PINS_PT)
    return _save(prs, d / "pins.pptx")


def full_deck(d: Path) -> Path:
    """One deck with a body slide, a table slide and an image slide, for the annotator and invariant tests."""
    from pptx.util import Inches
    prs = _prs()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Delivery plan"
    _bullets(s1, [("Go-live is set for Q2 2027.", 0), ("Uptime held at 99.9% all year.", 0),
                  ("The integration team owns cut-over.", 0)])
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "North region — FY2025"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for r, row in enumerate([("Metric", "Value"), ("Uptime", "99.9%")]):
        for c, text in enumerate(row):
            tbl.cell(r, c).text = text
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    s3.shapes.add_picture(str(_png(d / "pic.png")), Inches(1), Inches(1), Inches(8), Inches(5.5))
    return _save(prs, d / "full.pptx")


def _save(prs, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


# --- round 3 (spec §10): letter-spaced labels, navigation tabs, attribution lines, figure elements ----------------
SPACED_LABEL = "CURRENT STATE READOUT"
TABS = ("DISCOVER & CHOOSE", "SIGN UP & FIRST ORDER", "ONE-OFF ORDERS", "ISSUE RESOLUTION")
TAB_BODIES = tuple(f"Depot {c} ships {n} crates per week in 2025." for c, n in zip("ABCD", (1200, 900, 700, 400)))
ATTR_QUOTE = "“I would say more than 40% of our calls are about late deliveries.”"
ATTR_LINE = "Discovery Session – Logistics; 00:12:01–00:12:30"
ATTR_STATEMENT = "The nightly routing job skips depots that are closed for maintenance."
ATTR_LINE_2 = "Routing Workshop; Oct 3"
ELEMENT_PIC = (5 * EMU_PER_INCH, 1 * EMU_PER_INCH, 4 * EMU_PER_INCH, 2 * EMU_PER_INCH)   # x, y, w, h
ELEMENT_CHART = (1 * EMU_PER_INCH, 3 * EMU_PER_INCH, 4 * EMU_PER_INCH, 3 * EMU_PER_INCH)


def spaced_deck(d: Path) -> Path:
    """A letter-spaced eyebrow (rendered as single letters) over plain body text, nothing visual."""
    from pptx.util import Inches, Pt
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = _tb(s.shapes, Inches(0.7), Inches(0.4), Inches(6), Inches(0.4), SPACED_LABEL)
    for r in tb.text_frame.paragraphs[0].runs:
        r.font.size = Pt(12)
        r._r.get_or_add_rPr().set("spc", "600")
    _tb(s.shapes, Inches(0.7), Inches(1), Inches(8.5), Inches(0.6), "Returns are rising in the north")
    _tb(s.shapes, Inches(0.7), Inches(2), Inches(8.5), Inches(1.2),
        "Depot staff reported that weekend returns now arrive in two waves.").text_frame.word_wrap = True
    return _save(prs, d / "spaced.pptx")


def tabs_deck(d: Path) -> Path:
    """Four slides sharing a tab bar at the same position (page furniture), one body statement each."""
    from pptx.util import Inches, Pt
    prs = _prs()
    for body in TAB_BODIES:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for i, t in enumerate(TABS):
            sh = s.shapes.add_shape(1, Inches(0.3 + 2.35 * i), Inches(0.5), Inches(2.2), Inches(0.33))
            sh.text_frame.text = t
            sh.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
        _tb(s.shapes, Inches(0.7), Inches(1.4), Inches(8.5), Inches(0.5), body.split(" ships")[0] + " volumes")
        _tb(s.shapes, Inches(0.7), Inches(2.4), Inches(8.5), Inches(0.8), body)
    return _save(prs, d / "tabs.pptx")


def attribution_deck(d: Path) -> Path:
    """An evidence box: a quote, its attribution line, a statement and its attribution line."""
    from pptx.util import Inches, Pt
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _tb(s.shapes, Inches(0.7), Inches(0.4), Inches(8.5), Inches(0.6), "Late deliveries drive calls")
    tb = _tb(s.shapes, Inches(5.5), Inches(1.5), Inches(4), Inches(5), ATTR_QUOTE)
    tb.text_frame.word_wrap = True
    for t in (ATTR_LINE, ATTR_STATEMENT, ATTR_LINE_2):
        tb.text_frame.add_paragraph().text = t
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(9 if "; " in p.text else 14)
    return _save(prs, d / "attribution.pptx")


def elements_deck(d: Path) -> Path:
    """One slide with a title, a sizeable picture and a column chart at known boxes."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _tb(s.shapes, Inches(0.7), Inches(0.3), Inches(8.5), Inches(0.6), "Regional uptime review")
    pic = s.shapes.add_picture(str(_png(d / "callout.png", 200, 100)), *ELEMENT_PIC)
    pic.name = "Uptime callout"
    cd = CategoryChartData()
    cd.categories = ["North", "South", "East"]
    cd.add_series("Uptime %", tuple(float(v) for v in CHART_VALUES))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *ELEMENT_CHART, cd)
    return _save(prs, d / "elements.pptx")


CALLOUT_BOX = (5 * EMU_PER_INCH, int(0.2 * EMU_PER_INCH), int(3.5 * EMU_PER_INCH), int(0.45 * EMU_PER_INCH))  # ~2% of the slide
LOGO_BOX = (int(0.2 * EMU_PER_INCH), int(6.9 * EMU_PER_INCH), int(0.6 * EMU_PER_INCH), int(0.4 * EMU_PER_INCH))   # <0.5%


def callout_deck(d: Path) -> Path:
    """Slide 1: body text and a small picture callout (~2% of the slide) plus a logo-sized picture; slide 2: text and
    the logo only; slide 3: numbered step badges."""
    from pptx.util import Inches
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _tb(s.shapes, Inches(0.7), Inches(1), Inches(8.5), Inches(0.6), "Depot hours are set by each region.")
    s.shapes.add_picture(str(_png(d / "callout.png", 140, 18)), *CALLOUT_BOX).name = "Hours callout"
    s.shapes.add_picture(str(_png(d / "logo.png", 24, 16)), *LOGO_BOX)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    _tb(s2.shapes, Inches(0.7), Inches(1), Inches(8.5), Inches(0.6), "Returns are handled by the depot team.")
    s2.shapes.add_picture(str(_png(d / "logo2.png", 24, 16)), *LOGO_BOX)
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(1, 4):
        _tb(s3.shapes, Inches(0.5 + 3 * i - 3), Inches(2), Inches(0.4), Inches(0.4), str(i))
        _tb(s3.shapes, Inches(1 + 3 * i - 3), Inches(2), Inches(2.4), Inches(0.6), f"Step text {'ABC'[i - 1]}")
    return _save(prs, d / "callout.pptx")


# --- round 4 (spec §11): slide context and structural labels -----------------------------------------------------
CTX_TITLE = "A changed delivery can still miss the loaded truck"
CTX_LEDE = "A customer calls to change tomorrow's delivery and the change travels to the depot handheld."
CTX_KEY = "Reported flow; dashed boxes are not confirmed"
CTX_LABELS = ("2 ORDER PLACED", "COMMERCE & IPX (INTEGRATION)", "WHERE IT CAN BREAK")
CTX_FIGURE_LABEL = "UPTIME 99.9%"
CTX_CARD = "OrderHub saves the change as the master record"


def context_deck(d: Path) -> Path:
    """A diagram-like slide built from text boxes: title, lede, stage/lane labels, one card, an all-caps figure,
    and a small key line at the bottom."""
    from pptx.util import Inches, Pt
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = _tb(s.shapes, Inches(0.5), Inches(0.4), Inches(9), Inches(0.6), CTX_TITLE)
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    _tb(s.shapes, Inches(0.5), Inches(1.05), Inches(9), Inches(0.6), CTX_LEDE).text_frame.word_wrap = True
    for i, lab in enumerate(CTX_LABELS):
        _tb(s.shapes, Inches(0.5 + 3 * i), Inches(2), Inches(2.8), Inches(0.4), lab)
    _tb(s.shapes, Inches(0.5), Inches(3), Inches(4), Inches(0.6), CTX_CARD)
    _tb(s.shapes, Inches(5), Inches(3), Inches(4), Inches(0.6), CTX_FIGURE_LABEL)
    k = _tb(s.shapes, Inches(0.5), Inches(6.9), Inches(6), Inches(0.3), CTX_KEY)
    k.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    return _save(prs, d / "context.pptx")
