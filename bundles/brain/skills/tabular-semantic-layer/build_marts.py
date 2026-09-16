#!/usr/bin/env python3
"""Config-driven ETL: heterogeneous Excel reporting workbooks -> normalized long
facts (Parquet) -> SQLite views. The generic loader; all corpus specifics live in
the families JSON. Never loads a whole workbook into memory (openpyxl read_only).

Output long schema:
  facts(family, metric, grain, entity, month, value:double, source_file)

Usage:
  build_marts.py --root <reporting dir> --config <families.json> --out-dir <dir>
"""
import argparse, glob, json, os, re, sys, warnings
warnings.filterwarnings("ignore")
import openpyxl

MONTHS = {m: i+1 for i, m in enumerate(
    ["january","february","march","april","may","june","july","august",
     "september","october","november","december"])}
ABBR = {m[:3]: i+1 for m, i in [(k, v-1) for k, v in MONTHS.items()]}

def month_from_filename(fn, default_year=2026):
    """First month name by POSITION in the string (so 'June ... May' -> June)."""
    low = fn.lower()
    best_pos, found = 10**9, None
    for name, num in MONTHS.items():
        p = low.find(name)
        if p != -1 and p < best_pos:
            best_pos, found = p, num
    if found is None:
        for ab, num in ABBR.items():
            m = re.search(r"\b"+ab, low)
            if m and m.start() < best_pos:
                best_pos, found = m.start(), num
    yr = default_year
    ym = re.search(r"20(2\d)", fn)
    if ym: yr = int("20"+ym.group(1))
    return f"{yr}-{found:02d}" if found else None

def norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip().lower())

_TC_DIMS = set()   # dimensions to canonicalize by title-casing (set from config)

def smart_title(s):
    # Title-case each alnum run; preserve punctuation/parentheticals; keep short all-caps acronyms
    def cap(w):
        return w if (w.isupper() and len(w) <= 3) else w.capitalize()
    return re.sub(r"[A-Za-z0-9]+", lambda m: cap(m.group(0)), str(s))

def conform(dim_type, value, dim_map):
    v = re.sub(r"\s+", " ", str(value).strip())
    mapped = dim_map.get(dim_type, {}).get(norm(v))
    if mapped is not None:
        return mapped                      # explicit alias wins
    if dim_type in _TC_DIMS:
        return smart_title(v)              # canonical casing for branch/region/etc.
    return v

def get_sheet(wb, name):
    if name in wb.sheetnames:
        return wb[name]
    for s in wb.sheetnames:               # tolerant fallback: contains
        if norm(name) in norm(s):
            return wb[s]
    return None

def read_rows(ws, cap=100000):
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        out.append(list(row))
        if i >= cap: break
    return out

def alts(spec):
    """A measure spec is a header substring or a list of alternative substrings."""
    return spec if isinstance(spec, list) else [spec]

def header_matches(cell_norm, spec):      # substring (flexible; for find_header)
    return any(norm(a) in cell_norm for a in alts(spec))

def header_exact(cell_norm, spec):         # exact per alternative (for stable wide layouts)
    return any(norm(a) == cell_norm for a in alts(spec))

def find_header(rows, dim_header, measure_specs, hint=None):
    """Row index whose cells contain the dim header + >=1 measure header."""
    cand = range(len(rows)) if hint is None else [hint] + list(range(len(rows)))
    for i in cand:
        if i >= len(rows): continue
        cells = [norm(c) for c in rows[i]]
        if norm(dim_header) in cells and any(
                any(header_matches(c, ms) for c in cells) for ms in measure_specs):
            return i
    return hint

def col_index(header, spec):
    for a in alts(spec):                       # exact match first, per alternative
        na = norm(a)
        for j, c in enumerate(header):
            if norm(c) == na: return j
    for a in alts(spec):                        # then substring
        na = norm(a)
        for j, c in enumerate(header):
            if na in norm(c): return j
    return None

def to_float(v):
    import datetime
    if isinstance(v, bool): return None
    if isinstance(v, (int, float)): return float(v)
    if isinstance(v, datetime.time):
        return v.hour*3600 + v.minute*60 + v.second + v.microsecond/1e6
    if isinstance(v, str):
        s = v.replace(",", "").replace("$", "").replace("%", "").strip()
        try: return float(s)
        except ValueError: return None
    return None

def load_tolerant_sheet(rows, family, month, dim_candidates, measures, dim_map, entity_regex):
    """For template-less exports (e.g. NPS): find the header row by CONTENT (any dim
    candidate + >=1 measure), infer grain from the matched dim, map measures by name
    wherever they sit, and clean entity encodings (e.g. 'SE1: Clinton Poche' -> 'SE1')."""
    import re as _re
    hdr_i = matched_grain = dcol = None
    for i, row in enumerate(rows[:40]):
        cells = [norm(c) for c in row]
        if not any(any(header_matches(c, ms) for c in cells) for ms in measures.values()):
            continue
        for grain, cands in dim_candidates.items():
            for cand in cands:
                if norm(cand) in cells:
                    hdr_i, matched_grain, dcol = i, grain, cells.index(norm(cand)); break
            if hdr_i is not None: break
        if hdr_i is not None: break
    if hdr_i is None:
        return [], {"reason": "no header row with a dim candidate + a measure"}
    header = rows[hdr_i]
    mcols = {m: col_index(header, txt) for m, txt in measures.items()}
    missing = [m for m, j in mcols.items() if j is None]
    mcols = {m: j for m, j in mcols.items() if j is not None}
    facts = []
    for r in rows[hdr_i+1:]:
        if dcol >= len(r) or not is_dim_value(r[dcol]): continue
        raw = str(r[dcol]).strip()
        ent = raw
        if entity_regex:
            m = _re.match(entity_regex, raw)
            if m: ent = m.group(1)
        ent = conform(matched_grain, ent, dim_map)
        for m, j in mcols.items():
            if j < len(r):
                v = to_float(r[j])
                if v is not None:
                    facts.append((family["name"], m, matched_grain, ent, month, v))
    diag = {"reason": (f"partial: {missing}" if missing else (None if facts else "no data rows")),
            "grain": matched_grain, "missing_measures": missing}
    return facts, diag

def load_matrix(rows, family, year):
    """metric-rows x month-columns (e.g. workforce MOM sheet). Month header row maps
    column -> month; col A holds metric labels; take the FIRST occurrence of each
    (top consolidated block). Time cells (ASA/AHT) become seconds via to_float."""
    hr = family.get("month_header_row", 1)
    mcol = family.get("metric_col", 0)
    if hr >= len(rows): return [], {"reason": f"fewer rows than month_header_row={hr}"}
    # month -> FIRST (leftmost) column: the sheet repeats month headers across several
    # sections (monthly totals, daily avg, weekly); the leftmost block is monthly totals.
    month2col = {}
    for j, cell in enumerate(rows[hr]):
        if j == mcol: continue
        mm = month_from_filename(str(cell), year) if cell not in (None, "") else None
        if mm and mm not in month2col:
            month2col[mm] = j
    entity = list(family["grains"].values())[0].get("entity", "overall")
    facts = []
    matched, missing = [], []
    for m, label in family["measures"].items():
        found = False
        for r in rows[hr+1:]:                          # first matching row only
            if mcol < len(r) and r[mcol] and header_exact(norm(r[mcol]), label):
                for month, j in month2col.items():
                    if j < len(r):
                        val = to_float(r[j])
                        if val is not None:
                            facts.append((family["name"], m, "overall", entity, month, val))
                found = True; break
        (matched if found else missing).append(m)
    diag = {"reason": None, "measures_matched": matched, "missing_measures": missing}
    if not month2col: diag["reason"] = f"no month headers found in row {hr}"
    elif missing: diag["reason"] = f"partial: labels not matched: {missing}"
    return facts, diag

def is_dim_value(v):
    if v in (None, ""): return False
    s = str(v).strip().lower()
    return s not in ("total", "grand total", "all", "sum",
                     "#n/a", "n/a", "#ref!", "#value!", "#div/0!", "(blank)", "blank")

def load_long(rows, family, grain, dim_type, dim_header, measures, dim_map):
    hdr_i = find_header(rows, dim_header, measures.values(),
                        hint=family.get("header_row"))
    if hdr_i is None:
        return [], {"reason": "no header row containing dim + a measure"}
    header = rows[hdr_i]
    dcol = col_index(header, dim_header)
    mcols = {m: col_index(header, txt) for m, txt in measures.items()}
    missing = [m for m, j in mcols.items() if j is None]
    mcols = {m: j for m, j in mcols.items() if j is not None}
    if dcol is None:
        return [], {"reason": f"dim column '{dim_header}' not found in header", "header": hdr_i}
    if not mcols:
        return [], {"reason": "no measure columns matched", "missing_measures": missing, "header": hdr_i}
    facts = []
    month = family["_month"]
    started = False
    for r in rows[hdr_i+1:]:
        # stop at first fully-blank row once the table has started (table boundary)
        if all(c in (None, "") for c in r):
            if started: break
            continue
        if dcol >= len(r) or not is_dim_value(r[dcol]):
            continue
        started = True
        ent = conform(dim_type, r[dcol], dim_map)
        for m, j in mcols.items():
            if j < len(r):
                val = to_float(r[j])
                if val is not None:
                    facts.append((family["name"], m, grain, ent, month, val))
    diag = {"reason": None, "header": hdr_i, "measures_matched": list(mcols.keys()),
            "missing_measures": missing}
    if missing:
        diag["reason"] = f"partial: measures not matched: {missing}"
    return facts, diag

def load_wide_month(rows, family, grain, dim_type, dim_header, measures, dim_map):
    """Adjustments-style: banner row has month dates per block; header row repeats
    measure names. Forward-fill month across columns."""
    br = family.get("banner_row", 0); hr = family.get("header_row", 1)
    if hr >= len(rows): return [], {"reason": f"fewer rows than header_row={hr}"}
    banner, header = rows[br], rows[hr]
    # forward-fill month label across columns
    colmonth, cur = {}, None
    import datetime
    for j, b in enumerate(banner):
        if isinstance(b, datetime.datetime):
            cur = f"{b.year}-{b.month:02d}"          # new month block starts
        elif isinstance(b, str) and b.strip():
            cur = month_from_filename(b)              # non-date label -> month or None (resets)
        # blank banner cell -> forward-fill current block's month
        colmonth[j] = cur
    # dim label may sit on the banner row (r0) or the measure-header row (r1) depending
    # on the month's template — search both.
    dcol = col_index(header, dim_header)
    if dcol is None:
        dcol = col_index(banner, dim_header)
    if dcol is None:
        return [], {"reason": f"dim column '{dim_header}' not found in rows {br}/{hr}"}
    facts = []
    for r in rows[hr+1:]:
        if dcol >= len(r) or not is_dim_value(r[dcol]): continue
        ent = conform(dim_type, r[dcol], dim_map)
        for j, cell in enumerate(header):
            for m, txt in measures.items():
                if header_exact(norm(cell), txt) and colmonth.get(j) and j < len(r):
                    val = to_float(r[j])
                    if val is not None:
                        facts.append((family["name"], m, grain, ent, colmonth[j], val))
    return facts, {"reason": None if facts else "no month-banner columns matched measures"}

def _month_range(start, end):
    """Inclusive monthly range over 'YYYY-MM' strings."""
    ys, ms = map(int, start.split("-")); ye, me = map(int, end.split("-"))
    out, y, m = [], ys, ms
    while (y, m) <= (ye, me):
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m > 12: m, y = 1, y + 1
    return out

def apply_derived(df, cfg):
    """Append derived/ratio metrics: metricC = numerator / denominator at a grain,
    computed deterministically and cited "<derived>" (never a hand-precomputed source
    cell). Config: derived: [{family, metric, grain, numerator, denominator, scale?}].
    Honest by construction: operands are inner-joined on (entity, month) so a MISSING
    operand yields NO row, and denominator==0 rows are dropped — never a fabricated value."""
    import pandas as pd
    for spec in cfg.get("derived", []):
        fam, grain = spec["family"], spec["grain"]
        base = df[(df.family == fam) & (df.grain == grain)]
        num = base[base.metric == spec["numerator"]][["entity", "month", "value"]].rename(columns={"value": "num"})
        den = base[base.metric == spec["denominator"]][["entity", "month", "value"]].rename(columns={"value": "den"})
        if num.empty or den.empty:
            continue
        merged = num.merge(den, on=["entity", "month"], how="inner")
        merged = merged[merged["den"] != 0]
        scale = spec.get("scale", 1.0)
        rows = [(fam, spec["metric"], grain, r.entity, r.month, (r.num / r.den) * scale, "<derived>")
                for r in merged.itertuples()]
        if rows:
            df = pd.concat([df, pd.DataFrame(rows, columns=df.columns)], ignore_index=True)
    return df

def coverage_report(df, cfg):
    """Completeness against the observed grid (+ an optional expected roster). Distinct
    from build_audit (per-file PARSE health): this catches an entity/month that never
    appeared at all — the disappearing-division / missing-month silent gap.
    - intra_family_holes: an entity present in some of a family/grain's months but absent
      in others (e.g. a division that stops appearing after Jan). Detected with NO config.
    - expected_violations: months/entities named in cfg['coverage'] that are wholly absent
      (e.g. a month whose source file was never produced/ingested). Needs config.
    Also returns the grains each family carries (grain visibility)."""
    expected = cfg.get("coverage", {})
    grains, holes, violations = [], [], []
    fam_grains = {}
    for (fam, grain), g in df.groupby(["family", "grain"]):
        months = sorted(m for m in g["month"].dropna().unique())
        ents = sorted(e for e in g["entity"].dropna().unique())
        present = set(zip(g["entity"], g["month"]))
        ent_holes = {e: miss for e in ents
                     if (miss := [m for m in months if (e, m) not in present])}
        exp = expected.get(fam) or expected.get("*")
        if exp:
            exp_months = _month_range(*exp["month_range"]) if exp.get("month_range") else list(exp.get("months", []))
            miss_m = [m for m in exp_months if m not in set(months)]
            if miss_m: violations.append({"family": fam, "grain": grain, "missing_months": miss_m})
            miss_e = [e for e in exp.get("entities", []) if e not in set(ents)]
            if miss_e: violations.append({"family": fam, "grain": grain, "missing_entities": miss_e})
        grains.append({"family": fam, "grain": grain, "n_months": len(months),
                       "month_span": [months[0], months[-1]] if months else [],
                       "n_entities": len(ents), "holes": ent_holes})
        fam_grains.setdefault(fam, []).append(grain)
        if ent_holes: holes.append({"family": fam, "grain": grain, "holes": ent_holes})
    return {"grains": grains, "family_grains": fam_grains,
            "intra_family_holes": holes, "expected_violations": violations}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--config", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--db", default=None, help="SQLite file to write the `facts` table into (default <out-dir>/knowledge.sqlite); point it at your unified knowledge.sqlite")
    ap.add_argument("--strict", action="store_true",
                    help="exit non-zero if any globbed file produced zero facts (CI/prod guard)")
    a = ap.parse_args()
    cfg = json.load(open(a.config))
    dim_map = cfg.get("dimension_map", {})
    global _TC_DIMS
    _TC_DIMS = set(cfg.get("conform_titlecase", []))
    os.makedirs(a.out_dir, exist_ok=True)

    all_facts = []
    audit = []   # one row per (file, unit): status ok|partial|zero|benign|error + reason
    # per-family substrings for files that are EXPECTED to yield nothing (redundant/legacy)
    benign = {f["name"]: [p.lower() for p in f.get("allow_zero", [])] for f in cfg["families"]}

    def record(family, fn, unit, facts, diag):
        got = [t + (fn,) for t in facts]
        all_facts.extend(got)
        status = "ok"
        if not got:
            status = "zero"
            if fn and any(p in fn.lower() for p in benign.get(family, [])):
                status = "benign"   # expected-empty; won't fail --strict
        elif diag.get("reason"): status = "partial"
        audit.append({"family": family, "file": fn, "unit": unit, "facts": len(got),
                      "status": status, "reason": diag.get("reason")})

    for fam in cfg["families"]:
        files = sorted(glob.glob(os.path.join(a.root, fam["glob"])))
        if fam["layout"] == "matrix_month_cols" and files:   # snapshot holds full history
            files = [max(files, key=lambda p: month_from_filename(os.path.basename(p)) or "")]
        if not files:
            audit.append({"family": fam["name"], "file": None, "unit": "-", "facts": 0,
                          "status": "zero", "reason": f"glob matched no files: {fam['glob']}"})
        for fp in files:
            fn = os.path.basename(fp)
            fam["_month"] = month_from_filename(fn) if fam["month_from"] == "filename" else None
            try:
                wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
            except Exception as e:
                audit.append({"family": fam["name"], "file": fn, "unit": "-", "facts": 0,
                              "status": "error", "reason": f"open failed: {e}"}); continue
            if fam["layout"] == "matrix_month_cols":
                ws = get_sheet(wb, fam["sheet"])
                if ws is None:
                    record(fam["name"], fn, fam["sheet"], [], {"reason": f"no sheet '{fam['sheet']}'"})
                else:
                    yr = 2026
                    ym = re.search(r"20(2\d)", fn)
                    if ym: yr = int("20"+ym.group(1))
                    facts, diag = load_matrix(read_rows(ws), fam, yr)
                    record(fam["name"], fn, fam["sheet"], facts, diag)
                wb.close(); continue
            if fam["layout"] == "tolerant_long":
                total = []
                for sh in wb.sheetnames:
                    facts, _ = load_tolerant_sheet(read_rows(wb[sh]), fam, fam["_month"],
                                                   fam["dim_candidates"], fam["measures"],
                                                   dim_map, fam.get("entity_regex"))
                    total.extend(facts)
                record(fam["name"], fn, "auto", total,
                       {"reason": None if total else "no sheet yielded rows (dim+measure not found)"})
                wb.close(); continue
            for grain, gc in fam["grains"].items():
                ws = get_sheet(wb, gc["sheet"])
                if ws is None:
                    record(fam["name"], fn, grain, [], {"reason": f"no sheet '{gc['sheet']}'"}); continue
                dim_type = "region" if grain == "region" else grain
                loader = load_wide_month if fam["layout"] == "wide_month" else load_long
                facts, diag = loader(read_rows(ws), fam, grain, dim_type, gc["dim_header"], fam["measures"], dim_map)
                record(fam["name"], fn, grain, facts, diag)
            wb.close()

    # write parquet + sqlite `facts`
    import pandas as pd, sqlite3
    df = pd.DataFrame(all_facts, columns=["family","metric","grain","entity","month","value","source_file"])
    df = df.drop_duplicates(subset=["family","metric","grain","entity","month"], keep="last").reset_index(drop=True)

    # weighted rollups: aggregate a finer grain up to a coarser one (e.g. NPS region -> division),
    # weighting rate metrics by a count metric. Config: rollups: [{family, from, to, weight, map_prefix}]
    for spec in cfg.get("rollups", []):
        sub = df[(df.family == spec["family"]) & (df.grain == spec["from"])].copy()
        if sub.empty: continue
        pref = spec["map_prefix"]
        def to_coarse(e):
            for p, v in pref.items():
                if str(e).upper().startswith(p.upper()): return v
            return None
        sub["coarse"] = sub["entity"].map(to_coarse)
        sub = sub[sub["coarse"].notna()]
        wm = spec["weight"]
        w = sub[sub.metric == wm][["entity","month","value"]].rename(columns={"value":"wt"})
        rate = sub[sub.metric != wm].merge(w, on=["entity","month"], how="left")
        rate["wt"] = rate["wt"].fillna(0.0); rate["wv"] = rate["value"] * rate["wt"]
        g = rate.groupby(["coarse","metric","month"]).agg(wvsum=("wv","sum"), wsum=("wt","sum")).reset_index()
        g = g[g.wsum > 0]
        rows = [(spec["family"], r.metric, spec["to"], r.coarse, r.month, r.wvsum / r.wsum, "<rollup>")
                for r in g.itertuples()]
        gw = sub[sub.metric == wm].groupby(["coarse","month"]).agg(value=("value","sum")).reset_index()
        rows += [(spec["family"], wm, spec["to"], r.coarse, r.month, r.value, "<rollup>") for r in gw.itertuples()]
        df = pd.concat([df, pd.DataFrame(rows, columns=df.columns)], ignore_index=True)

    df = apply_derived(df, cfg)

    pq = os.path.join(a.out_dir, "facts.parquet")
    try: df.to_parquet(pq, index=False)
    except Exception: pq = None       # parquet optional (pyarrow); sqlite is the store
    db = a.db or os.path.join(a.out_dir, "knowledge.sqlite")
    con = sqlite3.connect(db)
    df.to_sql("facts", con, if_exists="replace", index=False)
    con.execute("CREATE INDEX IF NOT EXISTS idx_facts ON facts(family, metric, grain, entity, month)")
    con.commit(); con.close()

    # audit summary — silent skips are made LOUD here
    zeros = [x for x in audit if x["status"] in ("zero", "error")]
    partials = [x for x in audit if x["status"] == "partial"]
    benigns = [x for x in audit if x["status"] == "benign"]
    with open(os.path.join(a.out_dir, "build_audit.json"), "w") as f:
        json.dump({"summary": {"units": len(audit),
                               "ok": sum(1 for x in audit if x["status"]=="ok"),
                               "partial": len(partials), "benign": len(benigns),
                               "zero_or_error": len(zeros)},
                   "audit": audit}, f, indent=2)
    print(f"TOTAL {len(df)} facts -> SQLite: {db} (table: facts)" + (f"  [+parquet {pq}]" if pq else ""), file=sys.stderr)
    if len(df):
        cov = df.groupby(["family","grain"]).agg(
            facts=("value","size"), months=("month","nunique"), entities=("entity","nunique")).reset_index()
        print("\ncoverage:\n" + cov.to_string(index=False), file=sys.stderr)

    # coverage matrix — completeness against the observed grid (+ optional expected roster).
    # Catches the silent gap build_audit CANNOT: an entity/month that never appeared.
    covrep = coverage_report(df, cfg) if len(df) else {"grains": [], "family_grains": {},
                                                       "intra_family_holes": [], "expected_violations": []}
    with open(os.path.join(a.out_dir, "coverage.json"), "w") as f:
        json.dump(covrep, f, indent=2)
    if covrep["family_grains"]:
        print("\ngrains per family:", file=sys.stderr)
        for fam, gs in sorted(covrep["family_grains"].items()):
            print(f"   - {fam}: {', '.join(sorted(gs))}", file=sys.stderr)
    holes = covrep["intra_family_holes"]
    if holes:
        print(f"\n⚠️  {len(holes)} family/grain(s) with COVERAGE HOLES (entity present some months, missing others):", file=sys.stderr)
        for h in holes:
            for e, miss in h["holes"].items():
                print(f"   - {h['family']}/{h['grain']}: '{e}' missing {len(miss)} month(s): {', '.join(miss)}", file=sys.stderr)
    viols = covrep["expected_violations"]
    if viols:
        print(f"\n❌ {len(viols)} EXPECTED-ROSTER violation(s) (configured months/entities wholly absent):", file=sys.stderr)
        for v in viols:
            what = f"missing_months={v['missing_months']}" if "missing_months" in v else f"missing_entities={v['missing_entities']}"
            print(f"   - {v['family']}/{v['grain']}: {what}", file=sys.stderr)
    print(f"\ncoverage -> {os.path.join(a.out_dir,'coverage.json')}", file=sys.stderr)
    if benigns:
        print(f"\nℹ️  {len(benigns)} expected-empty (allow_zero) — not failures:", file=sys.stderr)
        for x in benigns: print(f"   - {x['family']} [{x['file']}]: {x['reason']}", file=sys.stderr)
    if partials:
        print(f"\n⚠️  {len(partials)} PARTIAL extractions (some measures/labels unmatched):", file=sys.stderr)
        for x in partials: print(f"   - {x['family']}/{x['unit']} [{x['file']}]: {x['reason']}", file=sys.stderr)
    if zeros:
        print(f"\n❌ {len(zeros)} ZERO-FACT units — a globbed file yielded NOTHING (likely a silent gap):", file=sys.stderr)
        for x in zeros: print(f"   - {x['family']}/{x['unit']} [{x['file']}]: {x['reason']}", file=sys.stderr)
    print(f"\naudit -> {os.path.join(a.out_dir,'build_audit.json')}", file=sys.stderr)
    if a.strict and (zeros or partials or viols):
        print(f"\nSTRICT: failing build ({len(zeros)} zero/error, {len(partials)} partial, "
              f"{len(viols)} expected-roster violation(s); {len(benigns)} benign ignored).", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
