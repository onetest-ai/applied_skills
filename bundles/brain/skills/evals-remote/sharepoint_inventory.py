"""Parse the SharePoint Power Query export (docs/kb-storage.xlsx, sheet 'Query')."""
import argparse
import json
import os
import re

import openpyxl


def normalize(filename):
    return re.sub(r"[^a-z0-9]+", "", str(filename).lower())


def read_inventory(xlsx_path, sheet="Query"):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    idx = {h: i for i, h in enumerate(header)}

    def g(row, col):
        i = idx.get(col)
        return row[i] if i is not None and i < len(row) else None

    out = []
    for row in rows_iter:
        name = g(row, "Name")
        if name is None:
            continue
        ext = str(g(row, "Extension") or "").strip()
        name = str(name)
        if not ext:
            # No Extension column or a blank cell: the Name suffix is the best evidence.
            ext = os.path.splitext(name)[1]
        # The SharePoint 'Name' column already carries the extension for most rows;
        # only append 'Extension' when it isn't already the suffix (avoids 'x.pdf.pdf').
        if ext and not name.lower().endswith(ext.lower()):
            filename = f"{name}{ext}"
        else:
            filename = name
        out.append({
            "name": str(name),
            "ext": str(ext),
            "modified": str(g(row, "Date modified") or ""),
            "created": str(g(row, "Date created") or ""),
            "folder": str(g(row, "Folder Path") or ""),
            "filename": filename,
            "match_key": normalize(filename),
        })
    return out


def main(argv=None):
    p = argparse.ArgumentParser()
    p.add_argument("xlsx")
    p.add_argument("--sheet", default="Query")
    a = p.parse_args(argv)
    print(json.dumps(read_inventory(a.xlsx, a.sheet), indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
