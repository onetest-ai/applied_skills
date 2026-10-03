import openpyxl

import sharepoint_inventory as si


def _make_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Query"
    ws.append(["Name", "Extension", "Date accessed", "Date modified", "Date created", "Folder Path"])
    ws.append(["Team_Demo_03.14.2022", ".pptx", None, "2022-03-14 10:00:00", "2022-03-14 10:00:00",
               "https://files.example.com/site/Prior engagement 2021 - 2023/"])
    ws.append(["Monthly Ops January 2026", ".xlsm", None, "2026-02-07 09:00:00", "2026-01-31 09:00:00",
               "https://files.example.com/site/Inputs/Ops Reporting/"])
    p = tmp_path / "kb.xlsx"
    wb.save(p)
    return str(p)


def test_read_inventory_parses_query_sheet(tmp_path):
    rows = si.read_inventory(_make_xlsx(tmp_path))
    assert len(rows) == 2
    r0 = rows[0]
    assert r0["filename"] == "Team_Demo_03.14.2022.pptx"
    assert "Prior engagement" in r0["folder"]
    assert r0["match_key"] == si.normalize("Team_Demo_03.14.2022.pptx")


def test_normalize_is_stable():
    assert si.normalize("A B.PPTX") == si.normalize("a  b.pptx")


def test_name_already_carrying_extension_is_not_doubled(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Query"
    ws.append(["Name", "Extension", "Date accessed", "Date modified", "Date created", "Folder Path"])
    ws.append(["Team.xlsx", ".xlsx", None, "2026-08-01", "2026-08-01", "https://x/Team/"])
    p = tmp_path / "kb.xlsx"
    wb.save(p)
    rows = si.read_inventory(str(p))
    assert rows[0]["filename"] == "Team.xlsx"  # not "Team.xlsx.xlsx"
    assert rows[0]["match_key"] == si.normalize("Team.xlsx")


def _xlsx_with(tmp_path, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Query"
    ws.append(header)
    for r in rows:
        ws.append(r)
    p = tmp_path / "kb_ext.xlsx"
    wb.save(p)
    return str(p)


def test_missing_extension_column_falls_back_to_the_name_suffix(tmp_path):
    p = _xlsx_with(tmp_path, ["Name", "Date modified", "Folder Path"],
                   [["Quarterly report.pdf", "2026-01-01", "https://x/a/"]])
    rows = si.read_inventory(p)
    assert rows[0]["ext"] == ".pdf"
    assert rows[0]["filename"] == "Quarterly report.pdf"


def test_blank_extension_cell_falls_back_to_the_name_suffix(tmp_path):
    p = _xlsx_with(tmp_path, ["Name", "Extension", "Folder Path"],
                   [["Quarterly report.pdf", None, "https://x/a/"],
                    ["Notes.docx", "", "https://x/a/"]])
    rows = si.read_inventory(p)
    assert [r["ext"] for r in rows] == [".pdf", ".docx"]
    assert rows[0]["filename"] == "Quarterly report.pdf"


def test_present_extension_wins_and_is_not_doubled(tmp_path):
    p = _xlsx_with(tmp_path, ["Name", "Extension", "Folder Path"],
                   [["Quarterly report.pdf", ".pdf", "https://x/a/"]])
    rows = si.read_inventory(p)
    assert rows[0]["ext"] == ".pdf"
    assert rows[0]["filename"] == "Quarterly report.pdf"
