"""Marts config policies (brain 0.12): collision merge, entity disambiguation,
sheet aliases, several globs per family. Synthetic workbooks, invented names."""
from __future__ import annotations

import json
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent.parent / "skills" / "tabular-semantic-layer"
sys.path.insert(0, str(HERE))

try:
    import openpyxl
    import pandas as pd  # noqa: F401
    import build_marts as B
    _DEPS = True
except Exception:  # pragma: no cover
    _DEPS = False


def _book(path: Path, sheets: dict[str, list[tuple]]):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(list(r))
    wb.save(path)


RANKER = {"name": "ranker", "glob": "*Ranker*.xlsx", "month_from": "filename", "layout": "tolerant_long",
          "dim_candidates": {"branch": ["Branch Name"]},
          "measures": {"nps": "Overall NPS", "n_records": "# of records"}}


@unittest.skipUnless(_DEPS, "requires pandas + openpyxl")
class PolicyBuild(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.root = self.tmp / "root"
        self.root.mkdir()
        self.out = self.tmp / "out"

    def build(self, families, strict=False, ok=True):
        cfg = self.tmp / "families.json"
        cfg.write_text(json.dumps({"families": families, "conform_titlecase": ["branch"]}))
        cmd = [sys.executable, str(HERE / "build_marts.py"), "--root", str(self.root),
               "--config", str(cfg), "--out-dir", str(self.out)] + (["--strict"] if strict else [])
        r = subprocess.run(cmd, capture_output=True, text=True)
        if ok:
            self.assertEqual(r.returncode, 0, r.stderr)
        self.r = r
        if not (self.out / "build_audit.json").exists():
            return None
        con = sqlite3.connect(self.out / "knowledge.sqlite")
        con.row_factory = sqlite3.Row
        self.audit = json.loads((self.out / "build_audit.json").read_text())
        return con

    def fact(self, con, metric, entity):
        return con.execute("SELECT value FROM facts WHERE metric=? AND entity=?", (metric, entity)).fetchone()[0]

    # ---- 1. collision policy -------------------------------------------------------
    def _ranker(self):
        _book(self.root / "April Ranker.xlsx", {"Sheet0": [
            ("Rank", "Branch Name", "# of records", "Overall NPS"),
            (1, "NORTHFIELD", 203, 18.7),
            (2, "EASTON", 40, 10),
            (3, "NORTHFIELD", 1, 100),          # stray row under the same name
        ]})

    def test_default_keeps_last_and_reports(self):
        self._ranker()
        con = self.build([RANKER])
        self.assertEqual(self.fact(con, "nps", "Northfield"), 100.0)
        c = [x for x in self.audit["collisions"] if x["metric"] == "nps"][0]
        self.assertEqual((c["policy"], c["values"]), ("last", [18.7, 100.0]))

    def test_weighted_mean_and_sum(self):
        self._ranker()
        fam = dict(RANKER, collision={"policy": "weighted_mean", "weight": "n_records", "sum": ["n_records"]})
        con = self.build([fam])
        self.assertAlmostEqual(self.fact(con, "nps", "Northfield"), (203 * 18.7 + 100) / 204)
        self.assertEqual(self.fact(con, "n_records", "Northfield"), 204.0)
        self.assertEqual(self.fact(con, "nps", "Easton"), 10.0)              # untouched
        m = {r["metric"]: dict(r) for r in con.execute("SELECT * FROM fact_merges")}
        self.assertEqual((m["nps"]["policy"], m["nps"]["n_rows"]), ("weighted_mean", 2))
        self.assertEqual(json.loads(m["nps"]["inputs"]),
                         [{"value": 18.7, "weight": 203.0}, {"value": 100.0, "weight": 1.0}])
        self.assertEqual(m["n_records"]["policy"], "sum")
        self.assertEqual(self.audit["summary"]["collisions_merged"], 2)

    def test_weighted_mean_skips_rows_without_value_or_weight(self):
        _book(self.root / "April Ranker.xlsx", {"Sheet0": [
            ("Branch Name", "# of records", "Overall NPS"),
            ("NORTHFIELD", 10, 50), ("NORTHFIELD", 5, None), ("NORTHFIELD", None, -100)]})
        fam = dict(RANKER, collision={"policy": "weighted_mean", "weight": "n_records"})
        con = self.build([fam])
        self.assertEqual(self.fact(con, "nps", "Northfield"), 50.0)

    def test_error_policy_fails_strict(self):
        self._ranker()
        fam = dict(RANKER, collision={"policy": "error"})
        self.build([fam], strict=True, ok=False)
        self.assertNotEqual(self.r.returncode, 0)
        self.assertIn("COLLISION", self.r.stderr)

    def test_unknown_policy_is_refused(self):
        self._ranker()
        self.build([dict(RANKER, collision={"policy": "median"})], ok=False)
        self.assertNotEqual(self.r.returncode, 0)
        self.assertIn("median", self.r.stderr)

    # ---- 2. entity disambiguation --------------------------------------------------
    def _repeat(self):
        _book(self.root / "March Repeat.xlsx", {
            "Branch Summary": [("Division", "Branch Name", "State", "% Repeat"),
                               ("West", "Riverton", "PA", 0.2), ("West", "RIVERTON  ", "OH", 0.08),
                               ("West", "EASTON", "PA", 0.1), ("West", "Millbrook", "IL", 0.3)],
            "Branch Mapping": [("Code", "Branch Name"), ("RIVERTON OH", "RIVERTON  "), ("RIVERTON PA", "Riverton"),
                               ("EASTON PA", "EASTON"),
                               # the mapping sheet lists most names under two spellings;
                               # only names that collide in the DATA sheet may be mapped
                               ("MILLBROOK IL", "MILLBROOK "), ("MILLBROOK IL", "Millbrook")]})
        return {"name": "repeat", "glob": "*Repeat*.xlsx", "month_from": "filename", "layout": "long",
                "grains": {"branch": {"sheet": "Branch Summary", "dim_header": "Branch Name"}},
                "measures": {"repeat_pct": "% Repeat"}}

    def test_entity_map_splits_only_ambiguous_names(self):
        fam = self._repeat()
        fam["grains"]["branch"]["entity_map"] = {"sheet": "Branch Mapping", "from": "Branch Name", "to": "Code"}
        con = self.build([fam], strict=False)
        ents = {r[0]: r[1] for r in con.execute("SELECT entity, value FROM facts")}
        # Easton and Millbrook keep their plain names (still join other families)
        self.assertEqual(ents, {"Riverton PA": 0.2, "Riverton OH": 0.08, "Easton": 0.1, "Millbrook": 0.3})
        self.assertEqual(self.audit["collisions"], [])

    def test_entity_map_always(self):
        fam = self._repeat()
        fam["grains"]["branch"]["entity_map"] = {"sheet": "Branch Mapping", "from": "Branch Name",
                                                 "to": "Code", "when": "always"}
        con = self.build([fam])
        self.assertIn("Easton PA", {r[0] for r in con.execute("SELECT entity FROM facts")})

    def test_entity_with_extra_column(self):
        fam = self._repeat()
        fam["grains"]["branch"]["entity_with"] = ["State"]
        con = self.build([fam])
        self.assertEqual({r[0] for r in con.execute("SELECT entity FROM facts")},
                         {"Riverton PA", "Riverton OH", "Easton PA", "Millbrook IL"})

    def test_entity_map_reports_a_name_mapped_differently_across_files(self):
        fam = self._repeat()
        fam["glob"] = "*Repeat*.xlsx"
        fam["grains"]["branch"]["entity_map"] = {"sheet": "Branch Mapping", "from": "Branch Name", "to": "Code"}
        _book(self.root / "April Repeat.xlsx", {
            "Branch Summary": [("Branch Name", "% Repeat"), ("Riverton", 0.3), ("RIVERTON  ", 0.1)],
            "Branch Mapping": [("Code", "Branch Name"), ("RIVERTON KY", "Riverton"), ("RIVERTON OH", "RIVERTON  ")]})
        self.build([fam])
        inc = self.audit["entity_map_inconsistent"]
        self.assertEqual([(i["family"], i["raw"], sorted(i["targets"])) for i in inc],
                         [("repeat", "Riverton", ["RIVERTON KY", "RIVERTON PA"])])

    # ---- 3. sheet aliases ------------------------------------------------------------
    def test_sheet_accepts_alternatives(self):
        fam = self._repeat()
        (self.root / "March Repeat.xlsx").unlink()
        _book(self.root / "March Repeat.xlsx", {"Branches": [("Branch Name", "% Repeat"), ("EASTON", 0.1)]})
        fam["grains"]["branch"]["sheet"] = ["Branch Summary", "Branches"]
        con = self.build([fam])
        self.assertEqual(self.fact(con, "repeat_pct", "Easton"), 0.1)

    def test_get_sheet_prefers_exact_over_contains(self):
        wb = openpyxl.Workbook()
        wb.active.title = "MOM old"
        wb.create_sheet("Month over Month")
        self.assertEqual(B.get_sheet(wb, ["Month over Month", "MOM"]).title, "Month over Month")
        self.assertEqual(B.get_sheet(wb, ["Nope", "MOM"]).title, "MOM old")
        self.assertIsNone(B.get_sheet(wb, ["Nope"]))

    # ---- 4. several globs --------------------------------------------------------------
    def test_glob_list_orders_by_vintage_across_patterns(self):
        (self.root / "sub").mkdir()
        _book(self.root / "sub" / "June 2026 Ranker.xlsx", {"S": [("Branch Name", "Overall NPS"), ("EASTON", 1)]})
        _book(self.root / "July 2026 Ranker.xlsx", {"S": [("Branch Name", "Overall NPS"), ("EASTON", 2)]})
        fam = dict(RANKER, glob=["sub/*Ranker*.xlsx", "*Ranker*.xlsx"], month_from="filename")
        con = self.build([fam])
        rows = {r[0]: r[1] for r in con.execute("SELECT month, value FROM facts WHERE metric='nps'")}
        self.assertEqual(rows, {"2026-06": 1.0, "2026-07": 2.0})
        self.assertEqual(len(self.audit["audit"]), 2)      # each file read once


if __name__ == "__main__":
    unittest.main()
